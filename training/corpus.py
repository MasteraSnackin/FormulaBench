"""Build and verify the development-only FormulaBench SFT corpus.

This module deliberately verifies the frozen public split from ``dataset.json``
and every initial workbook before it locates a golden workbook.  Once that
check succeeds, only development-task goldens may be opened.  The generated
JSONL and manifest are deterministic and hash-bound to their source files.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import math
import os
import tempfile
from collections import defaultdict
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from experiments.make_public_split import (
    build_public_split,
    initial_workbooks_aggregate_sha256,
    largest_remainder_allocation,
    target_size_proxy_bucket,
)
from formulabench.artifacts import DatasetTask, confined_path, load_dataset_manifest, sha256_file
from formulabench.contract import (
    SpreadsheetResponse,
    build_target_contract,
    parse_response,
    validate_response,
    write_response_atomic,
)
from formulabench.prompts import SYSTEM_PROMPT
from formulabench.provider import ANSWER_TOOL_NAME, QWEN_ANSWER_TOOL
from formulabench.runner import _build_prompt

SCHEMA_VERSION = 1
CORPUS_BUILDER = "FormulaBench/sft-corpus/v1"
DEFAULT_MAX_TARGET_CELLS = 500
VALIDATION_FRACTION = 0.20
PARTITION_SEED = "FormulaBench/sft-partition/v1"
CORPUS_NAME = "corpus.jsonl"
MANIFEST_NAME = "corpus_manifest.json"
ALLOWED_OUTPUTS = frozenset({CORPUS_NAME, MANIFEST_NAME})


class CorpusError(ValueError):
    """A reproducibility, leakage, or corpus-contract failure."""


def _canonical_compact(payload: object) -> str:
    return json.dumps(payload, ensure_ascii=False, separators=(",", ":"), sort_keys=True)


def _canonical_document(payload: object) -> str:
    return json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n"


def _sha256_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def _atomic_write(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        dir=path.parent,
        prefix=f".{path.name}.",
        suffix=".tmp",
    )
    temporary = Path(temporary_name)
    try:
        with os.fdopen(descriptor, "wb") as handle:
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _safe_output_directory(path: Path) -> Path:
    if path.is_symlink():
        raise CorpusError("output directory must not be a symbolic link")
    path.mkdir(parents=True, exist_ok=True)
    unexpected = sorted(entry.name for entry in path.iterdir() if entry.name not in ALLOWED_OUTPUTS)
    if unexpected:
        raise CorpusError("output directory contains unexpected entries: " + ", ".join(unexpected))
    for name in ALLOWED_OUTPUTS:
        candidate = path / name
        if candidate.is_symlink():
            raise CorpusError(f"output file must not be a symbolic link: {name}")
    return path.resolve(strict=True)


def _load_and_verify_split(dataset_dir: Path, split_manifest: Path) -> dict[str, Any]:
    if split_manifest.is_symlink() or not split_manifest.is_file():
        raise CorpusError("split manifest must be a regular non-symlink file")
    try:
        declared = json.loads(split_manifest.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise CorpusError("split manifest is not valid UTF-8 JSON") from exc
    if not isinstance(declared, dict):
        raise CorpusError("split manifest root must be an object")

    # This reproduction opens only dataset.json and initial workbooks.  No
    # golden path is resolved until the exact split payload has matched.
    reproduced = build_public_split(dataset_dir)
    if declared != reproduced:
        raise CorpusError("split manifest does not reproduce from the supplied public dataset")
    development = declared.get("development_ids")
    held_out = declared.get("held_out_ids")
    final_only = declared.get("final_only_ids")
    if not all(isinstance(ids, list) for ids in (development, held_out, final_only)):
        raise CorpusError("split manifest is missing its three task-ID lists")
    if len(development) != 80 or len(set(development)) != 80:
        raise CorpusError("development split must contain exactly 80 unique task IDs")
    if set(development).intersection(held_out) or set(development).intersection(final_only):
        raise CorpusError("development split overlaps a held-out split")
    return declared


def _golden_path(dataset_dir: Path, task: DatasetTask) -> Path:
    folder = confined_path(dataset_dir, task.spreadsheet_path, reject_symlinks=True)
    exact = folder / "golden.xlsx"
    candidates = (
        [exact] if exact.exists() or exact.is_symlink() else sorted(folder.glob("*_golden.xlsx"))
    )
    if len(candidates) != 1:
        raise CorpusError(f"development task {task.id} must have exactly one golden workbook")
    candidate = candidates[0]
    try:
        resolved = confined_path(
            dataset_dir,
            candidate.relative_to(dataset_dir),
            reject_symlinks=True,
        )
    except (ValueError, OSError) as exc:
        raise CorpusError(f"development task {task.id} has an unsafe golden workbook") from exc
    if not resolved.is_file():
        raise CorpusError(f"development task {task.id} golden workbook is not a regular file")
    return resolved


def _target_coordinates(
    task: DatasetTask,
    initial_workbook: Any,
    golden_workbook: Any,
) -> list[tuple[str, str]]:
    contract = build_target_contract(task.as_dict(), initial_workbook)
    coordinates: dict[tuple[str, str], None] = {}
    for target in contract.ranges:
        if target.sheet not in golden_workbook.sheetnames:
            raise CorpusError(
                f"development task {task.id} golden workbook lacks target sheet {target.sheet!r}"
            )
        worksheet = golden_workbook[target.sheet]
        min_column, min_row, max_column, max_row = range_boundaries(target.cell_range)
        min_column = min_column or 1
        max_column = max_column or worksheet.max_column
        min_row = min_row or 1
        if max_row is None:
            initial_extent = (
                initial_workbook[target.sheet].max_row
                if target.sheet in initial_workbook.sheetnames
                else 1
            )
            max_row = max(initial_extent, worksheet.max_row)
        for row in range(int(min_row), int(max_row) + 1):
            for column in range(int(min_column), int(max_column) + 1):
                coordinates.setdefault((target.sheet, f"{get_column_letter(column)}{row}"), None)
    if not coordinates:
        raise CorpusError(f"development task {task.id} has no resolved answer cells")
    return list(coordinates)


def _json_cell_value(value: object) -> object:
    if isinstance(value, ArrayFormula):
        if not isinstance(value.text, str) or not value.text.startswith("="):
            raise CorpusError("array formula has no serialisable formula text")
        return value.text
    if isinstance(value, DataTableFormula):
        raise CorpusError("data-table formulas are not representable by the response contract")
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float):
        if not math.isfinite(value):
            raise CorpusError("non-finite target number is not representable")
        return value
    if isinstance(value, dt.datetime):
        if value.tzinfo is not None or value.microsecond % 1_000:
            raise CorpusError("target datetime is not an exact timezone-free millisecond value")
        return {"type": "datetime", "value": value.isoformat(timespec="milliseconds")}
    if isinstance(value, dt.date):
        return {"type": "date", "value": value.isoformat()}
    if isinstance(value, dt.time):
        raise CorpusError("time-only target values are not representable by the response contract")
    raise CorpusError(f"unsupported target value type: {type(value).__name__}")


def _answer_payload(
    task: DatasetTask,
    initial_workbook: Any,
    golden_workbook: Any,
    coordinates: Sequence[tuple[str, str]],
) -> dict[str, object]:
    cells: list[dict[str, object]] = []
    for sheet, coordinate in coordinates:
        cell = golden_workbook[sheet][coordinate]
        value = None if isinstance(cell, MergedCell) else cell.value
        cells.append({"sheet": sheet, "cell": coordinate, "value": _json_cell_value(value)})

    coordinate_set = {
        (sheet, *coordinate_to_tuple(coordinate)) for sheet, coordinate in coordinates
    }
    unmerge_ranges: list[dict[str, str]] = []
    for worksheet in initial_workbook.worksheets:
        golden_merges = (
            {str(item) for item in golden_workbook[worksheet.title].merged_cells.ranges}
            if worksheet.title in golden_workbook.sheetnames
            else set()
        )
        ordered_merges = sorted(
            worksheet.merged_cells.ranges,
            key=lambda item: (
                range_boundaries(str(item))[1::-1] + range_boundaries(str(item))[3:1:-1]
            ),
        )
        for merged in ordered_merges:
            merged_text = str(merged)
            if merged_text in golden_merges:
                continue
            min_column, min_row, max_column, max_row = range_boundaries(merged_text)
            if any(
                sheet == worksheet.title
                and min_row <= row <= max_row
                and min_column <= column <= max_column
                for sheet, row, column in coordinate_set
            ):
                unmerge_ranges.append({"sheet": worksheet.title, "range": merged_text})
    payload: dict[str, object] = {
        "cells": cells,
        "fills": [],
        "preserve_ranges": [],
        "unmerge_ranges": unmerge_ranges,
    }
    parsed = parse_response(payload)
    if not isinstance(parsed, SpreadsheetResponse):  # pragma: no cover - defensive type check
        raise CorpusError("answer did not parse as a SpreadsheetResponse")
    contract = build_target_contract(task.as_dict(), initial_workbook)
    validation = validate_response(contract, parsed, workbook=initial_workbook)
    if not validation.ok:
        failures = ",".join(code.value for code in validation.failure_codes)
        raise CorpusError(f"golden answer fails target coverage: {failures}")
    return payload


def _runtime_check(task: DatasetTask, payload: Mapping[str, object]) -> None:
    with tempfile.TemporaryDirectory(prefix="formulabench-corpus-check-") as temporary:
        result = write_response_atomic(
            task.as_dict(),
            payload,
            task.init_xlsx,
            Path(temporary) / "candidate.xlsx",
        )
    if not result.success:
        failures = ",".join(code.value for code in result.failure_codes)
        raise CorpusError(f"golden answer fails the runtime contract: {failures}")


def _partition_rows(rows: list[dict[str, object]]) -> None:
    if not rows:
        raise CorpusError("no development examples passed the corpus gates")
    validation_count = round(len(rows) * VALIDATION_FRACTION)
    validation_count = max(1, min(len(rows) - 1, validation_count)) if len(rows) > 1 else 0

    strata: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    for row in rows:
        instruction_type = str(row["instruction_type"])
        target_bucket = target_size_proxy_bucket(int(row["target_cell_count"]))
        strata[(instruction_type, target_bucket)].append(row)
    quotas = largest_remainder_allocation(
        {stratum: len(items) for stratum, items in strata.items()},
        validation_count,
        split_name="sft_validation",
    )
    validation_ids: set[str] = set()
    for stratum, items in sorted(strata.items()):
        ranked = sorted(
            items,
            key=lambda row: (
                hashlib.sha256(
                    "\0".join(
                        (
                            PARTITION_SEED,
                            stratum[0],
                            stratum[1],
                            str(row["task_id"]),
                        )
                    ).encode("utf-8")
                ).digest(),
                str(row["task_id"]),
            ),
        )
        validation_ids.update(str(row["task_id"]) for row in ranked[: quotas[stratum]])
    for row in rows:
        row["partition"] = "validation" if str(row["task_id"]) in validation_ids else "train"


def build_corpus(
    *,
    dataset_dir: Path,
    split_manifest: Path,
    out_dir: Path,
    max_target_cells: int = DEFAULT_MAX_TARGET_CELLS,
) -> dict[str, Any]:
    """Build a deterministic, hash-bound corpus and return its manifest."""

    if isinstance(max_target_cells, bool) or max_target_cells < 1:
        raise CorpusError("max target cells must be a positive integer")
    dataset_dir = dataset_dir.resolve(strict=True)
    split_manifest = split_manifest.resolve(strict=True)
    split = _load_and_verify_split(dataset_dir, split_manifest)
    output = _safe_output_directory(out_dir)
    split_manifest_sha256 = sha256_file(split_manifest)
    dataset_manifest_sha256 = str(split["dataset_manifest_sha256"])

    tasks = load_dataset_manifest(dataset_dir)
    tasks_by_id = {task.id: task for task in tasks}
    development_ids = [str(task_id) for task_id in split["development_ids"]]
    if any(task_id not in tasks_by_id for task_id in development_ids):
        raise CorpusError("development split refers to a task outside dataset.json")

    candidate_rows: list[dict[str, object]] = []
    exclusions: list[dict[str, object]] = []
    initial_hashes: list[tuple[str, str]] = []
    golden_hashes: list[tuple[str, str]] = []

    for task_id in development_ids:
        task = tasks_by_id[task_id]
        golden_path = _golden_path(dataset_dir, task)
        initial_sha256 = sha256_file(task.init_xlsx)
        golden_sha256 = sha256_file(golden_path)
        initial_hashes.append((task.id, initial_sha256))
        golden_hashes.append((task.id, golden_sha256))

        initial_workbook = openpyxl.load_workbook(
            task.init_xlsx,
            data_only=False,
            read_only=False,
            keep_links=False,
        )
        golden_workbook = openpyxl.load_workbook(
            golden_path,
            data_only=False,
            read_only=False,
            keep_links=False,
        )
        try:
            coordinates = _target_coordinates(task, initial_workbook, golden_workbook)
            target_count = len(coordinates)
            if target_count > max_target_cells:
                exclusions.append(
                    {
                        "reason": "target_cell_limit_exceeded",
                        "task_id": task.id,
                        "target_cell_count": target_count,
                    }
                )
                continue
            answer = _answer_payload(task, initial_workbook, golden_workbook, coordinates)
            _runtime_check(task, answer)
            user_prompt, _context = _build_prompt(task)
        finally:
            initial_workbook.close()
            golden_workbook.close()

        assistant_response = _canonical_compact(answer)
        candidate_rows.append(
            {
                "answer_cell_count": target_count,
                "instruction_type": str(task.get("instruction_type", "")),
                "messages": [
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": user_prompt},
                    {"role": "assistant", "content": assistant_response},
                ],
                "provenance": {
                    "answer_sha256": _sha256_text(assistant_response),
                    "dataset_manifest_sha256": dataset_manifest_sha256,
                    "golden_workbook_sha256": golden_sha256,
                    "initial_workbook_sha256": initial_sha256,
                    "prompt_sha256": _sha256_text(user_prompt),
                    "split_manifest_sha256": split_manifest_sha256,
                },
                "schema_version": SCHEMA_VERSION,
                "split": "development",
                "target_cell_count": target_count,
                "task_id": task.id,
            }
        )

    _partition_rows(candidate_rows)
    for row in candidate_rows:
        del row["instruction_type"]
    rows = list(candidate_rows)
    corpus_text = "".join(_canonical_compact(row) + "\n" for row in rows)
    corpus_sha256 = _sha256_text(corpus_text)
    partition_counts = {
        "train": sum(row["partition"] == "train" for row in rows),
        "validation": sum(row["partition"] == "validation" for row in rows),
    }
    example_records: list[dict[str, object]] = []
    for row in rows:
        provenance = row["provenance"]
        if not isinstance(provenance, Mapping):  # pragma: no cover - local invariant
            raise AssertionError("row provenance is not a mapping")
        example_records.append(
            {
                "answer_cell_count": row["answer_cell_count"],
                "answer_sha256": provenance["answer_sha256"],
                "golden_workbook_sha256": provenance["golden_workbook_sha256"],
                "initial_workbook_sha256": provenance["initial_workbook_sha256"],
                "partition": row["partition"],
                "prompt_sha256": provenance["prompt_sha256"],
                "row_sha256": _sha256_text(_canonical_compact(row)),
                "target_cell_count": row["target_cell_count"],
                "task_id": row["task_id"],
            }
        )
    manifest: dict[str, Any] = {
        "builder": CORPUS_BUILDER,
        "corpus_file": CORPUS_NAME,
        "corpus_sha256": corpus_sha256,
        "dataset_manifest_sha256": dataset_manifest_sha256,
        "development_only": True,
        "development_golden_binding_sha256": initial_workbooks_aggregate_sha256(golden_hashes),
        "development_golden_records": [
            {"task_id": task_id, "sha256": digest} for task_id, digest in golden_hashes
        ],
        "development_initial_records": [
            {"task_id": task_id, "sha256": digest} for task_id, digest in initial_hashes
        ],
        "eligible_development_count": len(rows),
        "example_count": len(rows),
        "examples": example_records,
        "excluded_development_count": len(exclusions),
        "exclusions": exclusions,
        "initial_workbook_binding_sha256": split["initial_workbook_binding"]["aggregate_sha256"],
        "max_target_cells": max_target_cells,
        "ordered_task_ids": [row["task_id"] for row in rows],
        "partition_counts": partition_counts,
        "partition_method": "stratified-largest-remainder-sha256",
        "partition_seed": PARTITION_SEED,
        "schema_version": SCHEMA_VERSION,
        "split": "development",
        "split_manifest_sha256": split_manifest_sha256,
        "split_seed": str(split["algorithm"]["seed"]),
        "system_prompt_sha256": _sha256_text(SYSTEM_PROMPT),
        "training_contract": {
            "assistant_format": "qwen3.8 submit_spreadsheet_answer tool arguments",
            "renderer": "qwen3_8_disable_thinking",
            "system_prompt_sha256": _sha256_text(SYSTEM_PROMPT),
            "tool_name": ANSWER_TOOL_NAME,
            "tool_schema_sha256": _sha256_text(_canonical_compact(QWEN_ANSWER_TOOL)),
        },
        "verified": True,
    }
    _atomic_write(output / CORPUS_NAME, corpus_text.encode("utf-8"))
    _atomic_write(output / MANIFEST_NAME, _canonical_document(manifest).encode("utf-8"))
    return manifest


def validate_corpus(
    *,
    dataset_dir: Path,
    split_manifest: Path,
    corpus: Path,
    manifest_path: Path,
) -> dict[str, Any]:
    """Rebuild the expected corpus and reject any byte or manifest difference."""

    for label, path in (("corpus", corpus), ("manifest", manifest_path)):
        if path.is_symlink() or not path.is_file():
            raise CorpusError(f"{label} must be a regular non-symlink file")
    try:
        supplied_manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise CorpusError("corpus manifest is not valid UTF-8 JSON") from exc
    if not isinstance(supplied_manifest, dict):
        raise CorpusError("corpus manifest root must be an object")
    if supplied_manifest.get("verified") is not True:
        raise CorpusError("corpus manifest is not marked verified")
    if supplied_manifest.get("development_only") is not True:
        raise CorpusError("corpus manifest is not development-only")
    max_target_cells = supplied_manifest.get("max_target_cells")
    if isinstance(max_target_cells, bool) or not isinstance(max_target_cells, int):
        raise CorpusError("corpus manifest has an invalid target-cell limit")

    with tempfile.TemporaryDirectory(prefix="formulabench-corpus-rebuild-") as temporary:
        expected_manifest = build_corpus(
            dataset_dir=dataset_dir,
            split_manifest=split_manifest,
            out_dir=Path(temporary),
            max_target_cells=max_target_cells,
        )
        expected_corpus = (Path(temporary) / CORPUS_NAME).read_bytes()
    if corpus.read_bytes() != expected_corpus:
        raise CorpusError("corpus bytes do not reproduce from the verified development split")
    if supplied_manifest != expected_manifest:
        raise CorpusError("corpus manifest does not reproduce from the verified development split")
    if supplied_manifest["corpus_sha256"] != hashlib.sha256(expected_corpus).hexdigest():
        raise CorpusError("corpus SHA-256 does not match its bytes")
    return supplied_manifest


def _summary(manifest: Mapping[str, Any]) -> dict[str, object]:
    return {
        "corpus_sha256": manifest["corpus_sha256"],
        "excluded": len(manifest["exclusions"]),
        "partitions": manifest["partition_counts"],
        "rows": manifest["example_count"],
        "verified": manifest["verified"],
    }


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    build = subparsers.add_parser("build", help="build a verified development-only corpus")
    build.add_argument("--dataset-dir", required=True, type=Path)
    build.add_argument("--split-manifest", required=True, type=Path)
    build.add_argument("--out-dir", required=True, type=Path)
    build.add_argument("--max-target-cells", type=int, default=DEFAULT_MAX_TARGET_CELLS)

    validate = subparsers.add_parser("validate", help="rebuild and verify a corpus")
    validate.add_argument("--dataset-dir", required=True, type=Path)
    validate.add_argument("--split-manifest", required=True, type=Path)
    validate.add_argument("--corpus", required=True, type=Path)
    validate.add_argument("--manifest", required=True, type=Path)
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        if args.command == "build":
            manifest = build_corpus(
                dataset_dir=args.dataset_dir,
                split_manifest=args.split_manifest,
                out_dir=args.out_dir,
                max_target_cells=args.max_target_cells,
            )
        else:
            manifest = validate_corpus(
                dataset_dir=args.dataset_dir,
                split_manifest=args.split_manifest,
                corpus=args.corpus,
                manifest_path=args.manifest,
            )
    except (CorpusError, OSError, ValueError) as exc:
        raise SystemExit(f"corpus error: {exc}") from exc
    print(_canonical_compact(_summary(manifest)))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())


__all__ = [
    "CORPUS_NAME",
    "DEFAULT_MAX_TARGET_CELLS",
    "MANIFEST_NAME",
    "CorpusError",
    "build_corpus",
    "main",
    "validate_corpus",
]
