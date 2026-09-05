from __future__ import annotations

import datetime as dt
import hashlib
import json
import os
import re
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.formula import ArrayFormula

from formulabench.context import (
    ContextBudgetError,
    build_context,
    parse_data_positions,
)


def _records(document: str) -> list[dict]:
    return [json.loads(line) for line in document.splitlines()]


def _rewrite_formula_cache(path: Path, *, formula: str, cached_value: str) -> None:
    """Install a cached value that openpyxl intentionally does not write."""

    replacement_made = False
    rewritten = path.with_suffix(".rewritten.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(rewritten, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename.startswith("xl/worksheets/sheet") and info.filename.endswith(".xml"):
                formula_xml = formula.removeprefix("=").encode()
                pattern = b"<f>" + re.escape(formula_xml) + rb"</f><v(?:\s*/>|>\s*</v>)"
                replacement = b"<f>" + formula_xml + b"</f><v>" + cached_value.encode() + b"</v>"
                payload, replacements = re.subn(pattern, replacement, payload, count=1)
                replacement_made = replacement_made or bool(replacements)
            target.writestr(info, payload)
    assert replacement_made, "test workbook did not contain the expected formula XML"
    os.replace(rewritten, path)


@pytest.fixture
def evidence_workbook(tmp_path: Path) -> Path:
    path = tmp_path / "evidence.xlsx"
    workbook = openpyxl.Workbook()
    cash = workbook.active
    cash.title = "Cash, Inc"
    cash["A1"] = 2
    cash["A2"] = 3
    cash["Z200"] = "=SUM(A1:A2)"
    cash["Z201"] = None
    cash.merge_cells("B4:C4")
    cash["B4"] = "merged header"

    apostrophe = workbook.create_sheet("O'Brien")
    apostrophe["A1"] = ArrayFormula(ref="A1:A2", text="=ROW(A1:A2)")
    apostrophe["C3"] = "array source"
    workbook.save(path)
    workbook.close()
    _rewrite_formula_cache(path, formula="=SUM(A1:A2)", cached_value="5")
    return path


def _task(path: Path) -> dict:
    return {
        "id": "synthetic-1",
        "instruction_type": "Cell-Level Manipulation",
        "instruction": "Set Z201 using the existing calculation in Z200 and inspect O'Brien!A1:A2.",
        "answer_position": "Z201",
        "answer_sheet": "Cash, Inc",
        "data_position": "'Cash, Inc'!A1:A2','O''Brien'!A1:C3",
        "init_xlsx": str(path),
        "golden_xlsx": str(path.parent / "must-not-be-read-golden.xlsx"),
    }


@pytest.fixture
def source_output_workbook(tmp_path: Path) -> Path:
    path = tmp_path / "source-output.xlsx"
    workbook = openpyxl.Workbook()
    source = workbook.active
    source.title = "Source Data"
    source["A1"] = "alpha"
    source["B1"] = 7
    source["A2"] = dt.date(2024, 1, 2)
    source["B2"] = '=A1&"!"'
    output = workbook.create_sheet("Output")
    output["D4"] = "replace me"
    workbook.active = 1
    workbook.save(path)
    workbook.close()
    return path


def _source_output_task(path: Path) -> dict:
    return {
        "id": "source-output",
        "instruction_type": "Sheet-Level Manipulation",
        "instruction": "Copy the values from the 'Source Data' sheet into the Output sheet.",
        "answer_position": "D4",
        "answer_sheet": "Output",
        "data_position": "A1:B2",
        "init_xlsx": str(path),
    }


def test_data_position_parser_preserves_commas_and_apostrophes() -> None:
    refs = parse_data_positions(
        "'Cash, Inc'!A1:B2','O''Brien'!C3:C4,Plain!'A1: D2",
        sheet_names=("Cash, Inc", "O'Brien", "Plain"),
        default_sheet="Plain",
    )

    assert [(ref.sheet, ref.a1_range, ref.resolution) for ref in refs] == [
        ("Cash, Inc", "A1:B2", "explicit"),
        ("O'Brien", "C3:C4", "explicit"),
        ("Plain", "A1:D2", "explicit"),
    ]

    prose_ref = parse_data_positions(
        "Inspect O'Brien!C3 before calculating the result.",
        sheet_names=("Cash, Inc", "O'Brien", "Plain"),
        default_sheet="Plain",
    )
    assert [(ref.sheet, ref.a1_range) for ref in prose_ref] == [("O'Brien", "C3")]


def test_unqualified_data_uses_unique_instruction_named_non_answer_sheet(
    source_output_workbook: Path,
) -> None:
    document = build_context(_source_output_task(source_output_workbook), char_budget=6_000)
    records = _records(document.jsonl)
    data = next(record for record in records if record["type"] == "data")

    assert data == {
        "cell_count": 4,
        "declared": True,
        "dynamic": False,
        "range": "A1:B2",
        "requested_sheet": None,
        "resolution": "instruction_named_default",
        "sheet": "Source Data",
        "type": "data",
    }


def test_ambiguous_instruction_named_data_sheets_keep_answer_fallback(tmp_path: Path) -> None:
    path = tmp_path / "ambiguous-source.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Source One"
    workbook.create_sheet("Source Two")
    workbook.create_sheet("Output")
    workbook.active = 2
    workbook.save(path)
    workbook.close()
    task = {
        "id": "ambiguous-source",
        "instruction_type": "Sheet-Level Manipulation",
        "instruction": "Combine Source One and Source Two in Output.",
        "answer_position": "C3",
        "answer_sheet": "Output",
        "data_position": "A1:B2",
        "init_xlsx": str(path),
    }

    records = _records(build_context(task, char_budget=6_000).jsonl)
    data = next(record for record in records if record["type"] == "data")

    assert data["sheet"] == "Output"
    assert data["requested_sheet"] is None
    assert data["resolution"] == "default"


def test_explicit_data_sheet_overrides_instruction_named_default(
    source_output_workbook: Path,
) -> None:
    task = _source_output_task(source_output_workbook)
    task["data_position"] = "Output!A1:B2"

    records = _records(build_context(task, char_budget=6_000).jsonl)
    data = next(record for record in records if record["type"] == "data")

    assert data["sheet"] == "Output"
    assert data["requested_sheet"] == "Output"
    assert data["resolution"] == "explicit"


def test_compact_data_rows_preserve_typed_values_without_verbose_duplicates(
    source_output_workbook: Path,
) -> None:
    document = build_context(_source_output_task(source_output_workbook), char_budget=6_000)
    records = _records(document.jsonl)
    rows = [record for record in records if record["type"] == "data_row"]

    assert [(record["range"], record["sheet"]) for record in rows] == [
        ("A1:B1", "Source Data"),
        ("A2:B2", "Source Data"),
    ]
    assert rows[0]["values"] == ["alpha", 7]
    assert rows[1]["values"][0] == {
        "type": "datetime",
        "value": "2024-01-02T00:00:00",
    }
    formula = rows[1]["values"][1]
    assert formula["formula_view"] == {
        "formula_kind": "standard",
        "kind": "formula",
        "text": '=A1&"!"',
    }
    assert formula["cached_value"] == {
        "reason": "no_cached_result",
        "type": "unknown",
        "value": None,
    }
    assert not any(
        record["type"] == "cell" and record["sheet"] == "Source Data" and record["cell"] == "B2"
        for record in records
    )
    assert document.included_cells == 5
    assert document.omitted_cells == 0


def test_compact_answer_rows_preserve_formula_cache_and_typed_values(tmp_path: Path) -> None:
    path = tmp_path / "answer-semantics.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Output"
    sheet["A2"] = dt.date(2024, 6, 15)
    sheet["B2"] = "=40+2"
    workbook.save(path)
    workbook.close()
    _rewrite_formula_cache(path, formula="=40+2", cached_value="42")
    task = {
        "id": "answer-semantics",
        "instruction_type": "Cell-Level Manipulation",
        "instruction": "Complete the Output table.",
        "answer_position": "A2:B2",
        "answer_sheet": "Output",
        "data_position": None,
        "init_xlsx": str(path),
    }

    document = build_context(task, char_budget=5_000)
    records = _records(document.jsonl)
    answer_rows = [record for record in records if record["type"] == "answer_row"]

    assert answer_rows == [
        {
            "range": "A2:B2",
            "sheet": "Output",
            "type": "answer_row",
            "values": [
                {"type": "datetime", "value": "2024-06-15T00:00:00"},
                {
                    "cached_value": 42,
                    "cell_type": "f",
                    "formula_view": {
                        "formula_kind": "standard",
                        "kind": "formula",
                        "text": "=40+2",
                    },
                },
            ],
        }
    ]
    assert not any(record["type"] == "cell" for record in records)
    assert document.included_cells == 2
    assert document.omitted_cells == 0


def test_complete_120_cell_answer_template_precedes_compact_source_rows(tmp_path: Path) -> None:
    path = tmp_path / "template-and-source.xlsx"
    workbook = openpyxl.Workbook()
    source = workbook.active
    source.title = "RANGES"
    for row in range(1, 57):
        for col in range(1, 6):
            source.cell(row=row, column=col, value=row * 10 + col)
    output = workbook.create_sheet("LISTS")
    for row in range(3, 33):
        output.cell(row=row, column=1, value=row - 2)
        output.cell(row=row, column=2, value=dt.date(2024, 1, min(row, 28)))
        output.cell(row=row, column=3, value=f"REF-{row:02d}")
        output.cell(row=row, column=4, value=row * 100)
    workbook.save(path)
    workbook.close()
    task = {
        "id": "complete-template",
        "instruction_type": "Sheet-Level Manipulation",
        "instruction": "Use RANGES to complete the existing LISTS template.",
        "answer_position": "A3:D32",
        "answer_sheet": "LISTS",
        "data_position": "RANGES!A1:E56",
        "init_xlsx": str(path),
    }

    document = build_context(task, char_budget=30_000)
    assert build_context(task, char_budget=30_000) == document
    records = _records(document.jsonl)
    answer_rows = [record for record in records if record["type"] == "answer_row"]
    data_rows = [record for record in records if record["type"] == "data_row"]

    assert len(answer_rows) == 30
    assert answer_rows[0]["range"] == "A3:D3"
    assert answer_rows[-1]["range"] == "A32:D32"
    assert answer_rows[0]["values"] == [
        1,
        {"type": "datetime", "value": "2024-01-03T00:00:00"},
        "REF-03",
        300,
    ]
    assert len(data_rows) == 56
    assert max(records.index(record) for record in answer_rows) < min(
        records.index(record) for record in data_rows
    )
    assert not any(record["type"] == "cell" and record["sheet"] == "LISTS" for record in records)
    assert document.included_cells == 400
    assert document.omitted_cells == 0


def test_compact_answer_falls_back_atomically_under_tight_budget(tmp_path: Path) -> None:
    path = tmp_path / "tight-answer.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Output"
    for row in range(1, 31):
        for col in range(1, 5):
            sheet.cell(row=row, column=col, value=f"value-{row}-{col}")
    workbook.save(path)
    workbook.close()
    task = {
        "id": "tight-answer",
        "instruction_type": "Sheet-Level Manipulation",
        "instruction": "Replace the Output table.",
        "answer_position": "A1:D30",
        "answer_sheet": "Output",
        "data_position": None,
        "init_xlsx": str(path),
    }

    document = build_context(task, char_budget=2_400)
    records = _records(document.jsonl)

    assert not any(record["type"] == "answer_row" for record in records)
    assert any(
        record["type"] == "cell" and "answer_sample" in record["roles"] for record in records
    )
    assert 0 < document.included_cells < 120
    assert document.omitted_cells == 120 - document.included_cells
    assert document.truncated is True
    assert document.used_chars <= 2_400


def test_overlapping_answer_and_data_keep_both_evidence_roles_and_unique_counts(
    tmp_path: Path,
) -> None:
    path = tmp_path / "overlap.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Shared"
    sheet.append(["A", "B"])
    sheet.append([1, "=A2+1"])
    workbook.save(path)
    workbook.close()
    task = {
        "id": "overlap",
        "instruction_type": "Cell-Level Manipulation",
        "instruction": "Rebuild the shared range.",
        "answer_position": "A1:B2",
        "answer_sheet": "Shared",
        "data_position": "Shared!A1:B2",
        "init_xlsx": str(path),
    }

    document = build_context(task, char_budget=6_000)
    records = _records(document.jsonl)

    assert [record["range"] for record in records if record["type"] == "answer_row"] == [
        "A1:B1",
        "A2:B2",
    ]
    assert [record["range"] for record in records if record["type"] == "data_row"] == [
        "A1:B1",
        "A2:B2",
    ]
    assert not any(record["type"] == "cell" for record in records)
    assert document.included_cells == 4
    assert document.omitted_cells == 0


def test_formula_views_include_cached_value_and_compact_array_formula(
    evidence_workbook: Path,
) -> None:
    document = build_context(_task(evidence_workbook), char_budget=12_000)
    records = _records(document.jsonl)
    cells = {
        (record["sheet"], record["cell"]): record for record in records if record["type"] == "cell"
    }

    standard = cells[("Cash, Inc", "Z200")]
    assert standard["formula_view"] == {
        "kind": "formula",
        "formula_kind": "standard",
        "text": "=SUM(A1:A2)",
    }
    assert standard["cached_value"] == {"type": "integer", "value": 5}
    assert "near_answer_formula" in standard["roles"]

    array_row = next(
        record
        for record in records
        if record["type"] == "data_row"
        and record["sheet"] == "O'Brien"
        and record["range"] == "A1:C1"
    )
    array = array_row["values"][0]
    assert array["formula_view"] == {
        "kind": "formula",
        "formula_kind": "array",
        "ref": "A1:A2",
        "text": "=ROW(A1:A2)",
    }
    assert array["cached_value"] == {
        "reason": "no_cached_result",
        "type": "unknown",
        "value": None,
    }
    assert ("O'Brien", "A1") not in cells


def test_relevant_merged_range_and_anchor_are_exposed(evidence_workbook: Path) -> None:
    task = _task(evidence_workbook)
    task["answer_position"] = "B4:C4"
    document = build_context(task, char_budget=12_000)
    sheets = {
        record["sheet"]: record for record in _records(document.jsonl) if record["type"] == "sheet"
    }
    assert sheets["Cash, Inc"]["relevant_merged_ranges"] == [{"anchor": "B4", "range": "B4:C4"}]


def test_tail_formula_regresses_fixed_top_left_crop(evidence_workbook: Path) -> None:
    document = build_context(_task(evidence_workbook), char_budget=8_000)
    cell_coordinates = {
        (record["sheet"], record["cell"])
        for record in _records(document.jsonl)
        if record["type"] == "cell"
    }
    answer_ranges = {
        (record["sheet"], record["range"])
        for record in _records(document.jsonl)
        if record["type"] == "answer_row"
    }

    assert ("Cash, Inc", "Z200") in cell_coordinates
    assert ("Cash, Inc", "Z201:Z201") in answer_ranges


def test_context_is_deterministic_and_does_not_mutate_input(evidence_workbook: Path) -> None:
    task = _task(evidence_workbook)
    sha_before = hashlib.sha256(evidence_workbook.read_bytes()).hexdigest()

    first = build_context(task, char_budget=9_000)
    second = build_context(task, char_budget=9_000)

    assert first == second
    assert hashlib.sha256(evidence_workbook.read_bytes()).hexdigest() == sha_before
    assert first.input_sha256 == sha_before
    assert "must-not-be-read-golden.xlsx" not in first.jsonl


def test_hard_character_budget_and_final_cut_record(evidence_workbook: Path) -> None:
    budget = 2_200
    document = build_context(_task(evidence_workbook), char_budget=budget)
    records = _records(document.jsonl)

    assert len(document.jsonl) == document.used_chars <= budget
    assert records[-1]["type"] == "cut"
    assert records[-1]["used_chars"] == len(document.jsonl)
    assert records[-1]["budget_chars"] == budget
    assert records[-1]["included_cells"] == document.included_cells
    assert records[-1]["omitted_cells"] == document.omitted_cells
    assert records[-1]["truncated"] is True
    assert {record["type"] for record in records} >= {
        "task",
        "workbook",
        "sheet",
        "answer",
        "data",
        "cut",
    }


def test_budget_too_small_for_mandatory_records_is_explicit(evidence_workbook: Path) -> None:
    with pytest.raises(ContextBudgetError, match="mandatory workbook records"):
        build_context(_task(evidence_workbook), char_budget=64)


def test_sparse_huge_range_is_sampled_without_dense_enumeration(tmp_path: Path) -> None:
    path = tmp_path / "sparse-huge.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Sparse"
    sheet["A1"] = "head"
    sheet["XFD1048576"] = "tail"
    workbook.save(path)
    workbook.close()
    task = {
        "id": "sparse",
        "instruction_type": "Cell-Level Manipulation",
        "instruction": "Return XFD1048576.",
        "answer_position": "XFD1048576",
        "answer_sheet": "Sparse",
        "data_position": "A1:XFD1048576",
        "init_xlsx": str(path),
    }

    document = build_context(task, char_budget=6_000)
    records = _records(document.jsonl)
    data_record = next(record for record in records if record["type"] == "data")
    cells = {
        record["cell"]: record
        for record in records
        if record["type"] == "cell" and record["sheet"] == "Sparse"
    }
    answer_rows = {
        record["range"]: record
        for record in records
        if record["type"] == "answer_row" and record["sheet"] == "Sparse"
    }

    assert data_record["cell_count"] == 16_384 * 1_048_576
    assert "A1" in cells
    assert answer_rows["XFD1048576:XFD1048576"]["values"] == ["tail"]


def test_spread_selection_includes_both_ends_under_budget(tmp_path: Path) -> None:
    path = tmp_path / "spread.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in range(1, 301):
        sheet.cell(row=row, column=1, value=f"row-{row}")
    workbook.save(path)
    workbook.close()
    task = {
        "id": "spread",
        "instruction_type": "Cell-Level Manipulation",
        "instruction": "Fill B1 from the supplied data.",
        "answer_position": "B1",
        "answer_sheet": "Sheet",
        "data_position": "A1:A300",
        "init_xlsx": str(path),
    }

    document = build_context(task, char_budget=3_600)
    data_cells = {
        record["cell"]
        for record in _records(document.jsonl)
        if record["type"] == "cell" and "data" in record["roles"]
    }

    assert "A1" in data_cells
    assert "A300" in data_cells
