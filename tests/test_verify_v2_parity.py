from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

import tools.verify_v2_parity as verifier
from exactsource.artifacts import truncate_trace_record
from exactsource.context import build_context
from exactsource.contracts import SetValue, SolvePlan
from exactsource.dataset import load_tasks
from exactsource.plans import apply_operations
from exactsource.prompts import build_messages
from tools.verify_v2_parity import main

PROJECT_ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(autouse=True)
def _canonical_environment(monkeypatch: pytest.MonkeyPatch) -> None:
    """Keep synthetic parity tests independent of the developer environment."""

    original_find_spec = verifier.importlib.util.find_spec

    def canonical_find_spec(name: str):
        return None if name == "PIL" else original_find_spec(name)

    monkeypatch.setattr(verifier.importlib.util, "find_spec", canonical_find_spec)


def _tasks(*task_ids: str) -> list[SimpleNamespace]:
    return [SimpleNamespace(id=task_id) for task_id in task_ids]


def test_canonical_environment_accepts_absent_pillow(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(verifier, "_pillow_is_available", lambda: False)

    verifier._require_canonical_optional_dependencies()


def test_canonical_environment_rejects_present_pillow(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(verifier, "_pillow_is_available", lambda: True)

    with pytest.raises(verifier.ParityError, match="Pillow is installed"):
        verifier._require_canonical_optional_dependencies()


def test_verifier_fails_closed_before_reading_inputs_when_pillow_is_present(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    monkeypatch.setattr(verifier, "_pillow_is_available", lambda: True)

    result = main(
        [
            "--dataset-dir",
            str(tmp_path / "missing-dataset"),
            "--reference-run",
            str(tmp_path / "missing-reference"),
        ]
    )

    captured = capsys.readouterr()
    assert result == 1
    assert "Pillow is installed" in captured.err
    assert "checked=0 accepted=0 fallback=0 mismatches=1" in captured.out


@pytest.mark.parametrize(("spec", "expected"), [(None, False), (object(), True)])
def test_pillow_detection_uses_import_discovery(
    monkeypatch: pytest.MonkeyPatch,
    spec: object | None,
    expected: bool,
) -> None:
    observed: list[str] = []

    def fake_find_spec(name: str):
        observed.append(name)
        return spec

    monkeypatch.setattr(verifier.importlib.util, "find_spec", fake_find_spec)

    assert verifier._pillow_is_available() is expected
    assert observed == ["PIL"]


def test_ids_selection_matches_v2_semantics() -> None:
    tasks = _tasks("first", "second", "third")

    selected = verifier._selected_tasks(tasks, " third, ,first,third,, ")

    assert [task.id for task in selected] == ["first", "third"]
    with pytest.raises(verifier.ParityError, match="at least one"):
        verifier._selected_tasks(tasks, ", ,")
    with pytest.raises(verifier.ParityError, match=r"unknown dataset tasks: \['missing'\]"):
        verifier._selected_tasks(tasks, "missing,first")


def test_full_prediction_coverage_requires_exact_manifest_order() -> None:
    tasks = _tasks("first", "second")

    verifier._validate_prediction_coverage(
        tasks,
        tasks,
        {"first": {}, "second": {}},
        subset=False,
    )
    with pytest.raises(verifier.ParityError, match="exactly match.*manifest in order"):
        verifier._validate_prediction_coverage(
            tasks,
            tasks,
            {"second": {}, "first": {}},
            subset=False,
        )
    with pytest.raises(verifier.ParityError, match="exactly match.*manifest in order"):
        verifier._validate_prediction_coverage(
            tasks,
            tasks,
            {"first": {}, "second": {}, "unexpected": {}},
            subset=False,
        )


def test_subset_prediction_coverage_accepts_subset_or_full_reference() -> None:
    tasks = _tasks("first", "second", "third")
    selected = [tasks[0], tasks[2]]

    verifier._validate_prediction_coverage(
        tasks,
        selected,
        {"first": {}, "third": {}},
        subset=True,
    )
    verifier._validate_prediction_coverage(
        tasks,
        selected,
        {"first": {}, "second": {}, "third": {}},
        subset=True,
    )
    with pytest.raises(verifier.ParityError, match="full manifest or the selected subset"):
        verifier._validate_prediction_coverage(
            tasks,
            selected,
            {"first": {}, "second": {}},
            subset=True,
        )


def test_empty_dataset_is_rejected(tmp_path: Path, capsys: pytest.CaptureFixture[str]) -> None:
    dataset = tmp_path / "empty-dataset"
    dataset.mkdir()
    (dataset / "dataset.json").write_text("[]", encoding="utf-8")
    reference = tmp_path / "reference"
    reference.mkdir()

    result = main(["--dataset-dir", str(dataset), "--reference-run", str(reference)])

    captured = capsys.readouterr()
    assert result == 1
    assert "dataset contains no tasks" in captured.err
    assert "checked=0 accepted=0 fallback=0 mismatches=1" in captured.out


def test_empty_reference_predictions_are_rejected(
    tmp_path: Path,
    capsys: pytest.CaptureFixture[str],
) -> None:
    dataset, _init_path, _golden_path = _make_dataset(tmp_path)
    reference = tmp_path / "empty-reference"
    reference.mkdir()
    (reference / "predictions.jsonl").write_text("", encoding="utf-8")

    result = main(["--dataset-dir", str(dataset), "--reference-run", str(reference)])

    captured = capsys.readouterr()
    assert result == 1
    assert "reference predictions are empty" in captured.err
    assert "checked=0 accepted=0 fallback=0 mismatches=1" in captured.out


def test_verifier_forces_openpyxl_lxml_off_before_import(tmp_path: Path) -> None:
    environment = os.environ.copy()
    environment["OPENPYXL_LXML"] = "True"
    environment["PYTHONDONTWRITEBYTECODE"] = "1"
    source = (
        "import sys\n"
        f"sys.path.insert(0, {str(PROJECT_ROOT)!r})\n"
        "import tools.verify_v2_parity\n"
        "from openpyxl.xml.functions import LXML\n"
        "print(LXML)\n"
    )

    completed = subprocess.run(
        [sys.executable, "-c", source],
        cwd=tmp_path,
        env=environment,
        check=True,
        capture_output=True,
        text=True,
    )

    assert completed.stdout.strip() == "False"


def _make_dataset(tmp_path: Path) -> tuple[Path, Path, Path]:
    dataset = tmp_path / "dataset"
    task_dir = dataset / "spreadsheets" / "tiny"
    task_dir.mkdir(parents=True)

    init_path = task_dir / "tiny_init.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = 7
    sheet["B1"] = None
    workbook.save(init_path)
    workbook.close()

    # The verifier must never need this workbook. Its different content makes an
    # accidental substitution observable even if the access assertion is removed.
    golden_path = task_dir / "tiny_golden.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "secret golden marker"
    workbook.save(golden_path)
    workbook.close()

    records = [
        {
            "id": "tiny-1",
            "instruction_type": "Cell-Level Manipulation",
            "instruction": "Put the number 42 in Data!B1.",
            "spreadsheet_path": "spreadsheets/tiny",
            "answer_position": "B1",
            "answer_sheet": "Data",
            "data_position": "A1",
        }
    ]
    (dataset / "dataset.json").write_text(json.dumps(records), encoding="utf-8")
    return dataset, init_path, golden_path


def _plan() -> SolvePlan:
    return SolvePlan(
        route="operations",
        summary="Write the requested literal.",
        operations=[SetValue(op="set_value", sheet="Data", cell="B1", value=42)],
        python_code=None,
    )


def _make_reference(
    tmp_path: Path,
    dataset: Path,
    *,
    accepted: bool,
    trace_limit: int = 1_000_000,
) -> Path:
    reference = tmp_path / ("reference-ok" if accepted else "reference-fallback")
    (reference / "outputs").mkdir(parents=True)
    (reference / "traces").mkdir()

    task = load_tasks(dataset)[0]
    context = build_context(task)
    messages = build_messages(task, context)
    prompt = json.dumps(messages, ensure_ascii=False, separators=(",", ":"))
    trace = {
        "schema_version": 1,
        "task_id": task.id,
        "step": 1,
        "model": "retained-reference-model",
        "prompt": prompt,
        "context": {
            "original_chars": context.original_chars,
            "emitted_chars": len(context.text),
            "truncated": context.truncated,
            "sha256": context.sha256,
        },
    }

    output = reference / "outputs" / f"{task.id}.xlsx"
    if accepted:
        plan = _plan()
        trace.update(
            {
                "plan_status": "accepted",
                "tool": "apply_operations",
                "tool_input": plan.model_dump(mode="json"),
            }
        )
        apply_operations(plan, task, task.init_xlsx, output)
        status = "ok"
    else:
        trace.update(
            {
                "plan_status": "parse_rejected",
                "tool_input": {"action": "parse_and_apply_returned_solve_plan"},
            }
        )
        shutil.copyfile(task.init_xlsx, output)
        status = "error:plan"

    retained = truncate_trace_record(trace, trace_limit)
    (reference / "traces" / f"{task.id}.jsonl").write_text(
        json.dumps(retained, ensure_ascii=False, separators=(",", ":")) + "\n",
        encoding="utf-8",
    )
    prediction = {
        "id": task.id,
        "output": f"outputs/{task.id}.xlsx",
        "status": status,
    }
    (reference / "predictions.jsonl").write_text(
        json.dumps(prediction, separators=(",", ":")) + "\n",
        encoding="utf-8",
    )
    return reference


def _rewrite_zip_header_timestamps(path: Path) -> None:
    rewritten = path.with_name("timestamp-rewritten.xlsx")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(rewritten, "w") as target:
        for info in source.infolist():
            replacement = zipfile.ZipInfo(info.filename, date_time=(2001, 2, 3, 4, 5, 6))
            replacement.compress_type = info.compress_type
            replacement.external_attr = info.external_attr
            replacement.internal_attr = info.internal_attr
            replacement.create_system = info.create_system
            target.writestr(replacement, source.read(info.filename))
    rewritten.replace(path)


def _rewrite_core_element(path: Path, element: bytes, value: bytes) -> None:
    rewritten = path.with_name(f"core-{element.decode().replace(':', '-')}-rewritten.xlsx")
    expression = re.compile(
        rb"(?P<open><"
        + re.escape(element)
        + rb"(?:\s[^>]*)?>)[^<]*(?P<close></"
        + re.escape(element)
        + rb">)"
    )
    replacements = 0
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(rewritten, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                payload, count = expression.subn(
                    lambda match: match.group("open") + value + match.group("close"),
                    payload,
                )
                replacements += count
            target.writestr(info, payload)
    assert replacements == 1
    rewritten.replace(path)


def test_fallback_is_byte_identical_and_never_opens_golden(
    tmp_path: Path,
    monkeypatch,
    capsys,
) -> None:
    dataset, _init_path, golden_path = _make_dataset(tmp_path)
    reference = _make_reference(tmp_path, dataset, accepted=False)

    original_load = openpyxl.load_workbook
    opened: list[Path] = []

    def guarded_load(filename, *args, **kwargs):
        path = Path(filename)
        opened.append(path)
        assert path.resolve() != golden_path.resolve()
        return original_load(filename, *args, **kwargs)

    monkeypatch.setattr("exactsource.dataset.openpyxl.load_workbook", guarded_load)
    result = main(
        ["--dataset-dir", str(dataset), "--reference-run", str(reference), "--ids", "tiny-1"]
    )

    assert result == 0
    assert opened
    assert golden_path not in opened
    assert "checked=1 accepted=0 fallback=1 mismatches=0" in capsys.readouterr().out


def test_prompt_mismatch_is_reported(tmp_path: Path, capsys) -> None:
    dataset, _init_path, _golden_path = _make_dataset(tmp_path)
    reference = _make_reference(tmp_path, dataset, accepted=True)
    trace_path = reference / "traces" / "tiny-1.jsonl"
    trace = json.loads(trace_path.read_text(encoding="utf-8"))
    trace["prompt"] += "changed"
    trace_path.write_text(json.dumps(trace) + "\n", encoding="utf-8")

    result = main(["--dataset-dir", str(dataset), "--reference-run", str(reference)])

    captured = capsys.readouterr()
    assert result == 1
    assert "initial full prompt mismatch" in captured.err
    assert "checked=1 accepted=1 fallback=0 mismatches=1" in captured.out


def test_fallback_byte_mismatch_is_reported(tmp_path: Path, capsys) -> None:
    dataset, _init_path, _golden_path = _make_dataset(tmp_path)
    reference = _make_reference(tmp_path, dataset, accepted=False)
    output = reference / "outputs" / "tiny-1.xlsx"
    output.write_bytes(output.read_bytes() + b"not part of the initial workbook")

    result = main(["--dataset-dir", str(dataset), "--reference-run", str(reference)])

    captured = capsys.readouterr()
    assert result == 1
    assert "failure output is not byte-identical to init" in captured.err
    assert "checked=1 accepted=0 fallback=1 mismatches=1" in captured.out


def test_accepted_operations_replay_ignores_only_zip_header_timestamps(
    tmp_path: Path,
    capsys,
) -> None:
    dataset, _init_path, _golden_path = _make_dataset(tmp_path)
    reference = _make_reference(tmp_path, dataset, accepted=True, trace_limit=700)
    output = reference / "outputs" / "tiny-1.xlsx"
    before = output.read_bytes()
    _rewrite_zip_header_timestamps(output)
    assert output.read_bytes() != before

    result = main(["--dataset-dir", str(dataset), "--reference-run", str(reference)])

    assert result == 0
    assert "checked=1 accepted=1 fallback=0 mismatches=0" in capsys.readouterr().out


def test_accepted_replay_ignores_only_core_modified_value(tmp_path: Path, capsys) -> None:
    dataset, _init_path, _golden_path = _make_dataset(tmp_path)
    reference = _make_reference(tmp_path, dataset, accepted=True)
    output = reference / "outputs" / "tiny-1.xlsx"
    before = output.read_bytes()
    _rewrite_core_element(
        output,
        b"dcterms:modified",
        b"1999-12-31T23:59:59Z",
    )
    assert output.read_bytes() != before

    result = main(["--dataset-dir", str(dataset), "--reference-run", str(reference)])

    assert result == 0
    assert "checked=1 accepted=1 fallback=0 mismatches=0" in capsys.readouterr().out


def test_accepted_replay_ignores_only_core_created_value(tmp_path: Path, capsys) -> None:
    dataset, _init_path, _golden_path = _make_dataset(tmp_path)
    reference = _make_reference(tmp_path, dataset, accepted=True)
    output = reference / "outputs" / "tiny-1.xlsx"
    before = output.read_bytes()
    _rewrite_core_element(
        output,
        b"dcterms:created",
        b"1999-12-31T23:59:59Z",
    )
    assert output.read_bytes() != before

    result = main(["--dataset-dir", str(dataset), "--reference-run", str(reference)])

    assert result == 0
    assert "checked=1 accepted=1 fallback=0 mismatches=0" in capsys.readouterr().out


def test_accepted_replay_rejects_duplicate_core_timestamp(tmp_path: Path, capsys) -> None:
    dataset, _init_path, _golden_path = _make_dataset(tmp_path)
    reference = _make_reference(tmp_path, dataset, accepted=True)
    output = reference / "outputs" / "tiny-1.xlsx"
    _rewrite_core_element(
        output,
        b"dcterms:created",
        (
            b"1999-12-31T23:59:59Z</dcterms:created>"
            b'<dcterms:created xsi:type="dcterms:W3CDTF">2000-01-01T00:00:00Z'
        ),
    )

    result = main(["--dataset-dir", str(dataset), "--reference-run", str(reference)])

    captured = capsys.readouterr()
    assert result == 1
    assert "multiple dcterms:created elements" in captured.err


def test_accepted_replay_keeps_core_creator_exact(tmp_path: Path, capsys) -> None:
    dataset, _init_path, _golden_path = _make_dataset(tmp_path)
    reference = _make_reference(tmp_path, dataset, accepted=True)
    output = reference / "outputs" / "tiny-1.xlsx"
    _rewrite_core_element(output, b"dc:creator", b"different creator")

    result = main(["--dataset-dir", str(dataset), "--reference-run", str(reference)])

    captured = capsys.readouterr()
    assert result == 1
    assert "ZIP member content differs at 'docProps/core.xml'" in captured.err
    assert "accepted-plan replay mismatch" in captured.err


def test_accepted_replay_reports_exact_inner_content_mismatch(tmp_path: Path, capsys) -> None:
    dataset, _init_path, _golden_path = _make_dataset(tmp_path)
    reference = _make_reference(tmp_path, dataset, accepted=True)
    output = reference / "outputs" / "tiny-1.xlsx"

    rewritten = output.with_name("content-rewritten.xlsx")
    changed = False
    with zipfile.ZipFile(output, "r") as source, zipfile.ZipFile(rewritten, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if not changed and info.filename.endswith("sheet1.xml"):
                payload += b" "
                changed = True
            target.writestr(info, payload)
    assert changed
    rewritten.replace(output)

    result = main(["--dataset-dir", str(dataset), "--reference-run", str(reference)])

    captured = capsys.readouterr()
    assert result == 1
    assert "ZIP member content differs" in captured.err
    assert "accepted-plan replay mismatch" in captured.err
