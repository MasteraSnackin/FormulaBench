from __future__ import annotations

import datetime as dt
import json
import zipfile
from copy import copy
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

import formulabench.contract as contract_module
from formulabench.contract import (
    EXCEL_MAX_CELL_CHARACTERS,
    EXCEL_MAX_FORMULA_CHARACTERS,
    MAX_EXPANDED_RESPONSE_CELLS,
    MAX_EXPANDED_RESPONSE_CHARACTERS,
    CellKey,
    ContractError,
    ExcelDateTimeValue,
    ExcelDateValue,
    FailureCode,
    ResponseCell,
    ResponseFill,
    ResponsePreserveRange,
    ResponseUnmergeRange,
    SheetResolution,
    SpreadsheetResponse,
    TargetProvenance,
    build_target_contract,
    make_prediction_record,
    make_trace_record,
    normalise_coordinate,
    parse_response,
    validate_response,
    write_response_atomic,
)


def _workbook(path: Path, sheets: tuple[str, ...] = ("Input", "Results")) -> Path:
    workbook = openpyxl.Workbook()
    workbook.active.title = sheets[0]
    for sheet in sheets[1:]:
        workbook.create_sheet(sheet)
    workbook[sheets[0]]["A1"] = "source"
    workbook.save(path)
    workbook.close()
    return path


def _replace_sheet_xml(path: Path, old: bytes, new: bytes) -> None:
    """Replace one exact worksheet fragment without changing workbook semantics."""

    rewritten = path.with_name(f"{path.stem}-rewritten.xlsx")
    with zipfile.ZipFile(path, "r") as source:
        members = [(info, source.read(info.filename)) for info in source.infolist()]
    replacements = 0
    with zipfile.ZipFile(rewritten, "w") as destination:
        for info, payload in members:
            if info.filename == "xl/worksheets/sheet1.xml":
                replacements = payload.count(old)
                payload = payload.replace(old, new)
            destination.writestr(info, payload)
    assert replacements == 1
    rewritten.replace(path)


def _response(*cells: tuple[str, str, object]) -> SpreadsheetResponse:
    return SpreadsheetResponse(
        cells=[ResponseCell(sheet=sheet, cell=cell, value=value) for sheet, cell, value in cells]
    )


def _fill_response(
    *fills: tuple[str, str, object],
    cells: tuple[tuple[str, str, object], ...] = (),
    preserve_ranges: tuple[tuple[str, str], ...] = (),
    unmerge_ranges: tuple[tuple[str, str], ...] = (),
) -> SpreadsheetResponse:
    return SpreadsheetResponse(
        cells=[ResponseCell(sheet=sheet, cell=cell, value=value) for sheet, cell, value in cells],
        fills=[
            ResponseFill(sheet=sheet, range=cell_range, value=value)
            for sheet, cell_range, value in fills
        ],
        preserve_ranges=[
            ResponsePreserveRange(sheet=sheet, range=cell_range)
            for sheet, cell_range in preserve_ranges
        ],
        unmerge_ranges=[
            ResponseUnmergeRange(sheet=sheet, range=cell_range)
            for sheet, cell_range in unmerge_ranges
        ],
    )


def test_response_is_sheet_qualified_strict_and_normalises_absolute_coordinates() -> None:
    response = parse_response(
        '{"cells":[{"sheet":"Model Output","cell":"$b$2","value":"=SUM(A1:A2)"}]}'
    )

    assert response.cells == [ResponseCell(sheet="Model Output", cell="B2", value="=SUM(A1:A2)")]
    assert normalise_coordinate("$XFD$1048576") == "XFD1048576"


def test_response_normalises_finite_fill_preserve_and_unmerge_ranges() -> None:
    response = parse_response(
        '{"cells":[],"fills":[{"sheet":"Output","range":"$b$2:$C$3","value":7}],'
        '"preserve_ranges":[{"sheet":"Output","range":"$F$6:$G$7"}],'
        '"unmerge_ranges":[{"sheet":"Output","range":"$D$4:$E$5"}]}'
    )

    assert response.fills == [ResponseFill(sheet="Output", range="B2:C3", value=7)]
    assert response.preserve_ranges == [ResponsePreserveRange(sheet="Output", range="F6:G7")]
    assert response.unmerge_ranges == [ResponseUnmergeRange(sheet="Output", range="D4:E5")]


def test_response_accepts_strict_typed_date_and_datetime_objects() -> None:
    response = parse_response(
        {
            "cells": [
                {
                    "sheet": "Output",
                    "cell": "A1",
                    "value": {"type": "date", "value": "2024-02-29"},
                }
            ],
            "fills": [
                {
                    "sheet": "Output",
                    "range": "B1:B2",
                    "value": {
                        "type": "datetime",
                        "value": "2024-02-29T23:59:58.123000",
                    },
                }
            ],
        }
    )

    assert response.cells[0].value == ExcelDateValue(type="date", value="2024-02-29")
    assert response.fills[0].value == ExcelDateTimeValue(
        type="datetime",
        value="2024-02-29T23:59:58.123000",
    )


@pytest.mark.parametrize(
    "value",
    [
        {"type": "date", "value": "2023-02-29"},
        {"type": "date", "value": "2024-2-29"},
        {"type": "date", "value": "2024-02-29Z"},
        {"type": "datetime", "value": "2024-02-30T00:00:00"},
        {"type": "datetime", "value": "2024-02-29 00:00:00"},
        {"type": "datetime", "value": "2024-02-29T24:00:00"},
        {"type": "datetime", "value": "2024-02-29T00:00:00Z"},
        {"type": "datetime", "value": "2024-02-29T00:00:00+01:00"},
        {"type": "datetime", "value": "2024-02-29T00:00:00.123456"},
        {"type": "datetime", "value": "2024-02-29T00:00:00.1234567"},
        {"type": "time", "value": "00:00:00"},
        {"type": "date", "value": "2024-02-29", "timezone": "UTC"},
        {"type": "date", "value": 20240229},
        dt.date(2024, 2, 29),
        dt.datetime(2024, 2, 29, tzinfo=dt.UTC),
    ],
)
def test_response_rejects_invalid_aware_or_non_json_typed_dates(value: object) -> None:
    with pytest.raises(ContractError) as caught:
        parse_response({"cells": [{"sheet": "Output", "cell": "A1", "value": value}]})

    assert caught.value.code is FailureCode.INVALID_RESPONSE_SCHEMA


@pytest.mark.parametrize(
    ("payload", "code"),
    [
        (
            '```json\n{"cells":[{"sheet":"Sheet1","cell":"A1","value":1}]}\n```',
            FailureCode.INVALID_RESPONSE_JSON,
        ),
        (
            '{"cells":[{"sheet":"Sheet1","cell":"A1","value":1}],"comment":"done"}',
            FailureCode.INVALID_RESPONSE_SCHEMA,
        ),
        (
            '{"cells":[{"cell":"A1","value":1}]}',
            FailureCode.INVALID_RESPONSE_SCHEMA,
        ),
        (
            '{"cells":[{"sheet":"Sheet1","cell":"A1:B2","value":1}]}',
            FailureCode.INVALID_RESPONSE_SCHEMA,
        ),
        (
            '{"cells":[{"sheet":"Sheet1","cell":"XFE1","value":1}]}',
            FailureCode.INVALID_RESPONSE_SCHEMA,
        ),
        (
            '{"cells":[{"sheet":"Sheet1","cell":"A1","value":NaN}]}',
            FailureCode.INVALID_RESPONSE_JSON,
        ),
        (
            '{"cells":[],"cells":[]}',
            FailureCode.INVALID_RESPONSE_JSON,
        ),
    ],
)
def test_strict_response_parser_rejects_invalid_payloads(payload: str, code: FailureCode) -> None:
    with pytest.raises(ContractError) as caught:
        parse_response(payload)

    assert caught.value.code is code


def test_mapping_response_rejects_non_finite_number() -> None:
    with pytest.raises(ContractError) as caught:
        parse_response({"cells": [{"sheet": "Sheet1", "cell": "A1", "value": float("inf")}]})

    assert caught.value.code is FailureCode.INVALID_RESPONSE_SCHEMA


@pytest.mark.parametrize(
    "cell_range",
    ["A:A", "A1", "B2:A3", "A2:B1", "XFE1:XFE2", "A1048577:A1048578"],
)
def test_fill_rejects_non_finite_reversed_and_out_of_excel_ranges(cell_range: str) -> None:
    with pytest.raises(ContractError) as caught:
        parse_response(
            {
                "cells": [],
                "fills": [{"sheet": "Output", "range": cell_range, "value": 1}],
            }
        )

    assert caught.value.code is FailureCode.INVALID_RESPONSE_SCHEMA


@pytest.mark.parametrize(
    "cell_range",
    ["A:A", "A1", "B2:A3", "A2:B1", "XFE1:XFE2", "A1048577:A1048578"],
)
def test_preserve_rejects_non_finite_reversed_and_out_of_excel_ranges(
    cell_range: str,
) -> None:
    with pytest.raises(ContractError) as caught:
        parse_response(
            {
                "cells": [],
                "preserve_ranges": [{"sheet": "Output", "range": cell_range}],
            }
        )

    assert caught.value.code is FailureCode.INVALID_RESPONSE_SCHEMA


@pytest.mark.parametrize("value", ["x", "=1+"])
def test_response_rejects_strings_beyond_excel_cell_limit(value: str) -> None:
    oversized = value + ("x" * EXCEL_MAX_CELL_CHARACTERS)

    with pytest.raises(ContractError) as caught:
        parse_response({"cells": [{"sheet": "Output", "cell": "A1", "value": oversized}]})

    assert caught.value.code is FailureCode.INVALID_RESPONSE_SCHEMA

    with pytest.raises(ContractError) as fill_error:
        parse_response(
            {
                "cells": [],
                "fills": [{"sheet": "Output", "range": "A1:A2", "value": oversized}],
            }
        )

    assert fill_error.value.code is FailureCode.INVALID_RESPONSE_SCHEMA


def test_response_rejects_formula_beyond_excel_formula_limit() -> None:
    oversized_formula = "=" + ("A" * EXCEL_MAX_FORMULA_CHARACTERS)
    assert len(oversized_formula) == EXCEL_MAX_FORMULA_CHARACTERS + 1
    assert len(oversized_formula) < EXCEL_MAX_CELL_CHARACTERS

    for payload in (
        {"cells": [{"sheet": "Output", "cell": "A1", "value": oversized_formula}]},
        {
            "cells": [],
            "fills": [{"sheet": "Output", "range": "A1:A2", "value": oversized_formula}],
        },
    ):
        with pytest.raises(ContractError) as caught:
            parse_response(payload)
        assert caught.value.code is FailureCode.INVALID_RESPONSE_SCHEMA


@pytest.mark.parametrize(
    "formula",
    [
        '=WEBSERVICE("https://example.invalid")',
        '=_xlfn.WEBSERVICE("https://example.invalid")',
        '=@HYPERLINK("https://example.invalid","open")',
        '=RTD("example.prog",,"topic")',
        '=IMAGE("https://example.invalid/image.png")',
        '=STOCKHISTORY("MSFT")',
        '=CUBEVALUE("external-connection","[Measures].[Value]")',
        '=PY("print(1)")',
        '=CALL("library","procedure","J")',
        "=cmd|' /C calc'!A0",
        "='https://example.invalid/[Book.xlsx]Sheet1'!A1",
        "='C:\\Reports\\[Book.xlsx]Sheet1'!A1",
        "='[Book.xlsx]Sheet1'!A1",
        '=INDIRECT("[Book.xlsx]Sheet1!A1")',
        "=WEBSERVICE(A1)#",
        "=_xlfn.WEBSERVICE(A1)#",
        "=HYPERLINK(A1)#",
        '=INDIRECT("[Book.xlsx]Sheet1!A1")#',
        '=INDIRECT("https://example.invalid/[Book.xlsx]Sheet1!A1")#',
        "='[Book]Sheet1'!A1#",
    ],
)
def test_atomic_write_rejects_unsafe_formulas_and_restores_input(
    tmp_path: Path,
    formula: str,
) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Output",))
    source_bytes = source.read_bytes()
    destination = tmp_path / "output.xlsx"

    result = write_response_atomic(
        {"answer_position": "'Output'!A1", "answer_sheet": "Output"},
        _response(("Output", "A1", formula)),
        source,
        destination,
    )

    assert not result.success
    assert result.failure_codes == (FailureCode.UNSAFE_FORMULA,)
    assert destination.read_bytes() == source_bytes


def test_atomic_write_rejects_unsafe_fill_formula_and_restores_input(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Output",))
    source_bytes = source.read_bytes()
    destination = tmp_path / "output.xlsx"

    result = write_response_atomic(
        {"answer_position": "'Output'!A1:B1", "answer_sheet": "Output"},
        _fill_response(("Output", "A1:B1", '=WEBSERVICE("https://example.invalid")')),
        source,
        destination,
    )

    assert not result.success
    assert result.failure_codes == (FailureCode.UNSAFE_FORMULA,)
    assert destination.read_bytes() == source_bytes


@pytest.mark.parametrize(
    "formula",
    [
        "=SUM(Table1[Amount])",
        "=Sheet1!Table1[Amount]",
        '=INDIRECT("A"&1)',
        '=IF(A1,"https://example.invalid","")',
        '="a|b"',
        "=A1#",
        "=SUM(Table1[Amount])+Sheet2!A1",
        "=INDEX(Table1[Amount],MATCH(A1,Sheet2!A:A,0))",
        "=SUM('WEBSERVICE(A1)'!A:A)",
    ],
)
def test_atomic_write_preserves_benign_formula_syntax(tmp_path: Path, formula: str) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Output",))
    destination = tmp_path / "output.xlsx"

    result = write_response_atomic(
        {"answer_position": "'Output'!A1", "answer_sheet": "Output"},
        _response(("Output", "A1", formula)),
        source,
        destination,
    )

    assert result.success
    written = openpyxl.load_workbook(destination, data_only=False)
    assert written["Output"]["A1"].value == formula
    written.close()


def test_runner_failure_codes_are_distinct_from_workbook_write_failure() -> None:
    assert FailureCode.MODEL_CALL_FAILED.value == "model_call_failed"
    assert FailureCode.CONTEXT_BUILD_FAILED.value == "context_build_failed"
    assert FailureCode.INPUT_LOAD_FAILED.value == "input_load_failed"
    assert (
        len(
            {
                FailureCode.MODEL_CALL_FAILED,
                FailureCode.CONTEXT_BUILD_FAILED,
                FailureCode.INPUT_LOAD_FAILED,
                FailureCode.WORKBOOK_WRITE_FAILED,
            }
        )
        == 4
    )


def test_target_contract_preserves_provenance_and_exact_sheet_policy(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Active", "Existing"))
    workbook = openpyxl.load_workbook(source)
    task = {
        "answer_position": "'Existing'!$B$2,'New Output'!C3,D4",
        "answer_sheet": "Missing Output",
    }

    contract = build_target_contract(task, workbook)
    workbook.close()

    assert contract.active_sheet == "Active"
    assert contract.ranges[0].sheet == "Existing"
    assert contract.ranges[0].requested_sheet == "Existing"
    assert contract.ranges[0].cell_range == "B2"
    assert contract.ranges[0].provenance is TargetProvenance.EXPLICIT
    assert contract.ranges[0].resolution is SheetResolution.EXISTING

    assert contract.ranges[1].sheet == "New Output"
    assert contract.ranges[1].requested_sheet == "New Output"
    assert contract.ranges[1].provenance is TargetProvenance.EXPLICIT
    assert contract.ranges[1].resolution is SheetResolution.CREATED

    assert contract.ranges[2].sheet == "Active"
    assert contract.ranges[2].requested_sheet == "Missing Output"
    assert contract.ranges[2].provenance is TargetProvenance.ANSWER_SHEET
    assert contract.ranges[2].resolution is SheetResolution.ACTIVE_FALLBACK


def test_unqualified_existing_answer_sheet_is_used_exactly(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Active", "answer sheet"))
    workbook = openpyxl.load_workbook(source)

    contract = build_target_contract(
        {"answer_position": "$b$2:$c$3", "answer_sheet": "answer sheet"},
        workbook,
    )
    workbook.close()

    target = contract.ranges[0]
    assert target.sheet == "answer sheet"
    assert target.cell_range == "B2:C3"
    assert target.provenance is TargetProvenance.ANSWER_SHEET
    assert target.resolution is SheetResolution.EXISTING


def test_explicit_created_sheet_cannot_be_silently_renamed_for_case_collision(
    tmp_path: Path,
) -> None:
    source = _workbook(tmp_path / "input.xlsx")
    workbook = openpyxl.load_workbook(source)

    with pytest.raises(ContractError) as caught:
        build_target_contract(
            {
                "answer_position": "'New Output'!A1,'new output'!B1",
                "answer_sheet": "Input",
            },
            workbook,
        )
    workbook.close()

    assert caught.value.code is FailureCode.INVALID_TARGET_SHEET


def test_validation_reports_duplicate_missing_and_extra_cells_together(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("One", "Two"))
    workbook = openpyxl.load_workbook(source)
    contract = build_target_contract(
        {"answer_position": "'One'!B2:B3,'Two'!B2", "answer_sheet": "One"},
        workbook,
    )
    workbook.close()
    response = _response(
        ("One", "$B$2", 10),
        ("One", "B2", 11),
        ("Two", "B2", 12),
        ("Two", "C2", 13),
    )

    result = validate_response(contract, response)

    assert result.failure_codes == (
        FailureCode.DUPLICATE_CELLS,
        FailureCode.MISSING_CELLS,
        FailureCode.EXTRA_CELLS,
    )
    assert result.duplicates == (CellKey("One", "B2"),)
    assert result.missing == (CellKey("One", "B3"),)
    assert result.extra == (CellKey("Two", "C2"),)


def test_same_coordinate_on_different_sheets_is_distinct(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("One", "Two"))
    workbook = openpyxl.load_workbook(source)
    contract = build_target_contract(
        {"answer_position": "'One'!B2,'Two'!B2", "answer_sheet": "One"},
        workbook,
    )
    workbook.close()

    result = validate_response(
        contract,
        _response(("One", "B2", 1), ("Two", "B2", 2)),
    )

    assert result.ok
    assert set(result.expected) == {CellKey("One", "B2"), CellKey("Two", "B2")}


def test_dynamic_whole_column_range_uses_returned_maximum_row(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    workbook = openpyxl.load_workbook(source)
    contract = build_target_contract(
        {"answer_position": "'Output'!A:C", "answer_sheet": "Output"},
        workbook,
    )
    response = _response(
        *(("Output", f"{column}{row}", row) for row in range(1, 4) for column in "ABC")
    )

    result = validate_response(contract, response, workbook=workbook)
    workbook.close()

    assert result.ok
    assert len(result.expected) == 9
    assert result.dynamic_extents[0].max_row == 3
    assert result.expected[-1] == CellKey("Output", "C3")


def test_dynamic_whole_column_range_accepts_one_finite_fill_extent(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    workbook = openpyxl.load_workbook(source)
    contract = build_target_contract(
        {"answer_position": "'Output'!A:C", "answer_sheet": "Output"},
        workbook,
    )
    response = _fill_response(("Output", "A1:C3", 7))

    result = validate_response(contract, response, workbook=workbook)
    workbook.close()

    assert result.ok
    assert len(result.expected) == 9
    assert result.dynamic_extents[0].max_row == 3
    assert len(result.materialised_cells) == 9


def test_dynamic_range_requires_dense_coverage_through_initial_sheet_extent(
    tmp_path: Path,
) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    workbook = openpyxl.load_workbook(source)
    workbook["Output"]["E10"] = "existing"
    contract = build_target_contract(
        {"answer_position": "'Output'!A:C", "answer_sheet": "Output"},
        workbook,
    )

    result = validate_response(
        contract,
        _fill_response(("Output", "A1:C2", 7)),
        workbook=workbook,
    )
    workbook.close()

    assert result.failure_codes == (FailureCode.MISSING_CELLS,)
    assert result.dynamic_extents[0].max_row == 10
    assert len(result.expected) == 30
    assert set(result.missing) == {
        CellKey("Output", f"{column}{row}") for row in range(3, 11) for column in "ABC"
    }


def test_dynamic_range_preserves_model_extension_beyond_initial_extent(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    workbook = openpyxl.load_workbook(source)
    workbook["Output"]["E10"] = "existing"
    contract = build_target_contract(
        {"answer_position": "'Output'!A:C", "answer_sheet": "Output"},
        workbook,
    )

    result = validate_response(
        contract,
        _fill_response(("Output", "A1:C12", 7)),
        workbook=workbook,
    )
    workbook.close()

    assert result.ok
    assert result.dynamic_extents[0].max_row == 12
    assert len(result.expected) == 36


def test_dynamic_initial_sheet_extent_is_area_capped_before_expansion(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    workbook = openpyxl.load_workbook(source)
    workbook["Output"]["E100000"] = "existing"
    contract = build_target_contract(
        {"answer_position": "'Output'!A:C", "answer_sheet": "Output"},
        workbook,
    )

    with pytest.raises(ContractError) as caught:
        validate_response(
            contract,
            _fill_response(("Output", "A1:C1", 7)),
            workbook=workbook,
        )
    workbook.close()

    assert caught.value.code is FailureCode.EXPANSION_LIMIT_EXCEEDED


def test_dynamic_range_uses_exact_active_fallback_sheet_extent(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Active", "Other"))
    workbook = openpyxl.load_workbook(source)
    workbook["Active"]["E5"] = "existing"
    workbook["Other"]["E20"] = "must not affect active target"
    contract = build_target_contract(
        {"answer_position": "A:C", "answer_sheet": "Missing"},
        workbook,
    )

    result = validate_response(
        contract,
        _fill_response(("Active", "A1:C5", 7)),
        workbook=workbook,
    )
    workbook.close()

    assert contract.ranges[0].resolution is SheetResolution.ACTIVE_FALLBACK
    assert result.ok
    assert result.dynamic_extents[0].max_row == 5


def test_cell_fill_and_fill_fill_overlaps_are_duplicate_failures(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    workbook = openpyxl.load_workbook(source)
    contract = build_target_contract(
        {"answer_position": "'Output'!A1:C2", "answer_sheet": "Output"},
        workbook,
    )
    response = _fill_response(
        ("Output", "A1:B2", 1),
        ("Output", "B1:C2", 2),
        cells=(("Output", "A1", 3),),
    )

    result = validate_response(contract, response, workbook=workbook)
    workbook.close()

    assert result.failure_codes == (FailureCode.DUPLICATE_CELLS,)
    assert set(result.duplicates) == {
        CellKey("Output", "A1"),
        CellKey("Output", "B1"),
        CellKey("Output", "B2"),
    }


def test_preserves_cover_dynamic_targets_within_existing_extent(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    workbook = openpyxl.load_workbook(source)
    workbook["Output"]["B3"] = "existing extent"
    contract = build_target_contract(
        {"answer_position": "'Output'!A:B", "answer_sheet": "Output"},
        workbook,
    )
    response = _fill_response(
        preserve_ranges=(("Output", "A1:B3"),),
    )

    result = validate_response(contract, response, workbook=workbook)
    workbook.close()

    assert result.ok
    assert result.dynamic_extents[0].max_row == 3
    assert result.materialised_cells == ()
    assert len(result.preserved_cells) == 6
    assert result.preserved_cells[-1] == CellKey("Output", "B3")


def test_preserve_cannot_extend_blank_dynamic_target_extent(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    workbook = openpyxl.load_workbook(source)
    contract = build_target_contract(
        {"answer_position": "'Output'!A:B", "answer_sheet": "Output"},
        workbook,
    )
    response = _fill_response(preserve_ranges=(("Output", "A1:B3"),))

    with pytest.raises(ContractError) as caught:
        validate_response(contract, response, workbook=workbook)
    workbook.close()

    assert caught.value.code is FailureCode.INVALID_RESPONSE_SCHEMA
    assert "cannot extend a dynamic target" in str(caught.value)


def test_preserve_overlaps_cells_fills_and_other_preserves_as_duplicates(
    tmp_path: Path,
) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    workbook = openpyxl.load_workbook(source)
    contract = build_target_contract(
        {"answer_position": "'Output'!A1:D2", "answer_sheet": "Output"},
        workbook,
    )
    response = _fill_response(
        ("Output", "C1:D2", 9),
        cells=(("Output", "A1", 3),),
        preserve_ranges=(("Output", "A1:B2"), ("Output", "B1:C2")),
    )

    result = validate_response(contract, response, workbook=workbook)
    workbook.close()

    assert result.failure_codes == (FailureCode.DUPLICATE_CELLS,)
    assert set(result.duplicates) == {
        CellKey("Output", "A1"),
        CellKey("Output", "B1"),
        CellKey("Output", "B2"),
        CellKey("Output", "C1"),
        CellKey("Output", "C2"),
    }


def test_preserve_missing_and_outside_target_cells_are_reported(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output", "Other"))
    workbook = openpyxl.load_workbook(source)
    contract = build_target_contract(
        {"answer_position": "'Output'!A1:B2", "answer_sheet": "Output"},
        workbook,
    )
    response = _fill_response(
        cells=(("Output", "A2", 3),),
        preserve_ranges=(("Output", "A1:B1"), ("Other", "D4:D4")),
    )

    result = validate_response(contract, response, workbook=workbook)
    workbook.close()

    assert result.failure_codes == (
        FailureCode.MISSING_CELLS,
        FailureCode.EXTRA_CELLS,
    )
    assert result.missing == (CellKey("Output", "B2"),)
    assert result.extra == (CellKey("Other", "D4"),)


def test_preserve_requires_workbook_and_rejects_new_target_sheet(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input",))
    workbook = openpyxl.load_workbook(source)
    contract = build_target_contract(
        {"answer_position": "'New Output'!A1:B1", "answer_sheet": "Input"},
        workbook,
    )
    response = _fill_response(
        preserve_ranges=(("New Output", "A1:B1"),),
    )

    with pytest.raises(ContractError) as missing_workbook:
        validate_response(contract, response)
    assert missing_workbook.value.code is FailureCode.INVALID_INPUT_WORKBOOK

    with pytest.raises(ContractError) as new_sheet:
        validate_response(contract, response, workbook=workbook)
    workbook.close()

    assert new_sheet.value.code is FailureCode.INVALID_RESPONSE_SCHEMA

    destination = tmp_path / "output.xlsx"
    destination.write_bytes(b"stale output")
    written = write_response_atomic(
        {"answer_position": "'New Output'!A1:B1", "answer_sheet": "Input"},
        response,
        source,
        destination,
    )
    assert written.failure_codes == (FailureCode.INVALID_RESPONSE_SCHEMA,)
    assert destination.read_bytes() == source.read_bytes()


def test_preserve_expansion_shares_response_cell_limit(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    workbook = openpyxl.load_workbook(source)
    contract = build_target_contract(
        {"answer_position": "'Output'!A1", "answer_sheet": "Output"},
        workbook,
    )
    response = _fill_response(
        preserve_ranges=(("Output", "A1:XFD1048576"),),
    )

    with pytest.raises(ContractError) as caught:
        validate_response(contract, response, workbook=workbook)
    workbook.close()

    assert caught.value.code is FailureCode.EXPANSION_LIMIT_EXCEEDED


def test_fill_rejects_formula_that_cannot_be_translated(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    workbook = openpyxl.load_workbook(source)
    contract = build_target_contract(
        {"answer_position": "'Output'!B1:C1", "answer_sheet": "Output"},
        workbook,
    )

    with pytest.raises(ContractError) as caught:
        validate_response(
            contract,
            _fill_response(("Output", "B1:C1", '="unterminated')),
            workbook=workbook,
        )
    workbook.close()

    assert caught.value.code is FailureCode.INVALID_RESPONSE_SCHEMA


def test_formula_fill_reuses_one_translator_instance(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    workbook = openpyxl.load_workbook(source)
    contract = build_target_contract(
        {"answer_position": "'Output'!A1:A3", "answer_sheet": "Output"},
        workbook,
    )
    real_translator = contract_module.Translator
    constructions = 0

    class CountingTranslator:
        def __init__(self, formula: str, *, origin: str) -> None:
            nonlocal constructions
            constructions += 1
            self._delegate = real_translator(formula, origin=origin)

        def translate_formula(self, destination: str) -> str:
            return self._delegate.translate_formula(destination)

    monkeypatch.setattr(contract_module, "Translator", CountingTranslator)
    result = validate_response(
        contract,
        _fill_response(("Output", "A1:A3", "=ROW()")),
        workbook=workbook,
    )
    workbook.close()

    assert result.ok
    assert constructions == 1


def test_translated_formula_cannot_grow_beyond_excel_formula_limit(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    workbook = openpyxl.load_workbook(source)
    contract = build_target_contract(
        {"answer_position": "'Output'!A1:A10", "answer_sheet": "Output"},
        workbook,
    )
    formula = "=" + ("X" * (EXCEL_MAX_FORMULA_CHARACTERS - 4)) + "+A1"
    assert len(formula) == EXCEL_MAX_FORMULA_CHARACTERS

    with pytest.raises(ContractError) as caught:
        validate_response(
            contract,
            _fill_response(("Output", "A1:A10", formula)),
            workbook=workbook,
        )
    workbook.close()

    assert caught.value.code is FailureCode.INVALID_RESPONSE_SCHEMA


def test_repeated_strings_are_character_capped_before_fill_loop(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    workbook = openpyxl.load_workbook(source)
    contract = build_target_contract(
        {"answer_position": "'Output'!A1", "answer_sheet": "Output"},
        workbook,
    )
    response = _fill_response(("Output", "A1:A27000", "x" * 5_000))

    with pytest.raises(ContractError) as caught:
        validate_response(contract, response, workbook=workbook)
    workbook.close()

    assert MAX_EXPANDED_RESPONSE_CHARACTERS == 128 * 1024 * 1024
    assert caught.value.code is FailureCode.EXPANSION_LIMIT_EXCEEDED


def test_response_and_target_expansion_are_capped_before_large_materialisation(
    tmp_path: Path,
) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    workbook = openpyxl.load_workbook(source)
    small_contract = build_target_contract(
        {"answer_position": "'Output'!A1", "answer_sheet": "Output"},
        workbook,
    )
    oversized_fill = _fill_response(("Output", "A1:XFD1048576", None))
    with pytest.raises(ContractError) as response_error:
        validate_response(small_contract, oversized_fill, workbook=workbook)
    assert response_error.value.code is FailureCode.EXPANSION_LIMIT_EXCEEDED

    oversized_target = build_target_contract(
        {"answer_position": "'Output'!A1:A250001", "answer_sheet": "Output"},
        workbook,
    )
    with pytest.raises(ContractError) as target_error:
        validate_response(oversized_target, _response(("Output", "A1", 1)), workbook=workbook)
    workbook.close()

    assert MAX_EXPANDED_RESPONSE_CELLS == 250_000
    assert target_error.value.code is FailureCode.EXPANSION_LIMIT_EXCEEDED


def test_dynamic_whole_column_range_checks_coverage_and_unrelated_extras(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    workbook = openpyxl.load_workbook(source)
    contract = build_target_contract(
        {"answer_position": "'Output'!A:C", "answer_sheet": "Output"},
        workbook,
    )
    result = validate_response(
        contract,
        _response(("Output", "A1", 1), ("Output", "C2", 2), ("Output", "D1", 3)),
        workbook=workbook,
    )
    workbook.close()

    assert result.failure_codes == (
        FailureCode.MISSING_CELLS,
        FailureCode.EXTRA_CELLS,
    )
    assert set(result.missing) == {
        CellKey("Output", "B1"),
        CellKey("Output", "C1"),
        CellKey("Output", "A2"),
        CellKey("Output", "B2"),
    }
    assert result.extra == (CellKey("Output", "D1"),)


def test_dynamic_range_without_target_cells_still_uses_worksheet_extent(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    workbook = openpyxl.load_workbook(source)
    contract = build_target_contract(
        {"answer_position": "'Output'!A:C", "answer_sheet": "Output"},
        workbook,
    )
    result = validate_response(
        contract,
        _response(("Output", "D1", 1)),
        workbook=workbook,
    )
    workbook.close()

    assert result.failure_codes == (
        FailureCode.MISSING_CELLS,
        FailureCode.EXTRA_CELLS,
    )
    assert set(result.missing) == {
        CellKey("Output", "A1"),
        CellKey("Output", "B1"),
        CellKey("Output", "C1"),
    }


def test_atomic_write_creates_only_explicit_missing_sheet_and_preserves_formula(
    tmp_path: Path,
) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Existing"))
    source_bytes = source.read_bytes()
    destination = tmp_path / "nested" / "output.xlsx"
    task = {
        "answer_position": "'Existing'!B2,'New Output'!$C$3",
        "answer_sheet": "Existing",
    }
    response = _response(
        ("Existing", "B2", "=SUM(1,2)"),
        ("New Output", "C3", 7),
    )

    result = write_response_atomic(task, response, source, destination)

    assert result.success
    assert result.status == "ok"
    assert source.read_bytes() == source_bytes
    workbook = openpyxl.load_workbook(destination, data_only=False)
    assert workbook.sheetnames == ["Input", "Existing", "New Output"]
    assert workbook["Existing"]["B2"].value == "=SUM(1,2)"
    assert workbook["Existing"]["B2"].data_type == "f"
    assert workbook["New Output"]["C3"].value == 7
    workbook.close()


def test_atomic_fill_translates_relative_references_and_repeats_constants(
    tmp_path: Path,
) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    destination = tmp_path / "output.xlsx"
    task = {
        "answer_position": "'Output'!B2:C3,'Output'!D2:D3",
        "answer_sheet": "Output",
    }
    response = _fill_response(
        ("Output", "B2:C3", "=$A2+B$1+$A$1"),
        ("Output", "D2:D3", 9),
    )

    result = write_response_atomic(task, response, source, destination)

    assert result.success
    workbook = openpyxl.load_workbook(destination, data_only=False)
    assert workbook["Output"]["B2"].value == "=$A2+B$1+$A$1"
    assert workbook["Output"]["C2"].value == "=$A2+C$1+$A$1"
    assert workbook["Output"]["B3"].value == "=$A3+B$1+$A$1"
    assert workbook["Output"]["C3"].value == "=$A3+C$1+$A$1"
    assert workbook["Output"]["D2"].value == 9
    assert workbook["Output"]["D3"].value == 9
    assert all(workbook["Output"][cell].data_type == "f" for cell in ("B2", "C2", "B3", "C3"))
    workbook.close()


def test_typed_dates_and_typed_date_fills_save_as_native_excel_values(
    tmp_path: Path,
) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    source_bytes = source.read_bytes()
    destination = tmp_path / "output.xlsx"
    response = {
        "cells": [
            {
                "sheet": "Output",
                "cell": "A1",
                "value": {"type": "date", "value": "2024-02-29"},
            },
            {
                "sheet": "Output",
                "cell": "B1",
                "value": {
                    "type": "datetime",
                    "value": "2024-02-29T23:59:58.123000",
                },
            },
            {"sheet": "Output", "cell": "C1", "value": "2024-02-29"},
        ],
        "fills": [
            {
                "sheet": "Output",
                "range": "A2:B2",
                "value": {"type": "date", "value": "2025-01-02"},
            },
            {
                "sheet": "Output",
                "range": "C2:D2",
                "value": {
                    "type": "datetime",
                    "value": "2025-01-02T03:04:05.006000",
                },
            },
        ],
    }

    result = write_response_atomic(
        {"answer_position": "'Output'!A1:C1,'Output'!A2:D2", "answer_sheet": "Output"},
        response,
        source,
        destination,
    )

    assert result.success
    assert source.read_bytes() == source_bytes
    workbook = openpyxl.load_workbook(destination, data_only=False)
    sheet = workbook["Output"]
    assert sheet["A1"].value == dt.datetime(2024, 2, 29)
    assert type(sheet["A1"].value) is dt.datetime
    assert sheet["B1"].value == dt.datetime(2024, 2, 29, 23, 59, 58, 123000)
    assert sheet["C1"].value == "2024-02-29"
    assert sheet["C1"].data_type == "s"
    assert [sheet[cell].value for cell in ("A2", "B2")] == [
        dt.datetime(2025, 1, 2),
        dt.datetime(2025, 1, 2),
    ]
    assert [sheet[cell].value for cell in ("C2", "D2")] == [
        dt.datetime(2025, 1, 2, 3, 4, 5, 6000),
        dt.datetime(2025, 1, 2, 3, 4, 5, 6000),
    ]
    assert all(sheet[cell].data_type == "d" for cell in ("A1", "B1", "A2", "B2", "C2", "D2"))
    workbook.close()


def test_atomic_preserve_ranges_leave_existing_cells_unchanged(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    workbook = openpyxl.load_workbook(source)
    sheet = workbook["Output"]
    sheet["A1"] = "=SUM(1,2)"
    bold_font = copy(sheet["A1"].font)
    bold_font.bold = True
    sheet["A1"].font = bold_font
    sheet["B1"] = dt.datetime(2024, 2, 29, 12, 0)
    sheet["B1"].number_format = "yyyy-mm-dd hh:mm"
    workbook.save(source)
    workbook.close()
    destination = tmp_path / "output.xlsx"
    response = _fill_response(
        cells=(("Output", "C1", 9),),
        preserve_ranges=(("Output", "A1:B1"),),
    )

    result = write_response_atomic(
        {"answer_position": "'Output'!A1:C1", "answer_sheet": "Output"},
        response,
        source,
        destination,
    )

    assert result.success
    assert result.validation is not None
    assert result.validation.preserved_cells == (
        CellKey("Output", "A1"),
        CellKey("Output", "B1"),
    )
    before_workbook = openpyxl.load_workbook(source, data_only=False)
    before = before_workbook["Output"]
    after_workbook = openpyxl.load_workbook(destination, data_only=False)
    after = after_workbook["Output"]
    for coordinate in ("A1", "B1"):
        assert after[coordinate].value == before[coordinate].value
        assert after[coordinate].data_type == before[coordinate].data_type
        assert after[coordinate].number_format == before[coordinate].number_format
        assert after[coordinate]._style == before[coordinate]._style
    assert after["C1"].value == 9
    before_workbook.close()
    after_workbook.close()


def test_atomic_preserve_accepts_empty_string_round_trip_normalisation(
    tmp_path: Path,
) -> None:
    source = tmp_path / "input.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Output"
    workbook["Output"]["A1"] = "placeholder"
    workbook.save(source)
    workbook.close()
    _replace_sheet_xml(source, b"<t>placeholder</t>", b"<t></t>")

    loaded = openpyxl.load_workbook(source, data_only=False)
    before = loaded["Output"]["A1"]
    assert (before.value, before.data_type, tuple(before._style)) == (
        "",
        "s",
        (0, 0, 0, 0, 0, 0, 0, 0, 0),
    )
    loaded.close()

    destination = tmp_path / "output.xlsx"
    result = write_response_atomic(
        {"answer_position": "'Output'!A1:B1", "answer_sheet": "Output"},
        _fill_response(
            cells=(("Output", "B1", "written"),),
            preserve_ranges=(("Output", "A1:A1"),),
        ),
        source,
        destination,
    )

    assert result.success
    after = openpyxl.load_workbook(destination, data_only=False)
    assert after["Output"]["A1"].value is None
    assert after["Output"]["A1"].data_type == "inlineStr"
    assert after["Output"]["B1"].value == "written"
    after.close()


def test_atomic_preserve_accepts_writer_float_rounding(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Output"
    workbook["Output"]["A1"] = 1
    workbook.save(source)
    workbook.close()
    _replace_sheet_xml(source, b"<v>1</v>", b"<v>27338.850000000002</v>")

    loaded = openpyxl.load_workbook(source, data_only=False)
    assert loaded["Output"]["A1"].value == 27338.850000000002
    loaded.close()

    destination = tmp_path / "output.xlsx"
    result = write_response_atomic(
        {"answer_position": "'Output'!A1:B1", "answer_sheet": "Output"},
        _fill_response(
            cells=(("Output", "B1", "written"),),
            preserve_ranges=(("Output", "A1:A1"),),
        ),
        source,
        destination,
    )

    assert result.success
    after = openpyxl.load_workbook(destination, data_only=False)
    assert after["Output"]["A1"].value == 27338.85
    assert after["Output"]["B1"].value == "written"
    after.close()


def test_atomic_preserve_accepts_integral_float_round_trip_as_integer(
    tmp_path: Path,
) -> None:
    source = tmp_path / "input.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Output"
    workbook["Output"]["A1"] = 1
    workbook.save(source)
    workbook.close()
    _replace_sheet_xml(source, b"<v>1</v>", b"<v>1.0</v>")

    loaded = openpyxl.load_workbook(source, data_only=False)
    assert type(loaded["Output"]["A1"].value) is float
    loaded.close()

    destination = tmp_path / "output.xlsx"
    result = write_response_atomic(
        {"answer_position": "'Output'!A1:B1", "answer_sheet": "Output"},
        _fill_response(
            cells=(("Output", "B1", "written"),),
            preserve_ranges=(("Output", "A1:A1"),),
        ),
        source,
        destination,
    )

    assert result.success
    after = openpyxl.load_workbook(destination, data_only=False)
    assert type(after["Output"]["A1"].value) is int
    assert after["Output"]["A1"].value == 1
    assert after["Output"]["B1"].value == "written"
    after.close()


def test_atomic_preserve_accepts_zero_style_blank_normalisation(tmp_path: Path) -> None:
    source = tmp_path / "input.xlsx"
    workbook = openpyxl.Workbook()
    workbook.active.title = "Output"
    workbook["Output"]["A1"] = 1
    workbook.save(source)
    workbook.close()
    _replace_sheet_xml(source, b"<v>1</v>", b"<v></v>")

    loaded = openpyxl.load_workbook(source, data_only=False)
    before = loaded["Output"]["A1"]
    assert before.value is None
    assert before.data_type == "n"
    assert tuple(before._style) == (0, 0, 0, 0, 0, 0, 0, 0, 0)
    loaded.close()

    destination = tmp_path / "output.xlsx"
    result = write_response_atomic(
        {"answer_position": "'Output'!A1:B1", "answer_sheet": "Output"},
        _fill_response(
            cells=(("Output", "B1", "written"),),
            preserve_ranges=(("Output", "A1:A1"),),
        ),
        source,
        destination,
    )

    assert result.success
    after = openpyxl.load_workbook(destination, data_only=False)
    assert after["Output"]["A1"].value is None
    assert after["Output"]["A1"]._style is None
    assert after["Output"]["B1"].value == "written"
    after.close()


def test_atomic_assignment_accepts_writer_float_normalisation(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Output",))
    destination = tmp_path / "output.xlsx"

    result = write_response_atomic(
        {"answer_position": "'Output'!A1", "answer_sheet": "Output"},
        _response(("Output", "A1", 27338.850000000002)),
        source,
        destination,
    )

    assert result.success
    written = openpyxl.load_workbook(destination, data_only=False)
    assert written["Output"]["A1"].value == 27338.85
    assert written["Output"]["A1"].data_type == "n"
    written.close()


@pytest.mark.parametrize(
    ("serial", "number_format", "expected_value"),
    [
        (0.5, "h:mm:ss", dt.time(12)),
        (61, "yyyy-mm-dd", dt.datetime(1900, 3, 1)),
        (1.5, "[h]:mm", dt.timedelta(days=1, hours=12)),
    ],
)
def test_atomic_assignment_accepts_formatted_temporal_serials(
    tmp_path: Path,
    serial: int | float,
    number_format: str,
    expected_value: object,
) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Output",))
    workbook = openpyxl.load_workbook(source)
    workbook["Output"]["A1"].number_format = number_format
    workbook.save(source)
    workbook.close()
    destination = tmp_path / "output.xlsx"

    result = write_response_atomic(
        {"answer_position": "'Output'!A1", "answer_sheet": "Output"},
        _response(("Output", "A1", serial)),
        source,
        destination,
    )

    assert result.success
    written = openpyxl.load_workbook(destination, data_only=False)
    assert written["Output"]["A1"].value == expected_value
    assert written["Output"]["A1"].data_type == "d"
    written.close()


def test_atomic_assignment_uses_1904_epoch_for_formatted_serial(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Output",))
    workbook = openpyxl.load_workbook(source)
    workbook.epoch = CALENDAR_MAC_1904
    workbook["Output"]["A1"].number_format = "yyyy-mm-dd"
    workbook.save(source)
    workbook.close()
    destination = tmp_path / "output.xlsx"

    result = write_response_atomic(
        {"answer_position": "'Output'!A1", "answer_sheet": "Output"},
        _response(("Output", "A1", 1)),
        source,
        destination,
    )

    assert result.success
    written = openpyxl.load_workbook(destination, data_only=False)
    assert written.epoch == CALENDAR_MAC_1904
    assert written["Output"]["A1"].value == dt.datetime(1904, 1, 2)
    assert written["Output"]["A1"].data_type == "d"
    written.close()


def test_assigned_value_round_trip_rejects_non_equivalent_number() -> None:
    assert not contract_module._assigned_value_round_trips(
        1.25,
        "n",
        1.2500000000000002,
        "n",
        CALENDAR_WINDOWS_1900,
    )


def test_atomic_assignment_keeps_empty_string_comparison_strict(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Output",))
    source_bytes = source.read_bytes()
    destination = tmp_path / "output.xlsx"

    result = write_response_atomic(
        {"answer_position": "'Output'!A1", "answer_sheet": "Output"},
        _response(("Output", "A1", "")),
        source,
        destination,
    )

    assert not result.success
    assert result.failure_codes == (FailureCode.WORKBOOK_WRITE_FAILED,)
    assert destination.read_bytes() == source_bytes


@pytest.mark.parametrize(
    ("formula", "target_range", "preserve_range"),
    [
        (ArrayFormula(ref="A1:A2", text="=ROW(A1:A2)"), "A1:A2", "A1:A2"),
        (
            DataTableFormula(
                ref="A1:B2",
                ca=True,
                dt2D=True,
                dtr=True,
                r1="C1",
                r2="C2",
                del1=True,
            ),
            "A1:B2",
            "A1:B2",
        ),
    ],
)
def test_atomic_preserve_round_trips_special_formula_anchors(
    tmp_path: Path,
    formula: ArrayFormula | DataTableFormula,
    target_range: str,
    preserve_range: str,
) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Output",))
    workbook = openpyxl.load_workbook(source)
    workbook["Output"]["A1"] = formula
    workbook.save(source)
    workbook.close()
    destination = tmp_path / "output.xlsx"

    result = write_response_atomic(
        {
            "answer_position": f"'Output'!{target_range}",
            "answer_sheet": "Output",
        },
        _fill_response(preserve_ranges=(("Output", preserve_range),)),
        source,
        destination,
    )

    assert result.success
    before_workbook = openpyxl.load_workbook(source, data_only=False)
    after_workbook = openpyxl.load_workbook(destination, data_only=False)
    before = before_workbook["Output"]["A1"].value
    after = after_workbook["Output"]["A1"].value
    assert type(after) is type(before)
    if isinstance(before, ArrayFormula):
        assert (after.ref, after.text) == (before.ref, before.text)
    else:
        assert isinstance(before, DataTableFormula)
        assert tuple(after) == tuple(before)
    before_workbook.close()
    after_workbook.close()


def test_preserve_rejects_conflicting_unmerge_and_restores_input(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Output",))
    workbook = openpyxl.load_workbook(source)
    sheet = workbook["Output"]
    sheet.merge_cells("A1:B1")
    sheet["A1"] = "Heading"
    workbook.save(source)
    workbook.close()
    source_bytes = source.read_bytes()
    destination = tmp_path / "output.xlsx"

    result = write_response_atomic(
        {"answer_position": "'Output'!A1:B1", "answer_sheet": "Output"},
        _fill_response(
            preserve_ranges=(("Output", "A1:B1"),),
            unmerge_ranges=(("Output", "A1:B1"),),
        ),
        source,
        destination,
    )

    assert not result.success
    assert result.failure_codes == (FailureCode.WORKBOOK_WRITE_FAILED,)
    assert destination.read_bytes() == source_bytes


def test_null_non_anchor_merged_writes_are_skipped_without_unmerging(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Output",))
    workbook = openpyxl.load_workbook(source)
    workbook["Output"].merge_cells("A1:B1")
    workbook["Output"]["A1"] = "Heading"
    workbook.save(source)
    workbook.close()
    destination = tmp_path / "output.xlsx"

    result = write_response_atomic(
        {"answer_position": "'Output'!A1:B1", "answer_sheet": "Output"},
        _response(("Output", "A1", "Heading"), ("Output", "B1", None)),
        source,
        destination,
    )

    assert result.success
    written = openpyxl.load_workbook(destination, data_only=False)
    assert [str(item) for item in written["Output"].merged_cells.ranges] == ["A1:B1"]
    assert written["Output"]["A1"].value == "Heading"
    written.close()


def test_exact_intersecting_unmerge_action_runs_before_writes(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Output",))
    workbook = openpyxl.load_workbook(source)
    workbook["Output"].merge_cells("H2:O10")
    workbook.save(source)
    workbook.close()
    destination = tmp_path / "output.xlsx"
    response = _fill_response(
        ("Output", "H2:H10", "=ROW()"),
        unmerge_ranges=(("Output", "H2:O10"),),
    )

    result = write_response_atomic(
        {"answer_position": "'Output'!H2:H10", "answer_sheet": "Output"},
        response,
        source,
        destination,
    )

    assert result.success
    written = openpyxl.load_workbook(destination, data_only=False)
    assert "H2:O10" not in {str(item) for item in written["Output"].merged_cells.ranges}
    assert written["Output"]["H10"].value == "=ROW()"
    written.close()


def test_unmerge_actions_reject_duplicates_non_exact_and_non_target_ranges(
    tmp_path: Path,
) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Output",))
    workbook = openpyxl.load_workbook(source)
    workbook["Output"].merge_cells("A1:B2")
    workbook["Output"].merge_cells("C1:D2")
    contract = build_target_contract(
        {"answer_position": "'Output'!A1", "answer_sheet": "Output"},
        workbook,
    )
    response = _fill_response(
        cells=(("Output", "A1", 1),),
        unmerge_ranges=(
            ("Output", "A1:B2"),
            ("Output", "A1:B2"),
            ("Output", "A1:A2"),
            ("Output", "C1:D2"),
            ("output", "A1:B2"),
        ),
    )

    result = validate_response(contract, response, workbook=workbook)
    workbook.close()

    assert result.failure_codes == (
        FailureCode.DUPLICATE_UNMERGE_RANGES,
        FailureCode.EXTRA_UNMERGE_RANGES,
    )
    assert [item.as_text() for item in result.unmerge_duplicates] == ["Output!A1:B2"]
    assert {item.as_text() for item in result.unmerge_extras} == {
        "Output!A1:A2",
        "Output!C1:D2",
        "output!A1:B2",
    }


def test_unqualified_missing_answer_sheet_writes_to_active_without_creating_it(
    tmp_path: Path,
) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Active", "Existing"))
    destination = tmp_path / "output.xlsx"
    task = {"answer_position": "$B$2", "answer_sheet": "Absent"}

    result = write_response_atomic(
        task,
        _response(("Active", "B2", 42)),
        source,
        destination,
    )

    assert result.success
    workbook = openpyxl.load_workbook(destination)
    assert workbook.sheetnames == ["Active", "Existing"]
    assert workbook["Active"]["B2"].value == 42
    workbook.close()


def test_failed_atomic_write_is_a_byte_identical_pristine_copy(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    destination = tmp_path / "output.xlsx"
    destination.write_bytes(b"stale output")
    task = {"answer_position": "'Output'!A1:A2", "answer_sheet": "Output"}
    invalid = _response(("Output", "A1", 1), ("Output", "B1", 2))

    result = write_response_atomic(task, invalid, source, destination)

    assert not result.success
    assert result.failure_codes == (
        FailureCode.MISSING_CELLS,
        FailureCode.EXTRA_CELLS,
    )
    assert destination.read_bytes() == source.read_bytes()


def test_failed_round_trip_verification_never_replaces_with_candidate(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    destination = tmp_path / "output.xlsx"
    destination.write_bytes(b"stale output")

    def reject_candidate(*args: object, **kwargs: object) -> None:
        raise ContractError(FailureCode.WORKBOOK_WRITE_FAILED, "round-trip mismatch")

    monkeypatch.setattr(contract_module, "_verify_saved_workbook", reject_candidate)
    result = write_response_atomic(
        {"answer_position": "'Output'!A1", "answer_sheet": "Output"},
        _response(("Output", "A1", "=1+1")),
        source,
        destination,
    )

    assert not result.success
    assert result.failure_codes == (FailureCode.WORKBOOK_WRITE_FAILED,)
    assert destination.read_bytes() == source.read_bytes()


def test_invalid_json_also_produces_a_byte_identical_pristine_copy(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx")
    destination = tmp_path / "output.xlsx"

    result = write_response_atomic(
        {"answer_position": "A1", "answer_sheet": "Input"},
        "not json",
        source,
        destination,
    )

    assert result.failure_codes == (FailureCode.INVALID_RESPONSE_JSON,)
    assert destination.read_bytes() == source.read_bytes()


def test_missing_input_has_a_stable_failure_without_creating_output(tmp_path: Path) -> None:
    destination = tmp_path / "output.xlsx"

    result = write_response_atomic(
        {"answer_position": "A1", "answer_sheet": "Input"},
        _response(("Input", "A1", 1)),
        tmp_path / "missing.xlsx",
        destination,
    )

    assert result.failure_codes == (FailureCode.INVALID_INPUT_WORKBOOK,)
    assert not destination.exists()


def test_trace_and_prediction_helpers_use_stable_failure_codes(tmp_path: Path) -> None:
    source = _workbook(tmp_path / "input.xlsx", ("Input", "Output"))
    destination = tmp_path / "outputs" / "task-1.xlsx"
    result = write_response_atomic(
        {"answer_position": "'Output'!A1:A2", "answer_sheet": "Output"},
        _response(("Output", "A1", 1)),
        source,
        destination,
    )

    trace = make_trace_record(
        step=1,
        model="Qwen/Qwen3.8-27B",
        prompt="prompt",
        response=json.dumps({"cells": []}),
        input_tokens=10,
        output_tokens=4,
        latency_ms=25,
        write_result=result,
    )
    prediction = make_prediction_record("task-1", "outputs/task-1.xlsx", result)

    assert trace["failure_codes"] == ["missing_cells"]
    assert prediction == {
        "id": "task-1",
        "output": "outputs/task-1.xlsx",
        "status": "error:missing_cells",
        "failure_codes": ["missing_cells"],
    }
