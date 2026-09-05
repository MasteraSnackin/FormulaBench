from __future__ import annotations

import asyncio
import json
from pathlib import Path

import openpyxl
import pytest

import formulabench.artifacts as artifacts
import formulabench.runner as runner_module
from formulabench.artifacts import (
    ArtifactLayout,
    ResumeStateError,
    atomic_write_jsonl,
    atomic_write_text,
    load_dataset_manifest,
    load_resume_state,
    recover_retry_transactions,
    sha256_file,
)
from formulabench.constants import MODEL_ID, MODEL_PROVENANCE
from formulabench.provider import Completion, ProviderConfigurationError, ProviderResponseError
from formulabench.runner import TaskRun, run_tasks
from formulabench.validate_out import validate_output


class FakeProvider:
    model = MODEL_ID
    renderer_name = "fake-test-renderer"

    def __init__(self, response: str) -> None:
        self.response = response

    async def complete(self, *, system_prompt: str, user_prompt: str) -> Completion:
        assert system_prompt
        assert user_prompt
        return Completion(
            text=self.response,
            input_tokens=123,
            output_tokens=17,
            model=self.model,
            renderer=self.renderer_name,
            stop_reason="stop",
            parse_termination="stop_sequence",
            response_format="strict_json_fallback",
            model_provenance=MODEL_PROVENANCE,
        )


class FailingProvider:
    model = MODEL_ID
    renderer_name = "fake-test-renderer"

    async def complete(self, *, system_prompt: str, user_prompt: str) -> Completion:
        raise RuntimeError("provider failed with secret-value-never-log")


class FailingConfiguredProvider:
    model = MODEL_ID
    renderer_name = "fake-test-renderer"

    async def complete(self, *, system_prompt: str, user_prompt: str) -> Completion:
        raise ProviderConfigurationError(
            "secret provider response body",
            status_code=422,
            underlying_exception_class="Unsafe;secret-class",
        )


class RejectedResponseProvider:
    model = MODEL_ID
    renderer_name = "fake-test-renderer"

    async def complete(self, *, system_prompt: str, user_prompt: str) -> Completion:
        raise ProviderResponseError(
            "private malformed model output",
            stop_reason="stop",
            parse_termination="stop_sequence",
            response_rejection="rejected_tool_grammar",
            input_tokens=321,
            output_tokens=22,
        )


class CountingProvider(FakeProvider):
    def __init__(self, response: str) -> None:
        super().__init__(response)
        self.calls = 0

    async def complete(self, *, system_prompt: str, user_prompt: str) -> Completion:
        self.calls += 1
        return await super().complete(system_prompt=system_prompt, user_prompt=user_prompt)


def _dataset(tmp_path: Path) -> tuple[Path, Path]:
    dataset = tmp_path / "dataset"
    task_dir = dataset / "spreadsheet" / "task-1"
    task_dir.mkdir(parents=True)
    workbook_path = task_dir / "1_task-1_init.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Inputs"
    sheet["A1"] = 2
    sheet["A2"] = 3
    sheet["B2"] = None
    workbook.save(workbook_path)
    workbook.close()
    manifest = [
        {
            "id": "task-1",
            "instruction": "Put the sum of A1:A2 in B2.",
            "spreadsheet_path": "spreadsheet/task-1",
            "instruction_type": "Cell-Level Manipulation",
            "answer_position": "B2",
            "answer_sheet": "Inputs",
            "data_position": "Inputs!A1:A2",
        }
    ]
    (dataset / "dataset.json").write_text(json.dumps(manifest), encoding="utf-8")
    return dataset, workbook_path


def _multi_dataset(tmp_path: Path, task_ids: tuple[str, ...]) -> Path:
    dataset = tmp_path / "dataset"
    manifest = []
    for index, task_id in enumerate(task_ids, start=1):
        task_dir = dataset / "spreadsheet" / task_id
        task_dir.mkdir(parents=True)
        workbook_path = task_dir / f"{index}_{task_id}_init.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Inputs"
        sheet["A1"] = index
        sheet["A2"] = index + 1
        workbook.save(workbook_path)
        workbook.close()
        manifest.append(
            {
                "id": task_id,
                "instruction": "Put the sum of A1:A2 in B2.",
                "spreadsheet_path": f"spreadsheet/{task_id}",
                "answer_position": "B2",
                "answer_sheet": "Inputs",
                "data_position": "Inputs!A1:A2",
            }
        )
    (dataset / "dataset.json").write_text(json.dumps(manifest), encoding="utf-8")
    return dataset


def test_runner_writes_validated_workbook_and_submission_records(tmp_path: Path) -> None:
    dataset, input_path = _dataset(tmp_path)
    input_sha = sha256_file(input_path)
    tasks = load_dataset_manifest(dataset)
    layout = ArtifactLayout.initialise(tmp_path / "out")
    atomic_write_text(layout.run_log, "test run\n")
    provider = FakeProvider('{"cells":[{"sheet":"Inputs","cell":"B2","value":"=SUM(A1:A2)"}]}')

    results = asyncio.run(run_tasks(tasks=tasks, provider=provider, layout=layout, concurrency=1))

    assert results[0].success is True
    assert sha256_file(input_path) == input_sha
    output = openpyxl.load_workbook(layout.outputs / "task-1.xlsx", data_only=False)
    assert output["Inputs"]["B2"].value == "=SUM(A1:A2)"
    output.close()
    trace = json.loads((layout.traces / "task-1.jsonl").read_text(encoding="utf-8"))
    assert trace["model"] == MODEL_ID
    assert trace["stop_reason"] == "stop"
    assert trace["parse_termination"] == "stop_sequence"
    assert trace["response_format"] == "strict_json_fallback"
    assert trace["response_rejection"] is None
    assert trace["request"] == {
        "model": MODEL_ID,
        "model_provenance": "configured_base_model",
        "max_output_tokens": 8192,
        "max_retries": 0,
        "num_samples": 1,
        "reasoning_effort": None,
        "renderer": "transformers-chat-template/qwen3.8-disable-thinking-pinned",
        "streaming": False,
        "stop": "<|im_end|>",
        "temperature": 0.0,
        "telemetry": "disabled",
        "thinking": "disabled",
        "tool": "submit_spreadsheet_answer",
        "tool_enforcement": "rendered_prompt_plus_strict_parser",
        "tokenizer_revision": "1d4bf0f2ff6012fd82039f2fa52739d0dd7c60c0",
        "transport": "tinker-native-sampling/0.27.1",
    }
    report = validate_output(dataset, layout.root, expected_model=MODEL_ID)
    assert report.ok, report.as_dict()


def test_runner_records_only_safe_provider_configuration_diagnostics(tmp_path: Path) -> None:
    dataset, _ = _dataset(tmp_path)
    tasks = load_dataset_manifest(dataset)
    layout = ArtifactLayout.initialise(tmp_path / "out")
    atomic_write_text(layout.run_log, "test run\n")

    results = asyncio.run(
        run_tasks(
            tasks=tasks,
            provider=FailingConfiguredProvider(),
            layout=layout,
            concurrency=1,
        )
    )

    assert results[0].success is False
    trace_text = (layout.traces / "task-1.jsonl").read_text(encoding="utf-8")
    trace = json.loads(trace_text)
    assert trace["error"] == (
        "model_call_failed:ProviderConfigurationError:http_status=422:"
        "underlying_exception_class=HTTPError"
    )
    assert trace["stop_reason"] is None
    assert trace["request"]["tool_enforcement"] == "rendered_prompt_plus_strict_parser"
    assert trace["request"]["thinking"] == "disabled"
    assert "secret provider response body" not in trace_text


def test_runner_preserves_safe_native_termination_facts_for_rejected_output(
    tmp_path: Path,
) -> None:
    dataset, _ = _dataset(tmp_path)
    tasks = load_dataset_manifest(dataset)
    layout = ArtifactLayout.initialise(tmp_path / "out")
    atomic_write_text(layout.run_log, "test run\n")

    results = asyncio.run(
        run_tasks(
            tasks=tasks,
            provider=RejectedResponseProvider(),
            layout=layout,
            concurrency=1,
        )
    )

    assert results[0].success is False
    trace_text = (layout.traces / "task-1.jsonl").read_text(encoding="utf-8")
    trace = json.loads(trace_text)
    assert trace["input_tokens"] == 321
    assert trace["output_tokens"] == 22
    assert trace["stop_reason"] == "stop"
    assert trace["parse_termination"] == "stop_sequence"
    assert trace["response_format"] is None
    assert trace["response_rejection"] == "rejected_tool_grammar"
    assert "private malformed model output" not in trace_text


def test_runner_copies_pristine_input_and_does_not_reflect_provider_error(
    tmp_path: Path,
) -> None:
    dataset, input_path = _dataset(tmp_path)
    tasks = load_dataset_manifest(dataset)
    layout = ArtifactLayout.initialise(tmp_path / "out")
    atomic_write_text(layout.run_log, "test run\n")

    results = asyncio.run(
        run_tasks(
            tasks=tasks,
            provider=FailingProvider(),
            layout=layout,
            concurrency=1,
        )
    )

    assert results[0].success is False
    assert sha256_file(layout.outputs / "task-1.xlsx") == sha256_file(input_path)
    trace_text = (layout.traces / "task-1.jsonl").read_text(encoding="utf-8")
    assert "secret-value-never-log" not in trace_text
    assert "model_call_failed:RuntimeError" in trace_text
    report = validate_output(
        dataset,
        layout.root,
        expected_model=MODEL_ID,
        secret_values=("secret-value-never-log",),
    )
    assert report.ok, report.as_dict()


def test_interruption_keeps_atomic_checkpoint_and_resume_skips_completed_task(
    tmp_path: Path,
) -> None:
    dataset = _multi_dataset(tmp_path, ("task-1", "task-2"))
    tasks = load_dataset_manifest(dataset)
    layout = ArtifactLayout.initialise(tmp_path / "out")
    atomic_write_text(layout.run_log, "attempt one\n")
    response = '{"cells":[{"sheet":"Inputs","cell":"B2","value":"=SUM(A1:A2)"}]}'

    class FirstThenBlockProvider(FakeProvider):
        def __init__(self) -> None:
            super().__init__(response)
            self.calls = 0
            self.second_started = asyncio.Event()
            self.never_release = asyncio.Event()

        async def complete(self, *, system_prompt: str, user_prompt: str) -> Completion:
            self.calls += 1
            if self.calls == 1:
                return await super().complete(
                    system_prompt=system_prompt,
                    user_prompt=user_prompt,
                )
            self.second_started.set()
            await self.never_release.wait()
            raise AssertionError("unreachable")

    async def interrupt_after_checkpoint() -> None:
        provider = FirstThenBlockProvider()
        run = asyncio.create_task(
            run_tasks(tasks=tasks, provider=provider, layout=layout, concurrency=1)
        )
        await asyncio.wait_for(provider.second_started.wait(), timeout=3)
        for _ in range(100):
            if layout.predictions.stat().st_size:
                break
            await asyncio.sleep(0.01)
        else:
            pytest.fail("first task did not reach the durable checkpoint")
        run.cancel()
        with pytest.raises(asyncio.CancelledError):
            await run

    asyncio.run(interrupt_after_checkpoint())

    checkpoint = [json.loads(line) for line in layout.predictions.read_text().splitlines()]
    assert [record["id"] for record in checkpoint] == ["task-1"]
    state = load_resume_state(layout, tasks, expected_model=MODEL_ID)
    assert state.completed_ids == frozenset({"task-1"})

    resumed_provider = CountingProvider(response)
    results = asyncio.run(
        run_tasks(
            tasks=tasks,
            provider=resumed_provider,
            layout=layout,
            concurrency=2,
            existing_predictions=state.predictions,
        )
    )

    assert resumed_provider.calls == 1
    assert [result.task_id for result in results] == ["task-1", "task-2"]
    final_records = [json.loads(line) for line in layout.predictions.read_text().splitlines()]
    assert [record["id"] for record in final_records] == ["task-1", "task-2"]
    assert all(result.success for result in results)


def test_explicit_retry_replaces_only_failure_and_preserves_success(
    tmp_path: Path,
) -> None:
    dataset = _multi_dataset(tmp_path, ("task-1", "task-2"))
    tasks = load_dataset_manifest(dataset)
    layout = ArtifactLayout.initialise(tmp_path / "out")
    atomic_write_text(layout.run_log, "attempt one\n")
    response = '{"cells":[{"sheet":"Inputs","cell":"B2","value":"=SUM(A1:A2)"}]}'

    class SuccessThenFailureProvider(FakeProvider):
        def __init__(self) -> None:
            super().__init__(response)
            self.calls = 0

        async def complete(self, *, system_prompt: str, user_prompt: str) -> Completion:
            self.calls += 1
            if self.calls == 2:
                raise RuntimeError("first-attempt failure")
            return await super().complete(system_prompt=system_prompt, user_prompt=user_prompt)

    first_results = asyncio.run(
        run_tasks(
            tasks=tasks,
            provider=SuccessThenFailureProvider(),
            layout=layout,
            concurrency=1,
        )
    )
    assert [result.success for result in first_results] == [True, False]

    state = load_resume_state(layout, tasks, expected_model=MODEL_ID)
    assert state.successful_ids == frozenset({"task-1"})
    task_one_prediction = dict(state.predictions["task-1"])
    task_one_output = (layout.outputs / "task-1.xlsx").read_bytes()
    task_one_trace = (layout.traces / "task-1.jsonl").read_bytes()

    retry_provider = CountingProvider(response)
    retry_results = asyncio.run(
        run_tasks(
            tasks=tasks,
            provider=retry_provider,
            layout=layout,
            concurrency=1,
            existing_predictions=state.predictions,
            retry_task_ids=state.completed_ids.difference(state.successful_ids),
        )
    )

    assert retry_provider.calls == 1
    assert all(result.success for result in retry_results)
    records = [json.loads(line) for line in layout.predictions.read_text().splitlines()]
    assert records[0] == task_one_prediction
    assert records[1]["status"] == "ok"
    assert (layout.outputs / "task-1.xlsx").read_bytes() == task_one_output
    assert (layout.traces / "task-1.jsonl").read_bytes() == task_one_trace


def test_failed_explicit_retry_safely_replaces_the_failure_trace(tmp_path: Path) -> None:
    dataset, input_path = _dataset(tmp_path)
    tasks = load_dataset_manifest(dataset)
    layout = ArtifactLayout.initialise(tmp_path / "out")
    atomic_write_text(layout.run_log, "attempt one\n")
    asyncio.run(run_tasks(tasks=tasks, provider=FailingProvider(), layout=layout, concurrency=1))
    first_trace = (layout.traces / "task-1.jsonl").read_text(encoding="utf-8")
    assert "model_call_failed:RuntimeError" in first_trace
    state = load_resume_state(layout, tasks, expected_model=MODEL_ID)

    class DifferentFailingProvider:
        calls = 0

        async def complete(self, *, system_prompt: str, user_prompt: str) -> Completion:
            del system_prompt, user_prompt
            type(self).calls += 1
            raise ValueError("second-attempt private details")

    results = asyncio.run(
        run_tasks(
            tasks=tasks,
            provider=DifferentFailingProvider(),
            layout=layout,
            concurrency=1,
            existing_predictions=state.predictions,
            retry_task_ids=frozenset({"task-1"}),
        )
    )

    assert DifferentFailingProvider.calls == 1
    assert results[0].success is False
    second_trace = (layout.traces / "task-1.jsonl").read_text(encoding="utf-8")
    assert "model_call_failed:ValueError" in second_trace
    assert "second-attempt private details" not in second_trace
    assert second_trace != first_trace
    assert sha256_file(layout.outputs / "task-1.xlsx") == sha256_file(input_path)
    replaced_state = load_resume_state(layout, tasks, expected_model=MODEL_ID)
    assert replaced_state.completed_ids == frozenset({"task-1"})
    assert replaced_state.successful_ids == frozenset()


def test_interrupted_failure_retry_keeps_pending_old_checkpoint(tmp_path: Path) -> None:
    dataset = _multi_dataset(tmp_path, ("task-1", "task-2"))
    tasks = load_dataset_manifest(dataset)
    layout = ArtifactLayout.initialise(tmp_path / "out")
    atomic_write_text(layout.run_log, "attempt one\n")
    asyncio.run(run_tasks(tasks=tasks, provider=FailingProvider(), layout=layout, concurrency=1))
    original_state = load_resume_state(layout, tasks, expected_model=MODEL_ID)
    original_task_two_prediction = dict(original_state.predictions["task-2"])
    original_task_two_output = (layout.outputs / "task-2.xlsx").read_bytes()
    original_task_two_trace = (layout.traces / "task-2.jsonl").read_bytes()
    response = '{"cells":[{"sheet":"Inputs","cell":"B2","value":"=SUM(A1:A2)"}]}'

    class FirstSucceedsThenBlocks(FakeProvider):
        def __init__(self) -> None:
            super().__init__(response)
            self.calls = 0
            self.second_started = asyncio.Event()
            self.never_release = asyncio.Event()

        async def complete(self, *, system_prompt: str, user_prompt: str) -> Completion:
            self.calls += 1
            if self.calls == 1:
                return await super().complete(
                    system_prompt=system_prompt,
                    user_prompt=user_prompt,
                )
            self.second_started.set()
            await self.never_release.wait()
            raise AssertionError("unreachable")

    async def interrupt_second_retry() -> None:
        provider = FirstSucceedsThenBlocks()
        run = asyncio.create_task(
            run_tasks(
                tasks=tasks,
                provider=provider,
                layout=layout,
                concurrency=1,
                existing_predictions=original_state.predictions,
                retry_task_ids=original_state.completed_ids,
            )
        )
        await asyncio.wait_for(provider.second_started.wait(), timeout=3)
        for _ in range(100):
            records = [json.loads(line) for line in layout.predictions.read_text().splitlines()]
            if records[0]["status"] == "ok":
                break
            await asyncio.sleep(0.01)
        else:
            pytest.fail("first failure replacement did not reach the durable checkpoint")
        run.cancel()
        with pytest.raises(asyncio.CancelledError):
            await run

    asyncio.run(interrupt_second_retry())

    records = [json.loads(line) for line in layout.predictions.read_text().splitlines()]
    assert records[0]["status"] == "ok"
    assert records[1] == original_task_two_prediction
    assert (layout.outputs / "task-2.xlsx").read_bytes() == original_task_two_output
    assert (layout.traces / "task-2.jsonl").read_bytes() == original_task_two_trace
    interrupted_state = load_resume_state(layout, tasks, expected_model=MODEL_ID)
    assert interrupted_state.completed_ids == frozenset({"task-1", "task-2"})
    assert interrupted_state.successful_ids == frozenset({"task-1"})


@pytest.mark.parametrize("phase", ["prepared", "output", "trace", "prediction"])
def test_retry_transaction_recovers_every_publication_phase(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    phase: str,
) -> None:
    dataset, _ = _dataset(tmp_path)
    tasks = load_dataset_manifest(dataset)
    layout = ArtifactLayout.initialise(tmp_path / "out")
    atomic_write_text(layout.run_log, "attempt one\n")
    asyncio.run(run_tasks(tasks=tasks, provider=FailingProvider(), layout=layout, concurrency=1))
    old_state = load_resume_state(layout, tasks, expected_model=MODEL_ID)
    response = '{"cells":[{"sheet":"Inputs","cell":"B2","value":"=SUM(A1:A2)"}]}'

    class SimulatedHardStop(BaseException):
        pass

    original_publish = artifacts._publish_retry_file
    original_cleanup = artifacts._remove_retry_workspace
    publish_calls = 0

    def interrupted_publish(source: Path, destination: Path, publish_path: Path) -> None:
        nonlocal publish_calls
        if phase == "prepared" and publish_calls == 0:
            raise SimulatedHardStop
        original_publish(source, destination, publish_path)
        publish_calls += 1
        if phase == "output" and publish_calls == 1:
            raise SimulatedHardStop
        if phase == "trace" and publish_calls == 2:
            raise SimulatedHardStop

    def interrupted_cleanup(
        current_layout: ArtifactLayout, workspace: artifacts.RetryWorkspace
    ) -> None:
        if phase == "prediction":
            raise SimulatedHardStop
        original_cleanup(current_layout, workspace)

    monkeypatch.setattr(artifacts, "_publish_retry_file", interrupted_publish)
    monkeypatch.setattr(artifacts, "_remove_retry_workspace", interrupted_cleanup)
    with pytest.raises(SimulatedHardStop):
        asyncio.run(
            run_tasks(
                tasks=tasks,
                provider=FakeProvider(response),
                layout=layout,
                concurrency=1,
                existing_predictions=old_state.predictions,
                retry_task_ids=frozenset({"task-1"}),
            )
        )
    assert (layout.root / ".retry" / "task-1" / "meta.json").is_file()
    interrupted_prediction = json.loads(layout.predictions.read_text(encoding="utf-8"))
    if phase == "prediction":
        assert interrupted_prediction["status"] == "ok"
    else:
        assert interrupted_prediction == dict(old_state.predictions["task-1"])

    monkeypatch.setattr(artifacts, "_publish_retry_file", original_publish)
    monkeypatch.setattr(artifacts, "_remove_retry_workspace", original_cleanup)
    recover_retry_transactions(layout, tasks, expected_model=MODEL_ID)

    recovered = load_resume_state(layout, tasks, expected_model=MODEL_ID)
    assert recovered.successful_ids == frozenset({"task-1"})
    workbook = openpyxl.load_workbook(layout.outputs / "task-1.xlsx", data_only=False)
    assert workbook["Inputs"]["B2"].value == "=SUM(A1:A2)"
    workbook.close()
    assert {path.name for path in layout.outputs.iterdir()} == {"task-1.xlsx"}
    assert {path.name for path in layout.traces.iterdir()} == {"task-1.jsonl"}
    assert not (layout.root / ".retry").exists()
    recover_retry_transactions(layout, tasks, expected_model=MODEL_ID)


def test_cleanup_failure_cannot_reinstall_old_failure_during_later_checkpoint(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    dataset = _multi_dataset(tmp_path, ("task-1", "task-2"))
    tasks = load_dataset_manifest(dataset)
    layout = ArtifactLayout.initialise(tmp_path / "out")
    atomic_write_text(layout.run_log, "attempt one\n")
    asyncio.run(
        run_tasks(tasks=tasks[:1], provider=FailingProvider(), layout=layout, concurrency=1)
    )
    old_state = load_resume_state(layout, tasks, expected_model=MODEL_ID)
    response = '{"cells":[{"sheet":"Inputs","cell":"B2","value":"=SUM(A1:A2)"}]}'
    original_cleanup = artifacts._remove_retry_workspace
    cleanup_calls = 0

    def fail_once_after_marker_removal(
        current_layout: ArtifactLayout, workspace: artifacts.RetryWorkspace
    ) -> None:
        nonlocal cleanup_calls
        cleanup_calls += 1
        if cleanup_calls == 1:
            (workspace.directory / "meta.json").unlink()
            artifacts._fsync_directory(workspace.directory)
            raise ResumeStateError("simulated_cleanup_failure", task_id=workspace.task_id)
        original_cleanup(current_layout, workspace)

    monkeypatch.setattr(artifacts, "_remove_retry_workspace", fail_once_after_marker_removal)
    results = asyncio.run(
        run_tasks(
            tasks=tasks,
            provider=FakeProvider(response),
            layout=layout,
            concurrency=2,
            existing_predictions=old_state.predictions,
            retry_task_ids=frozenset({"task-1"}),
        )
    )

    assert all(result.success for result in results)
    final_state = load_resume_state(layout, tasks, expected_model=MODEL_ID)
    assert final_state.successful_ids == frozenset({"task-1", "task-2"})
    assert not (layout.root / ".retry").exists()


def test_worker_exception_still_cancels_siblings_and_cleans_unprepared_retries(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    dataset = _multi_dataset(tmp_path, ("task-1", "task-2"))
    tasks = load_dataset_manifest(dataset)
    layout = ArtifactLayout.initialise(tmp_path / "out")
    atomic_write_text(layout.run_log, "attempt one\n")
    asyncio.run(run_tasks(tasks=tasks, provider=FailingProvider(), layout=layout, concurrency=1))
    old_state = load_resume_state(layout, tasks, expected_model=MODEL_ID)

    async def exercise() -> None:
        sibling_started = asyncio.Event()
        sibling_cancelled = asyncio.Event()
        never_release = asyncio.Event()

        async def exceptional_run_task(**kwargs: object) -> TaskRun:
            task = kwargs["task"]
            assert isinstance(task, artifacts.DatasetTask)
            if task.id == "task-1":
                await sibling_started.wait()
                raise OSError("simulated artefact write failure")
            sibling_started.set()
            try:
                await never_release.wait()
            finally:
                sibling_cancelled.set()
            raise AssertionError("unreachable")

        monkeypatch.setattr(runner_module, "run_task", exceptional_run_task)
        with pytest.raises(OSError, match="simulated artefact write failure"):
            await run_tasks(
                tasks=tasks,
                provider=FakeProvider("{}"),
                layout=layout,
                concurrency=2,
                existing_predictions=old_state.predictions,
                retry_task_ids=old_state.completed_ids,
            )
        assert sibling_cancelled.is_set()

    asyncio.run(exercise())
    assert not (layout.root / ".retry").exists()
    unchanged = load_resume_state(layout, tasks, expected_model=MODEL_ID)
    assert unchanged.successful_ids == frozenset()


def test_worker_exception_remains_primary_when_retry_cleanup_also_fails(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    dataset = _multi_dataset(tmp_path, ("task-1", "task-2"))
    tasks = load_dataset_manifest(dataset)
    layout = ArtifactLayout.initialise(tmp_path / "out")
    atomic_write_text(layout.run_log, "attempt one\n")
    asyncio.run(run_tasks(tasks=tasks, provider=FailingProvider(), layout=layout, concurrency=1))
    old_state = load_resume_state(layout, tasks, expected_model=MODEL_ID)
    discard_attempts: list[str] = []
    original_discard = runner_module.discard_retry_workspace

    def fail_first_discard(
        current_layout: ArtifactLayout,
        workspace: artifacts.RetryWorkspace,
    ) -> None:
        discard_attempts.append(workspace.task_id)
        if workspace.task_id == "task-1":
            raise RuntimeError("simulated cleanup failure")
        original_discard(current_layout, workspace)

    async def exercise() -> None:
        sibling_started = asyncio.Event()
        sibling_cleanup_finished = asyncio.Event()
        never_release = asyncio.Event()

        async def exceptional_run_task(**kwargs: object) -> TaskRun:
            task = kwargs["task"]
            assert isinstance(task, artifacts.DatasetTask)
            if task.id == "task-1":
                await sibling_started.wait()
                raise OSError("primary worker failure")
            sibling_started.set()
            try:
                await never_release.wait()
            finally:
                # An awaited cleanup step proves run_tasks settles the cancelled
                # sibling rather than merely requesting cancellation.
                await asyncio.sleep(0)
                sibling_cleanup_finished.set()
            raise AssertionError("unreachable")

        monkeypatch.setattr(runner_module, "run_task", exceptional_run_task)
        monkeypatch.setattr(runner_module, "discard_retry_workspace", fail_first_discard)
        with pytest.raises(OSError, match="primary worker failure"):
            await run_tasks(
                tasks=tasks,
                provider=FakeProvider("{}"),
                layout=layout,
                concurrency=2,
                existing_predictions=old_state.predictions,
                retry_task_ids=old_state.completed_ids,
            )
        assert sibling_cleanup_finished.is_set()

    asyncio.run(exercise())
    assert discard_attempts == ["task-1", "task-2"]
    assert (layout.root / ".retry" / "task-1").is_dir()
    assert not (layout.root / ".retry" / "task-2").exists()


def test_resume_rejects_malformed_prediction_trace_and_stale_orphan(tmp_path: Path) -> None:
    dataset = _multi_dataset(tmp_path, ("task-1", "task-2"))
    tasks = load_dataset_manifest(dataset)
    layout = ArtifactLayout.initialise(tmp_path / "out")
    atomic_write_text(layout.run_log, "test run\n")
    provider = FakeProvider('{"cells":[{"sheet":"Inputs","cell":"B2","value":"=SUM(A1:A2)"}]}')
    asyncio.run(run_tasks(tasks=tasks[:1], provider=provider, layout=layout, concurrency=1))

    original_predictions = layout.predictions.read_bytes()
    record = json.loads(original_predictions)
    record["output"] = "outputs/stale.xlsx"
    atomic_write_jsonl(layout.predictions, [record])
    with pytest.raises(ResumeStateError, match="prediction_output_layout"):
        load_resume_state(layout, tasks, expected_model=MODEL_ID)

    layout.predictions.write_bytes(original_predictions)
    original_trace = (layout.traces / "task-1.jsonl").read_bytes()
    (layout.traces / "task-1.jsonl").write_text("not-json\n", encoding="utf-8")
    with pytest.raises(ResumeStateError, match="trace_invalid_json"):
        load_resume_state(layout, tasks, expected_model=MODEL_ID)

    (layout.traces / "task-1.jsonl").write_bytes(original_trace)
    orphan = layout.outputs / "task-2.xlsx"
    orphan.write_bytes(tasks[1].init_xlsx.read_bytes())
    with pytest.raises(ResumeStateError, match="orphan_output"):
        load_resume_state(layout, tasks, expected_model=MODEL_ID)
    assert orphan.exists()
