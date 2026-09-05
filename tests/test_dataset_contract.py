from __future__ import annotations

import json
import os
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from formulabench.context import build_context
from formulabench.contract import (
    ResponseFill,
    SheetResolution,
    SpreadsheetResponse,
    TargetProvenance,
    build_target_contract,
    validate_response,
)
from formulabench.prompts import (
    prompt_scaffold,
    target_summary_lines,
    workbook_char_budget,
)
from sb import answer_cells


def _dataset_dir() -> Path:
    configured = os.environ.get("FORMULABENCH_DATASET_DIR")
    if configured:
        return Path(configured)
    return Path(__file__).resolve().parents[1] / "data" / "spreadsheetbench_verified_400"


def _metadata_and_initial_workbook(task: dict[str, object], dataset_dir: Path) -> Path:
    spreadsheet_path = task["spreadsheet_path"]
    assert isinstance(spreadsheet_path, str)
    candidates = sorted((dataset_dir / spreadsheet_path).glob("*init*.xlsx"))
    assert len(candidates) == 1, f"expected one initial workbook for task {task['id']}"
    return candidates[0]


def test_every_dataset_target_contract_resolves_from_metadata_and_initial_workbook() -> None:
    """Exercise public target metadata without opening or locating golden files."""

    dataset_dir = _dataset_dir()
    metadata_path = dataset_dir / "dataset.json"
    if not metadata_path.exists():
        pytest.skip("SpreadsheetBench Verified 400 is not downloaded")

    tasks = json.loads(metadata_path.read_text(encoding="utf-8"))
    assert len(tasks) == 400
    dynamic_ranges = []
    created_ranges = []
    fallback_ranges = []

    for task in tasks:
        initial_path = _metadata_and_initial_workbook(task, dataset_dir)
        workbook = openpyxl.load_workbook(initial_path, read_only=True, data_only=False)
        original_sheets = tuple(workbook.sheetnames)

        contract = build_target_contract(task, workbook)

        assert contract.ranges
        assert tuple(workbook.sheetnames) == original_sheets
        for target_range in contract.ranges:
            assert target_range.sheet
            if target_range.provenance is TargetProvenance.EXPLICIT:
                assert target_range.requested_sheet == target_range.sheet
            if target_range.resolution is SheetResolution.CREATED:
                assert target_range.sheet not in original_sheets
                created_ranges.append((str(task["id"]), target_range))
            elif target_range.resolution is SheetResolution.ACTIVE_FALLBACK:
                assert target_range.sheet == contract.active_sheet
                fallback_ranges.append((str(task["id"]), target_range))
            else:
                assert target_range.sheet in original_sheets
            if target_range.dynamic:
                dynamic_ranges.append((str(task["id"]), target_range))
        workbook.close()

    # These assertions ensure the integration test covers all three unusual
    # policies, rather than silently passing on finite existing-sheet targets.
    assert created_ranges
    assert fallback_ranges
    assert dynamic_ranges
    assert {item.sheet for _, item in dynamic_ranges} == {"Sheet3", "Sheet4"}


def test_finite_targets_match_the_shipped_evaluator_sheet_and_cell_resolution() -> None:
    """Prove our stricter contract selects the same finite cells as the evaluator."""

    dataset_dir = _dataset_dir()
    metadata_path = dataset_dir / "dataset.json"
    if not metadata_path.exists():
        pytest.skip("SpreadsheetBench Verified 400 is not downloaded")

    tasks = json.loads(metadata_path.read_text(encoding="utf-8"))
    checked = 0
    for task in tasks:
        initial_path = _metadata_and_initial_workbook(task, dataset_dir)
        workbook = openpyxl.load_workbook(initial_path, read_only=False, data_only=False)
        try:
            contract = build_target_contract(task, workbook)
            if any(target.dynamic for target in contract.ranges):
                continue
            for target in contract.ranges:
                if target.resolution is SheetResolution.CREATED:
                    workbook.create_sheet(target.sheet)

            contract_cells: set[tuple[str, str]] = set()
            for target in contract.ranges:
                min_col, min_row, max_col, max_row = range_boundaries(target.cell_range)
                contract_cells.update(
                    (target.sheet, f"{get_column_letter(column)}{row}")
                    for row in range(min_row, max_row + 1)
                    for column in range(min_col, max_col + 1)
                )

            evaluator_cells = set()
            for sheet, coordinate in answer_cells(task, workbook):
                worksheet = (
                    workbook[sheet] if sheet and sheet in workbook.sheetnames else workbook.active
                )
                evaluator_cells.add((worksheet.title, coordinate))
            assert contract_cells == evaluator_cells, f"target mismatch for task {task['id']}"
            checked += 1
        finally:
            workbook.close()

    assert checked == 399


def test_dynamic_targets_match_evaluator_initial_worksheet_extent() -> None:
    """Cover the one public whole-column task without reading its golden workbook."""

    dataset_dir = _dataset_dir()
    metadata_path = dataset_dir / "dataset.json"
    if not metadata_path.exists():
        pytest.skip("SpreadsheetBench Verified 400 is not downloaded")

    tasks = json.loads(metadata_path.read_text(encoding="utf-8"))
    checked = 0
    for task in tasks:
        initial_path = _metadata_and_initial_workbook(task, dataset_dir)
        workbook = openpyxl.load_workbook(initial_path, read_only=False, data_only=False)
        try:
            contract = build_target_contract(task, workbook)
            dynamic = [target for target in contract.ranges if target.dynamic]
            if not dynamic:
                continue
            fills = []
            for target in dynamic:
                worksheet = (
                    workbook[target.sheet]
                    if target.sheet in workbook.sheetnames
                    else workbook.active
                )
                min_col, _, max_col, _ = range_boundaries(target.cell_range)
                fills.append(
                    ResponseFill(
                        sheet=target.sheet,
                        range=(
                            f"{get_column_letter(min_col)}1:"
                            f"{get_column_letter(max_col)}{worksheet.max_row}"
                        ),
                        value=None,
                    )
                )
            result = validate_response(
                contract,
                SpreadsheetResponse(cells=[], fills=fills),
                workbook=workbook,
            )
            evaluator_cells = {
                (
                    workbook[sheet].title
                    if sheet and sheet in workbook.sheetnames
                    else workbook.active.title,
                    coordinate,
                )
                for sheet, coordinate in answer_cells(task, workbook)
            }
            assert result.ok, (task["id"], result.failure_codes)
            assert {(key.sheet, key.cell) for key in result.expected} == evaluator_cells
            checked += 1
        finally:
            workbook.close()

    assert checked == 1


def test_public_13_1_context_selects_and_retains_ranges_source() -> None:
    """Regress source selection and dense evidence using only the public initial file."""

    dataset_dir = _dataset_dir()
    metadata_path = dataset_dir / "dataset.json"
    if not metadata_path.exists():
        pytest.skip("SpreadsheetBench Verified 400 is not downloaded")

    tasks = json.loads(metadata_path.read_text(encoding="utf-8"))
    task = next(item for item in tasks if str(item["id"]) == "13-1")
    initial_path = _metadata_and_initial_workbook(task, dataset_dir)
    workbook = openpyxl.load_workbook(initial_path, read_only=False, data_only=False)
    try:
        contract = build_target_contract(task, workbook)
        assert [(target.sheet, target.cell_range) for target in contract.ranges] == [
            ("LISTS", "A3:D32")
        ]
        source = workbook["RANGES"]
        source_values = {
            f"{get_column_letter(cell.column)}{cell.row}": cell.value
            for cell in source._cells.values()
            if cell.value is not None and cell.row <= 56 and cell.column <= 5
        }
        prefix, suffix = prompt_scaffold(
            instruction=str(task["instruction"]),
            target_lines=target_summary_lines(contract.ranges),
        )
    finally:
        workbook.close()

    context_task = dict(task)
    context_task["init_xlsx"] = str(initial_path)
    budget = workbook_char_budget(prefix=prefix, suffix=suffix)
    document = build_context(context_task, char_budget=budget)
    records = [json.loads(line) for line in document.jsonl.splitlines()]
    data = next(record for record in records if record["type"] == "data")
    data_rows = [record for record in records if record["type"] == "data_row"]
    compact_values: dict[str, object] = {}
    for record in data_rows:
        min_col, min_row, max_col, max_row = range_boundaries(record["range"])
        assert min_row == max_row
        assert len(record["values"]) == max_col - min_col + 1
        for offset, value in enumerate(record["values"]):
            coordinate = f"{get_column_letter(min_col + offset)}{min_row}"
            compact_values[coordinate] = value

    assert data == {
        "cell_count": 280,
        "declared": True,
        "dynamic": False,
        "range": "A1:E56",
        "requested_sheet": None,
        "resolution": "instruction_named_default",
        "sheet": "RANGES",
        "type": "data",
    }
    assert len(data_rows) == 56
    assert data_rows[0]["range"] == "A1:E1"
    assert data_rows[-1]["range"] == "A56:E56"
    assert len(source_values) == 243
    assert all(compact_values[coordinate] is not None for coordinate in source_values)
    assert compact_values["C1"] == "STAGE"
    assert compact_values["C16"] == "DATA"
    assert compact_values["C35"] == "OPERATION"
    assert compact_values["E55"] == 150
    assert document.used_chars <= budget
    assert document.included_cells >= 280
