from __future__ import annotations

import copy
import hashlib
import json
import os
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest
from openpyxl.worksheet.formula import ArrayFormula

from formulabench.artifacts import DatasetTask
from training.corpus import (
    CORPUS_NAME,
    CorpusError,
    _answer_payload,
    _json_cell_value,
    _load_and_verify_split,
    _partition_rows,
    _runtime_check,
    build_corpus,
)


def test_array_formulas_and_required_unmerges_are_representable(tmp_path: Path) -> None:
    initial_path = tmp_path / "initial.xlsx"
    initial = openpyxl.Workbook()
    initial.active.title = "Sheet1"
    initial.active.merge_cells("H2:O3")
    initial.save(initial_path)

    golden = openpyxl.Workbook()
    golden.active.title = "Sheet1"
    golden.active["H2"] = ArrayFormula(ref="H2", text="=ROW()")
    golden.active["H3"] = ArrayFormula(ref="H3", text="=ROW()")
    task = DatasetTask(
        id="development-unmerge",
        instruction="Unmerge the output and add the formulas.",
        spreadsheet_path="spreadsheet/development-unmerge",
        init_xlsx=initial_path,
        metadata={
            "id": "development-unmerge",
            "instruction": "Unmerge the output and add the formulas.",
            "spreadsheet_path": "spreadsheet/development-unmerge",
            "answer_position": "H2:H3",
            "answer_sheet": "Sheet1",
        },
    )

    payload = _answer_payload(task, initial, golden, [("Sheet1", "H2"), ("Sheet1", "H3")])

    assert payload["unmerge_ranges"] == [{"sheet": "Sheet1", "range": "H2:O3"}]
    assert [cell["value"] for cell in payload["cells"]] == ["=ROW()", "=ROW()"]
    _runtime_check(task, payload)
    initial.close()
    golden.close()


def test_unsupported_time_only_values_fail_closed() -> None:
    import datetime as dt

    with pytest.raises(CorpusError, match="time-only"):
        _json_cell_value(dt.time(12, 30))


def test_partition_is_deterministic_and_keeps_validation_separate() -> None:
    rows = [
        {
            "instruction_type": "Cell-Level Manipulation" if index % 2 else "Sheet-Level",
            "target_cell_count": 1 if index % 3 else 20,
            "task_id": f"task-{index:02d}",
        }
        for index in range(74)
    ]
    first = copy.deepcopy(rows)
    second = copy.deepcopy(rows)

    _partition_rows(first)
    _partition_rows(second)

    assert first == second
    assert sum(row["partition"] == "train" for row in first) == 59
    assert sum(row["partition"] == "validation" for row in first) == 15


def test_split_must_reproduce_before_development_ids_are_accepted(tmp_path: Path) -> None:
    payload = {
        "development_ids": [f"development-{index}" for index in range(80)],
        "held_out_ids": [f"held-out-{index}" for index in range(80)],
        "final_only_ids": [f"final-{index}" for index in range(240)],
    }
    split_path = tmp_path / "public_split.json"
    split_path.write_text(json.dumps(payload), encoding="utf-8")

    with patch("training.corpus.build_public_split", return_value=copy.deepcopy(payload)) as build:
        assert _load_and_verify_split(tmp_path, split_path) == payload
    build.assert_called_once_with(tmp_path)

    changed = copy.deepcopy(payload)
    changed["development_ids"][0] = "held-out-0"
    with (
        patch("training.corpus.build_public_split", return_value=changed),
        pytest.raises(CorpusError, match="does not reproduce"),
    ):
        _load_and_verify_split(tmp_path, split_path)


def test_configured_dataset_builds_only_the_frozen_development_corpus(
    tmp_path: Path,
) -> None:
    configured = os.environ.get("FORMULABENCH_DATASET_DIR")
    if not configured or not Path(configured).is_dir():
        pytest.skip("FORMULABENCH_DATASET_DIR is not configured")
    dataset_dir = Path(configured)
    split_path = Path(__file__).resolve().parents[1] / "experiments" / "public_split.json"
    split = json.loads(split_path.read_text(encoding="utf-8"))

    manifest = build_corpus(
        dataset_dir=dataset_dir,
        split_manifest=split_path,
        out_dir=tmp_path / "generated",
    )
    corpus_bytes = (tmp_path / "generated" / CORPUS_NAME).read_bytes()
    rows = [json.loads(line) for line in corpus_bytes.splitlines()]

    assert manifest["verified"] is True
    assert manifest["split"] == "development"
    assert manifest["example_count"] == 74
    assert manifest["partition_counts"] == {"train": 59, "validation": 15}
    assert manifest["excluded_development_count"] == 6
    assert {item["task_id"] for item in manifest["exclusions"]} == {
        "24-23",
        "91-34",
        "209-30",
        "290-27",
        "455-35",
        "48921",
    }
    assert all(row["task_id"] in set(split["development_ids"]) for row in rows)
    assert all(row["split"] == "development" for row in rows)
    assert hashlib.sha256(corpus_bytes).hexdigest() == manifest["corpus_sha256"]
    unmerge_row = next(row for row in rows if row["task_id"] == "38703")
    answer = json.loads(unmerge_row["messages"][-1]["content"])
    assert answer["unmerge_ranges"] == [
        {"range": "H2:O10", "sheet": "Sheet1"},
        {"range": "H11:O20", "sheet": "Sheet1"},
        {"range": "H21:P24", "sheet": "Sheet1"},
    ]
