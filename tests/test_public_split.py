from __future__ import annotations

import hashlib
import json
import os
from collections import Counter
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest

from experiments.make_public_split import (
    SPLIT_SIZES,
    SplitTask,
    assign_splits,
    build_public_split,
    canonical_json,
    initial_workbook_target_size_proxy,
    initial_workbooks_aggregate_sha256,
    largest_remainder_allocation,
    split_payload,
    target_size_proxy_bucket,
)
from formulabench.artifacts import confined_path, load_dataset_manifest, sha256_file


@pytest.mark.parametrize(
    ("size", "expected"),
    [
        (1, "1"),
        (2, "2-10"),
        (10, "2-10"),
        (11, "11-100"),
        (100, "11-100"),
        (101, "101-500"),
        (500, "101-500"),
        (501, "501-1000"),
        (1_000, "501-1000"),
        (1_001, "1001+"),
    ],
)
def test_target_size_proxy_bucket_boundaries(size: int, expected: str) -> None:
    assert target_size_proxy_bucket(size) == expected


def test_initial_workbook_target_size_proxy_counts_unique_and_dynamic_cells() -> None:
    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"
    worksheet["A5"] = "sets the used extent"
    try:
        overlapping = {
            "answer_position": "Data!A1:B2,Data!B2:C3",
            "answer_sheet": "Data",
        }
        dynamic = {"answer_position": "Data!D:E", "answer_sheet": "Data"}

        assert initial_workbook_target_size_proxy(overlapping, workbook) == 7
        assert initial_workbook_target_size_proxy(dynamic, workbook) == 10
    finally:
        workbook.close()


def _synthetic_tasks() -> list[SplitTask]:
    specifications = [
        ("Cell-Level Manipulation", 1, 36),
        ("Cell-Level Manipulation", 5, 116),
        ("Cell-Level Manipulation", 50, 104),
        ("Sheet-Level Manipulation", 250, 50),
        ("Sheet-Level Manipulation", 750, 24),
        ("Sheet-Level Manipulation", 2_000, 70),
    ]
    tasks: list[SplitTask] = []
    for instruction_type, target_size, count in specifications:
        for _ in range(count):
            index = len(tasks)
            tasks.append(
                SplitTask(
                    id=f"task-{index:03d}",
                    instruction_type=instruction_type,
                    target_size_proxy=target_size,
                    manifest_index=index,
                )
            )
    return tasks


def test_split_is_deterministic_disjoint_exhaustive_and_proportional() -> None:
    tasks = _synthetic_tasks()

    first = assign_splits(tasks)
    second = assign_splits(tasks)
    assert first == second

    ids_by_split = {
        name: {task.id for task in tasks if first[task.id] == name} for name, _ in SPLIT_SIZES
    }
    assert {name: len(ids) for name, ids in ids_by_split.items()} == {
        "development": 80,
        "held_out": 80,
        "final_only": 240,
    }
    assert ids_by_split["development"].isdisjoint(ids_by_split["held_out"])
    assert ids_by_split["development"].isdisjoint(ids_by_split["final_only"])
    assert ids_by_split["held_out"].isdisjoint(ids_by_split["final_only"])
    assert set().union(*ids_by_split.values()) == {task.id for task in tasks}

    stratum_totals = Counter(task.stratum for task in tasks)
    development_counts = Counter(task.stratum for task in tasks if first[task.id] == "development")
    for stratum, total in stratum_totals.items():
        assert abs(development_counts[stratum] - total * 80 / 400) < 1

    remaining = {
        stratum: total - development_counts[stratum] for stratum, total in stratum_totals.items()
    }
    held_out_counts = Counter(task.stratum for task in tasks if first[task.id] == "held_out")
    for stratum, capacity in remaining.items():
        assert abs(held_out_counts[stratum] - capacity * 80 / 320) < 1


def test_payload_ids_remain_in_manifest_order() -> None:
    tasks = _synthetic_tasks()
    assignments = assign_splits(tasks)

    initial_hashes = [(task.id, hashlib.sha256(task.id.encode()).hexdigest()) for task in tasks]
    payload = split_payload(
        tasks,
        manifest_sha256="a" * 64,
        initial_workbook_hashes=initial_hashes,
    )

    for split_name, _ in SPLIT_SIZES:
        expected = [task.id for task in tasks if assignments[task.id] == split_name]
        assert payload[f"{split_name}_ids"] == expected
    assert payload["dataset_manifest_sha256"] == "a" * 64
    assert payload["initial_workbook_binding"][
        "aggregate_sha256"
    ] == initial_workbooks_aggregate_sha256(initial_hashes)
    assert payload["initial_workbook_binding"]["record_count"] == len(tasks)
    assert payload["method"]["strata"] == [
        "instruction_type",
        "initial-workbook-resolved target-size proxy bucket",
    ]
    assert json.loads(canonical_json(payload)) == payload


def test_initial_workbook_binding_uses_documented_canonical_bytes() -> None:
    records = [
        ("task-a", hashlib.sha256(b"workbook-a").hexdigest()),
        ("task-b", hashlib.sha256(b"workbook-b").hexdigest()),
    ]
    canonical_records = b"".join(
        task_id.encode("utf-8") + b"\0" + digest.encode("ascii") + b"\n"
        for task_id, digest in records
    )

    assert (
        initial_workbooks_aggregate_sha256(records) == hashlib.sha256(canonical_records).hexdigest()
    )


def test_initial_workbook_binding_changes_with_ids_hashes_and_order() -> None:
    first_hash = hashlib.sha256(b"first").hexdigest()
    second_hash = hashlib.sha256(b"second").hexdigest()
    baseline = [("task-a", first_hash), ("task-b", second_hash)]

    baseline_digest = initial_workbooks_aggregate_sha256(baseline)
    assert initial_workbooks_aggregate_sha256(list(reversed(baseline))) != baseline_digest
    assert (
        initial_workbooks_aggregate_sha256([("task-renamed", first_hash), ("task-b", second_hash)])
        != baseline_digest
    )
    assert (
        initial_workbooks_aggregate_sha256([("task-a", second_hash), ("task-b", second_hash)])
        != baseline_digest
    )


def test_split_payload_requires_binding_records_in_task_order() -> None:
    tasks = _synthetic_tasks()
    ordered = [(task.id, "a" * 64) for task in tasks]

    with pytest.raises(ValueError, match="manifest order"):
        split_payload(
            tasks,
            manifest_sha256="b" * 64,
            initial_workbook_hashes=list(reversed(ordered)),
        )


def test_largest_remainder_uses_exact_capacity_and_size() -> None:
    capacities = {
        ("A", "1"): 1,
        ("A", "2-10"): 1,
        ("B", "1"): 1,
        ("B", "2-10"): 1,
    }

    first = largest_remainder_allocation(capacities, 2, split_name="development")
    second = largest_remainder_allocation(capacities, 2, split_name="development")

    assert first == second
    assert sum(first.values()) == 2
    assert all(0 <= first[stratum] <= capacity for stratum, capacity in capacities.items())


def test_configured_public_dataset_reproduces_checked_in_split_without_golden_access() -> None:
    configured = os.environ.get("FORMULABENCH_DATASET_DIR")
    if not configured or not Path(configured).is_dir():
        pytest.skip("FORMULABENCH_DATASET_DIR is not configured")
    dataset_dir = Path(configured)
    tasks = load_dataset_manifest(dataset_dir)
    initial_path_sequence = [task.init_xlsx.resolve(strict=True) for task in tasks]
    initial_paths = set(initial_path_sequence)
    manifest_path = confined_path(dataset_dir, "dataset.json", reject_symlinks=True)
    opened: list[Path] = []
    hashed: list[Path] = []
    original_load_workbook = openpyxl.load_workbook
    original_sha256_file = sha256_file

    def guarded_load_workbook(filename: object, *args: object, **kwargs: object) -> object:
        path = Path(filename).resolve(strict=True)  # type: ignore[arg-type]
        assert path in initial_paths, f"splitter tried to open a non-initial workbook: {path}"
        opened.append(path)
        return original_load_workbook(filename, *args, **kwargs)

    def guarded_sha256_file(filename: object) -> str:
        path = Path(filename).resolve(strict=True)  # type: ignore[arg-type]
        assert path in initial_paths or path == manifest_path
        hashed.append(path)
        return original_sha256_file(path)

    with (
        patch(
            "experiments.make_public_split.openpyxl.load_workbook",
            side_effect=guarded_load_workbook,
        ),
        patch(
            "experiments.make_public_split.sha256_file",
            side_effect=guarded_sha256_file,
        ),
    ):
        actual = build_public_split(dataset_dir)

    committed_path = Path(__file__).resolve().parents[1] / "experiments" / "public_split.json"
    committed_text = committed_path.read_text(encoding="utf-8")
    assert actual == json.loads(committed_text)
    assert canonical_json(actual) == committed_text
    assert len(opened) == 400
    assert opened == initial_path_sequence
    assert len(hashed) == 401
    assert Counter(hashed) == Counter([*initial_path_sequence, manifest_path])
    assert actual["initial_workbook_binding"]["record_count"] == 400
    expected_records = [(task.id, original_sha256_file(task.init_xlsx)) for task in tasks]
    assert actual["initial_workbook_binding"][
        "aggregate_sha256"
    ] == initial_workbooks_aggregate_sha256(expected_records)
