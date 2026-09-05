"""Audit public SpreadsheetBench structure without publishing task-level answers.

This is an analysis-only tool. It deliberately opens the public golden
workbooks, so its output is descriptive evidence rather than a held-out score.
Only aggregate counts, rates, and cryptographic bindings leave this module.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import re
from collections import Counter
from collections.abc import Mapping
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.datetime import CALENDAR_MAC_1904, CALENDAR_WINDOWS_1900

from experiments.make_public_split import initial_workbooks_aggregate_sha256
from formulabench.artifacts import (
    DatasetTask,
    confined_path,
    load_dataset_manifest,
    sha256_file,
)
from formulabench.contract import build_target_contract
from sb import answer_cells, values_equal

SCHEMA_VERSION = 1
ANALYSIS_VERSION = "FormulaBench/public-benchmark-audit/v1"
SIZE_BUCKETS = ("1-100", "101-500", "501-1000", "1001-10000", "10001+")
LEVELS = ("cell_level", "sheet_level")
LEVEL_LABELS = {
    "cell_level": "Cell-Level Manipulation",
    "sheet_level": "Sheet-Level Manipulation",
}

_MONTH_NAME = (
    r"(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|"
    r"Jul(?:y)?|Aug(?:ust)?|Sep(?:t(?:ember)?)?|Oct(?:ober)?|"
    r"Nov(?:ember)?|Dec(?:ember)?)"
)
_NUMERIC_DATE_TEXT_RE = re.compile(
    r"^(?:[0-9]{4}[-/.][0-9]{1,2}[-/.][0-9]{1,2}|"
    r"[0-9]{1,2}[-/.][0-9]{1,2}[-/.][0-9]{2,4})"
    r"(?:[ T][0-9]{1,2}:[0-9]{2}(?::[0-9]{2}(?:\.[0-9]+)?)?)?$"
)
_MONTH_NAME_DATE_TEXT_RE = re.compile(
    rf"^(?:[0-9]{{1,2}}[- ]{_MONTH_NAME}[- ,]+[0-9]{{2,4}}|"
    rf"{_MONTH_NAME}[- ][0-9]{{1,2}}(?:,)?[- ][0-9]{{2,4}}|"
    rf"[0-9]{{4}}[- ]{_MONTH_NAME}[- ][0-9]{{1,2}})"
    rf"(?:[ T][0-9]{{1,2}}:[0-9]{{2}}(?::[0-9]{{2}})?)?$",
    re.IGNORECASE,
)
_YEAR_MONTH_TEXT_RE = re.compile(r"^(?:19|20|21)[0-9]{2}/(?:0?[1-9]|1[0-2])$")
_NON_ISO_SLASH_DATE_TEXT_RE = re.compile(
    r"^[0-9]{1,2}/[0-9]{1,2}[/.][0-9]{2,4}"
    r"(?:[ T][0-9]{1,2}:[0-9]{2}(?::[0-9]{2}(?:\.[0-9]+)?)?)?$"
)


@dataclass(slots=True)
class LevelMetrics:
    """Mutable task aggregates for one instruction level."""

    tasks: int = 0
    finite_tasks: int = 0
    dynamic_tasks: int = 0
    finite_cells: int = 0
    dynamic_cells: int = 0
    scored_cells: int = 0
    size_buckets: Counter[str] = field(default_factory=Counter)
    blank_cells: int = 0
    tasks_with_blank_cells: int = 0
    initial_nonblank_target_cells: int = 0
    tasks_with_initial_nonblank_target_cells: int = 0
    correct_nonblank_prefill_cells: int = 0
    tasks_with_correct_nonblank_prefill_and_remaining_changes: int = 0
    cells_requiring_clearing: int = 0
    tasks_requiring_clearing: int = 0
    untouched_correct_cells: int = 0
    untouched_passes: int = 0
    date_or_datetime_cells: int = 0
    tasks_with_date_or_datetime_cells: int = 0
    non_midnight_datetime_cells: int = 0
    tasks_with_non_midnight_datetime_cells: int = 0
    time_cells: int = 0
    tasks_with_time_cells: int = 0
    initial_typed_date_time_cells: int = 0
    tasks_with_initial_typed_date_time_cells: int = 0
    initial_date_like_text_cells: int = 0
    tasks_with_initial_date_like_text_cells: int = 0
    tasks_mixing_initial_typed_dates_and_date_like_text: int = 0
    tasks_mixing_typed_dates_and_non_iso_slash_text: int = 0


def _instruction_level(value: object) -> str:
    if value == LEVEL_LABELS["cell_level"]:
        return "cell_level"
    if value == LEVEL_LABELS["sheet_level"]:
        return "sheet_level"
    raise ValueError(f"unsupported instruction_type: {value!r}")


def _size_bucket(size: int) -> str:
    if size <= 100:
        return "1-100"
    if size <= 500:
        return "101-500"
    if size <= 1_000:
        return "501-1000"
    if size <= 10_000:
        return "1001-10000"
    return "10001+"


def _ratio(numerator: int, denominator: int) -> float | None:
    return round(numerator / denominator, 4) if denominator else None


def _is_blank(value: object) -> bool:
    return value is None or value == ""


def _is_date_or_datetime(value: object) -> bool:
    return isinstance(value, (dt.datetime, dt.date))


def _is_non_midnight_datetime(value: object) -> bool:
    return isinstance(value, dt.datetime) and value.time() != dt.time()


def _is_date_like_text(value: object) -> bool:
    if not isinstance(value, str):
        return False
    text = value.strip()
    return bool(
        _NUMERIC_DATE_TEXT_RE.fullmatch(text)
        or _MONTH_NAME_DATE_TEXT_RE.fullmatch(text)
        or _YEAR_MONTH_TEXT_RE.fullmatch(text)
    )


def _is_non_iso_slash_date_text(value: object) -> bool:
    return isinstance(value, str) and bool(_NON_ISO_SLASH_DATE_TEXT_RE.fullmatch(value.strip()))


def _answer_values(
    task: Mapping[str, object], workbook: openpyxl.Workbook
) -> dict[tuple[str, str], object]:
    """Read evaluator-selected cells using the shipped target resolver."""

    result: dict[tuple[str, str], object] = {}
    for requested_sheet, coordinate in answer_cells(task, workbook):
        worksheet = (
            workbook[requested_sheet]
            if requested_sheet and requested_sheet in workbook.sheetnames
            else workbook.active
        )
        result[(worksheet.title, coordinate)] = worksheet[coordinate].value
    return result


def _populated_values(workbook: openpyxl.Workbook) -> list[object]:
    """Return instantiated values without iterating empty worksheet rectangles."""

    values: list[object] = []
    for worksheet in workbook.worksheets:
        for cell in worksheet._cells.values():
            if not isinstance(cell, MergedCell) and cell.value is not None:
                values.append(cell.value)
    return values


def _date_system(workbook: openpyxl.Workbook) -> str:
    if workbook.epoch == CALENDAR_WINDOWS_1900:
        return "1900"
    if workbook.epoch == CALENDAR_MAC_1904:
        return "1904"
    return "other"


def _golden_workbook_path(dataset_dir: Path, task: DatasetTask) -> Path:
    """Resolve exactly one public golden workbook beneath its validated task folder."""

    root = dataset_dir.resolve(strict=True)
    folder = confined_path(root, task.spreadsheet_path, reject_symlinks=True)
    candidates = sorted(folder.glob("*_golden.xlsx"), key=lambda path: path.name)
    exact = folder / "golden.xlsx"
    if exact.exists():
        candidates.append(exact)
    unique = sorted(set(candidates), key=lambda path: path.name)
    if len(unique) != 1:
        raise ValueError(f"task {task.id} must have exactly one golden workbook")
    relative = unique[0].relative_to(root).as_posix()
    resolved = confined_path(root, relative, reject_symlinks=True)
    if not resolved.is_file():
        raise ValueError(f"task {task.id} golden workbook is not a regular file")
    return resolved


def _binding_digest(manifest: str, initial: str, golden: str) -> str:
    payload = (
        f"{ANALYSIS_VERSION}\nmanifest\0{manifest}\ninitial\0{initial}\ngolden\0{golden}\n"
    ).encode("ascii")
    return hashlib.sha256(payload).hexdigest()


def _level_target_payload(level: str, metrics: LevelMetrics) -> dict[str, Any]:
    return {
        "dynamic_cells": metrics.dynamic_cells,
        "dynamic_tasks": metrics.dynamic_tasks,
        "finite_cells": metrics.finite_cells,
        "finite_size_buckets": {bucket: metrics.size_buckets[bucket] for bucket in SIZE_BUCKETS},
        "finite_tasks": metrics.finite_tasks,
        "instruction_type": LEVEL_LABELS[level],
        "scored_cells_including_dynamic": metrics.scored_cells,
        "tasks": metrics.tasks,
    }


def _target_state_payload(metrics: LevelMetrics) -> dict[str, Any]:
    return {
        "correct_nonblank_prefill_cells": metrics.correct_nonblank_prefill_cells,
        "initial_nonblank_target_cells": metrics.initial_nonblank_target_cells,
        "nonblank_cells_requiring_clearing": metrics.cells_requiring_clearing,
        "tasks_requiring_nonblank_cell_clearing": metrics.tasks_requiring_clearing,
        "tasks_with_correct_nonblank_prefill_and_remaining_changes": (
            metrics.tasks_with_correct_nonblank_prefill_and_remaining_changes
        ),
        "tasks_with_initial_nonblank_target_cells": (
            metrics.tasks_with_initial_nonblank_target_cells
        ),
    }


def _golden_target_payload(metrics: LevelMetrics) -> dict[str, Any]:
    return {
        "blank_cell_fraction": _ratio(metrics.blank_cells, metrics.scored_cells),
        "blank_cells": metrics.blank_cells,
        "cells": metrics.scored_cells,
        "tasks": metrics.tasks,
        "tasks_with_blank_cells": metrics.tasks_with_blank_cells,
    }


def _baseline_payload(metrics: LevelMetrics) -> dict[str, Any]:
    return {
        "cell_accuracy": _ratio(metrics.untouched_correct_cells, metrics.scored_cells),
        "cells": metrics.scored_cells,
        "correct_cells": metrics.untouched_correct_cells,
        "pass_rate": _ratio(metrics.untouched_passes, metrics.tasks),
        "passed_tasks": metrics.untouched_passes,
        "tasks": metrics.tasks,
    }


def _golden_date_payload(metrics: LevelMetrics) -> dict[str, Any]:
    return {
        "date_or_datetime_cells": metrics.date_or_datetime_cells,
        "non_midnight_datetime_cells": metrics.non_midnight_datetime_cells,
        "tasks_with_date_or_datetime_cells": metrics.tasks_with_date_or_datetime_cells,
        "tasks_with_non_midnight_datetime_cells": (metrics.tasks_with_non_midnight_datetime_cells),
        "tasks_with_time_cells": metrics.tasks_with_time_cells,
        "time_cells": metrics.time_cells,
    }


def _initial_date_payload(metrics: LevelMetrics) -> dict[str, Any]:
    return {
        "date_like_text_cells": metrics.initial_date_like_text_cells,
        "tasks_mixing_typed_dates_and_date_like_text": (
            metrics.tasks_mixing_initial_typed_dates_and_date_like_text
        ),
        "tasks_mixing_typed_dates_and_non_iso_slash_date_text": (
            metrics.tasks_mixing_typed_dates_and_non_iso_slash_text
        ),
        "tasks_with_date_like_text_cells": metrics.tasks_with_initial_date_like_text_cells,
        "tasks_with_typed_date_or_time_cells": (metrics.tasks_with_initial_typed_date_time_cells),
        "typed_date_or_time_cells": metrics.initial_typed_date_time_cells,
    }


def _sum_levels(metrics_by_level: Mapping[str, LevelMetrics]) -> LevelMetrics:
    total = LevelMetrics()
    counter_fields = {"size_buckets"}
    for metrics in metrics_by_level.values():
        for field_name in LevelMetrics.__dataclass_fields__:
            if field_name in counter_fields:
                total.size_buckets.update(metrics.size_buckets)
            else:
                value = getattr(total, field_name) + getattr(metrics, field_name)
                setattr(total, field_name, value)
    return total


def analyse(dataset_dir: Path) -> dict[str, Any]:
    """Return an aggregate-only audit of one public benchmark directory."""

    root = dataset_dir.resolve(strict=True)
    tasks = load_dataset_manifest(root)
    metrics_by_level = {level: LevelMetrics() for level in LEVELS}

    initial_hashes: list[tuple[str, str]] = []
    golden_hashes: list[tuple[str, str]] = []
    initial_date_systems: Counter[str] = Counter()
    golden_date_systems: Counter[str] = Counter()
    maximum_finite_cells = 0

    for task in tasks:
        metadata = task.as_dict()
        level = _instruction_level(metadata.get("instruction_type"))
        metrics = metrics_by_level[level]
        metrics.tasks += 1
        golden_path = _golden_workbook_path(root, task)
        initial_hashes.append((task.id, sha256_file(task.init_xlsx)))
        golden_hashes.append((task.id, sha256_file(golden_path)))

        initial_workbook = openpyxl.load_workbook(
            task.init_xlsx,
            data_only=True,
            read_only=False,
            keep_links=False,
        )
        golden_workbook = openpyxl.load_workbook(
            golden_path,
            data_only=True,
            read_only=False,
            keep_links=False,
        )
        try:
            initial_date_systems[_date_system(initial_workbook)] += 1
            golden_date_systems[_date_system(golden_workbook)] += 1
            contract = build_target_contract(metadata, initial_workbook)
            dynamic = any(target.dynamic for target in contract.ranges)
            golden_values = _answer_values(metadata, golden_workbook)
            initial_values = _answer_values(metadata, initial_workbook)
            scored_cells = len(golden_values)
            if scored_cells < 1:
                raise ValueError(f"task {task.id} resolves to no golden target cells")
            metrics.scored_cells += scored_cells
            if dynamic:
                metrics.dynamic_tasks += 1
                metrics.dynamic_cells += scored_cells
            else:
                metrics.finite_tasks += 1
                metrics.finite_cells += scored_cells
                metrics.size_buckets[_size_bucket(scored_cells)] += 1
                maximum_finite_cells = max(maximum_finite_cells, scored_cells)

            blank_cells = sum(_is_blank(value) for value in golden_values.values())
            metrics.blank_cells += blank_cells
            metrics.tasks_with_blank_cells += blank_cells > 0

            initial_at_target = {key: initial_values.get(key) for key in golden_values}
            nonblank_initial_keys = {
                key for key, value in initial_at_target.items() if not _is_blank(value)
            }
            correct_keys = {
                key
                for key, golden_value in golden_values.items()
                if values_equal(golden_value, initial_at_target[key])
            }
            correct_nonblank_keys = nonblank_initial_keys.intersection(correct_keys)
            clearing_keys = {key for key in nonblank_initial_keys if _is_blank(golden_values[key])}
            metrics.initial_nonblank_target_cells += len(nonblank_initial_keys)
            metrics.tasks_with_initial_nonblank_target_cells += bool(nonblank_initial_keys)
            metrics.correct_nonblank_prefill_cells += len(correct_nonblank_keys)
            metrics.tasks_with_correct_nonblank_prefill_and_remaining_changes += (
                bool(correct_nonblank_keys) and len(correct_keys) < scored_cells
            )
            metrics.cells_requiring_clearing += len(clearing_keys)
            metrics.tasks_requiring_clearing += bool(clearing_keys)
            metrics.untouched_correct_cells += len(correct_keys)
            metrics.untouched_passes += len(correct_keys) == scored_cells

            golden_date_values = [
                value for value in golden_values.values() if _is_date_or_datetime(value)
            ]
            golden_non_midnight_values = [
                value for value in golden_date_values if _is_non_midnight_datetime(value)
            ]
            golden_time_values = [
                value for value in golden_values.values() if isinstance(value, dt.time)
            ]
            metrics.date_or_datetime_cells += len(golden_date_values)
            metrics.tasks_with_date_or_datetime_cells += bool(golden_date_values)
            metrics.non_midnight_datetime_cells += len(golden_non_midnight_values)
            metrics.tasks_with_non_midnight_datetime_cells += bool(golden_non_midnight_values)
            metrics.time_cells += len(golden_time_values)
            metrics.tasks_with_time_cells += bool(golden_time_values)

            initial_populated = _populated_values(initial_workbook)
            typed_date_time_values = [
                value
                for value in initial_populated
                if _is_date_or_datetime(value) or isinstance(value, dt.time)
            ]
            date_like_text_values = [
                value for value in initial_populated if _is_date_like_text(value)
            ]
            non_iso_slash_values = [
                value for value in initial_populated if _is_non_iso_slash_date_text(value)
            ]
            metrics.initial_typed_date_time_cells += len(typed_date_time_values)
            metrics.tasks_with_initial_typed_date_time_cells += bool(typed_date_time_values)
            metrics.initial_date_like_text_cells += len(date_like_text_values)
            metrics.tasks_with_initial_date_like_text_cells += bool(date_like_text_values)
            metrics.tasks_mixing_initial_typed_dates_and_date_like_text += bool(
                typed_date_time_values and date_like_text_values
            )
            metrics.tasks_mixing_typed_dates_and_non_iso_slash_text += bool(
                typed_date_time_values and non_iso_slash_values
            )
        finally:
            initial_workbook.close()
            golden_workbook.close()

    total = _sum_levels(metrics_by_level)
    manifest_path = confined_path(root, "dataset.json", reject_symlinks=True)
    manifest_digest = sha256_file(manifest_path)
    initial_digest = initial_workbooks_aggregate_sha256(initial_hashes)
    golden_digest = initial_workbooks_aggregate_sha256(golden_hashes)

    by_level_targets = {
        level: _level_target_payload(level, metrics_by_level[level]) for level in LEVELS
    }
    by_level_golden = {level: _golden_target_payload(metrics_by_level[level]) for level in LEVELS}
    by_level_initial = {level: _target_state_payload(metrics_by_level[level]) for level in LEVELS}
    by_level_baseline = {level: _baseline_payload(metrics_by_level[level]) for level in LEVELS}
    by_level_golden_dates = {
        level: _golden_date_payload(metrics_by_level[level]) for level in LEVELS
    }
    by_level_initial_dates = {
        level: _initial_date_payload(metrics_by_level[level]) for level in LEVELS
    }

    return {
        "date_and_time": {
            "golden_target_cells": {
                **_golden_date_payload(total),
                "by_instruction_level": by_level_golden_dates,
            },
            "initial_workbooks": {
                **_initial_date_payload(total),
                "by_instruction_level": by_level_initial_dates,
            },
            "workbook_date_systems": {
                "golden": dict(sorted(golden_date_systems.items())),
                "initial": dict(sorted(initial_date_systems.items())),
            },
        },
        "golden_targets": {
            **_golden_target_payload(total),
            "by_instruction_level": by_level_golden,
        },
        "hashes": {
            "algorithm": "SHA-256",
            "dataset_binding_sha256": _binding_digest(
                manifest_digest, initial_digest, golden_digest
            ),
            "dataset_manifest_sha256": manifest_digest,
            "golden_workbooks": {
                "aggregate_sha256": golden_digest,
                "record_count": len(golden_hashes),
            },
            "initial_workbooks": {
                "aggregate_sha256": initial_digest,
                "record_count": len(initial_hashes),
            },
        },
        "initial_target_state": {
            **_target_state_payload(total),
            "by_instruction_level": by_level_initial,
        },
        "limitations": [
            (
                "This audit opens every public golden workbook. It is descriptive public "
                "validation and must not be presented as an untouched held-out result."
            ),
            (
                "No LibreOffice recalculation is performed. The untouched-input baseline "
                "compares the stored data-only values visible to openpyxl."
            ),
            (
                "Date-like text is a documented syntax heuristic; it does not infer locale "
                "or claim that every matching identifier is semantically a date."
            ),
            (
                "The audit measures evaluator-visible values, not formula quality, style, "
                "formatting, macros, charts, or workbook usability."
            ),
        ],
        "methodology": {
            "aggregation": (
                "Only dataset-wide and instruction-level totals are emitted. No task ID, "
                "sheet name, coordinate, formula, answer value, or per-task result is emitted."
            ),
            "blank_definition": "A raw golden target value is null or an empty string.",
            "date_like_text_definition": (
                "Trimmed three-part numeric dates, day/month-name/year dates, and YYYY/MM "
                "periods. Timezone-bearing timestamps and free-form text are excluded."
            ),
            "dynamic_range_resolution": (
                "Whole-column answer ranges use the golden workbook's max_row, matching the "
                "shipped evaluator's answer_cells behaviour."
            ),
            "hash_binding": (
                "Initial and golden aggregates use manifest-order records of UTF-8 task ID, "
                "NUL, lowercase workbook SHA-256, and LF. The dataset binding hashes the "
                "analysis version plus labelled manifest, initial, and golden digests."
            ),
            "inputs": ["dataset.json", "initial workbooks", "golden workbooks"],
            "preexisting_target_definition": (
                "Initial values are looked up by the evaluator-resolved golden sheet and "
                "coordinate. A correct nonblank prefill is a nonblank initial value accepted "
                "by sb.values_equal; 'and remaining changes' also requires another mismatch."
            ),
            "required_clearing_definition": (
                "A nonblank initial target value corresponds to a blank golden target value."
            ),
            "target_resolution": (
                "sb.answer_cells selects evaluator-visible targets; overlapping coordinates "
                "are counted once by the same sheet-and-coordinate mapping used by scoring."
            ),
            "untouched_input_comparison": (
                "Each golden target value is compared with the untouched initial value at "
                "the same resolved key using the shipped sb.values_equal function."
            ),
        },
        "schema_version": SCHEMA_VERSION,
        "scope": {
            "instruction_levels": {
                level: {
                    "instruction_type": LEVEL_LABELS[level],
                    "tasks": metrics_by_level[level].tasks,
                }
                for level in LEVELS
            },
            "tasks": total.tasks,
        },
        "targets": {
            "by_instruction_level": by_level_targets,
            "dynamic_cells": total.dynamic_cells,
            "dynamic_tasks": total.dynamic_tasks,
            "finite_cells": total.finite_cells,
            "finite_size_buckets": {bucket: total.size_buckets[bucket] for bucket in SIZE_BUCKETS},
            "finite_tasks": total.finite_tasks,
            "largest_finite_task_cells": maximum_finite_cells,
            "scored_cells_including_dynamic": total.scored_cells,
        },
        "untouched_input_baseline": {
            **_baseline_payload(total),
            "by_instruction_level": by_level_baseline,
        },
    }


def canonical_json(payload: Mapping[str, Any]) -> str:
    return json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n"


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dataset-dir", required=True, type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    text = canonical_json(analyse(args.dataset_dir))
    if args.output is None:
        print(text, end="")
    else:
        args.output.write_text(text, encoding="utf-8")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
