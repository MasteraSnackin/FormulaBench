"""Create the fixed public development, held-out, and final-only task split.

The splitter deliberately uses only ``dataset.json`` and the initial workbook
selected by :func:`formulabench.artifacts.load_dataset_manifest`.  It neither
searches for nor opens golden workbooks.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter, defaultdict
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from itertools import pairwise
from pathlib import Path
from typing import Any, TypeAlias

import openpyxl
from openpyxl.utils.cell import range_boundaries

from formulabench.artifacts import confined_path, load_dataset_manifest, sha256_file
from formulabench.contract import build_target_contract

SCHEMA_VERSION = 2
PUBLIC_TASK_COUNT = 400
SPLIT_SEED = "FormulaBench/public-split/v1"
SPLIT_SIZES = (("development", 80), ("held_out", 80), ("final_only", 240))
TARGET_SIZE_PROXY_BUCKETS = ("1", "2-10", "11-100", "101-500", "501-1000", "1001+")

Stratum: TypeAlias = tuple[str, str]


@dataclass(frozen=True, slots=True)
class SplitTask:
    """The public metadata needed to place one task in a split."""

    id: str
    instruction_type: str
    target_size_proxy: int
    manifest_index: int

    @property
    def target_size_proxy_bucket(self) -> str:
        return target_size_proxy_bucket(self.target_size_proxy)

    @property
    def stratum(self) -> Stratum:
        return self.instruction_type, self.target_size_proxy_bucket


def target_size_proxy_bucket(size: int) -> str:
    """Bucket a positive initial-workbook-resolved target-size proxy."""

    if isinstance(size, bool) or not isinstance(size, int) or size < 1:
        raise ValueError("target size must be a positive integer")
    if size == 1:
        return "1"
    if size <= 10:
        return "2-10"
    if size <= 100:
        return "11-100"
    if size <= 500:
        return "101-500"
    if size <= 1_000:
        return "501-1000"
    return "1001+"


def _union_interval_length(intervals: Iterable[tuple[int, int]]) -> int:
    ordered = sorted(intervals)
    if not ordered:
        return 0
    total = 0
    current_start, current_end = ordered[0]
    for start, end in ordered[1:]:
        if start > current_end + 1:
            total += current_end - current_start + 1
            current_start, current_end = start, end
        else:
            current_end = max(current_end, end)
    return total + current_end - current_start + 1


def _rectangle_union_area(rectangles: Sequence[tuple[int, int, int, int]]) -> int:
    """Count covered cells without expanding worksheet-sized rectangles."""

    if not rectangles:
        return 0
    row_edges = sorted({edge for _, start, _, end in rectangles for edge in (start, end + 1)})
    area = 0
    for start_row, end_row_exclusive in pairwise(row_edges):
        intervals = [
            (min_column, max_column)
            for min_column, rectangle_start, max_column, rectangle_end in rectangles
            if rectangle_start <= start_row <= rectangle_end
        ]
        area += _union_interval_length(intervals) * (end_row_exclusive - start_row)
    return area


def initial_workbook_target_size_proxy(task: Mapping[str, Any], workbook: Any) -> int:
    """Count targeted cells using only extents in the initial workbook.

    Whole-column targets are capped at the initial worksheet's used extent, so
    this is a stratification proxy rather than the evaluator's final target
    size. Finite ranges are counted exactly and overlapping cells only once.
    """

    contract = build_target_contract(task, workbook)
    rectangles_by_sheet: dict[str, list[tuple[int, int, int, int]]] = defaultdict(list)
    for target in contract.ranges:
        min_column, min_row, max_column, max_row = range_boundaries(target.cell_range)
        if min_column is None or max_column is None:
            raise ValueError(f"target has unresolved columns: {target.cell_range}")
        if min_row is None or max_row is None:
            worksheet = workbook[target.sheet] if target.sheet in workbook.sheetnames else None
            min_row = min_row or 1
            max_row = max_row or max(1, worksheet.max_row if worksheet is not None else 1)
        rectangles_by_sheet[target.sheet].append(
            (int(min_column), int(min_row), int(max_column), int(max_row))
        )

    size = sum(_rectangle_union_area(rectangles) for rectangles in rectangles_by_sheet.values())
    if size < 1:
        raise ValueError("resolved target contract contains no cells")
    return size


def initial_workbooks_aggregate_sha256(
    records: Sequence[tuple[str, str]],
) -> str:
    """Bind ordered task IDs to exact initial-workbook file bytes.

    Each record is ``UTF-8 task ID || NUL || lowercase ASCII workbook SHA-256
    || LF``. The returned digest is SHA-256 over the concatenated records in
    their supplied order.
    """

    aggregate = hashlib.sha256()
    for task_id, workbook_sha256 in records:
        if not isinstance(task_id, str) or not task_id or "\0" in task_id or "\n" in task_id:
            raise ValueError("binding task ID must be a non-empty single-line string")
        if (
            not isinstance(workbook_sha256, str)
            or len(workbook_sha256) != 64
            or any(character not in "0123456789abcdef" for character in workbook_sha256)
        ):
            raise ValueError("workbook SHA-256 must be a lowercase hexadecimal digest")
        aggregate.update(task_id.encode("utf-8"))
        aggregate.update(b"\0")
        aggregate.update(workbook_sha256.encode("ascii"))
        aggregate.update(b"\n")
    return aggregate.hexdigest()


def _hash_rank(*parts: str) -> bytes:
    payload = "\0".join((SPLIT_SEED, *parts)).encode("utf-8")
    return hashlib.sha256(payload).digest()


def _stratum_token(stratum: Stratum) -> str:
    return json.dumps(stratum, ensure_ascii=False, separators=(",", ":"))


def largest_remainder_allocation(
    capacities: Mapping[Stratum, int],
    size: int,
    *,
    split_name: str,
) -> dict[Stratum, int]:
    """Allocate ``size`` proportionally with Hamilton largest remainders."""

    if isinstance(size, bool) or not isinstance(size, int) or size < 0:
        raise ValueError("split size must be a non-negative integer")
    if any(
        isinstance(value, bool) or not isinstance(value, int) or value < 0
        for value in capacities.values()
    ):
        raise ValueError("stratum capacities must be non-negative integers")
    total = sum(capacities.values())
    if size > total:
        raise ValueError("split size exceeds remaining task capacity")
    if total == 0:
        return {stratum: 0 for stratum in capacities}

    allocation: dict[Stratum, int] = {}
    remainders: dict[Stratum, int] = {}
    for stratum, capacity in capacities.items():
        numerator = capacity * size
        allocation[stratum], remainders[stratum] = divmod(numerator, total)

    unassigned = size - sum(allocation.values())
    ranked = sorted(
        capacities,
        key=lambda stratum: (
            -remainders[stratum],
            _hash_rank("allocation", split_name, _stratum_token(stratum)),
            stratum,
        ),
    )
    for stratum in ranked:
        if unassigned == 0:
            break
        if allocation[stratum] < capacities[stratum]:
            allocation[stratum] += 1
            unassigned -= 1
    if unassigned:
        raise AssertionError("largest-remainder allocation did not reach the requested size")
    return allocation


def assign_splits(
    tasks: Sequence[SplitTask],
    *,
    split_sizes: Sequence[tuple[str, int]] = SPLIT_SIZES,
) -> dict[str, str]:
    """Assign unique task IDs sequentially in the declared split order."""

    task_ids = [task.id for task in tasks]
    if len(set(task_ids)) != len(task_ids):
        raise ValueError("task IDs must be unique")
    split_names = [name for name, _ in split_sizes]
    if len(set(split_names)) != len(split_names) or any(not name for name in split_names):
        raise ValueError("split names must be unique non-empty strings")
    if any(
        isinstance(size, bool) or not isinstance(size, int) or size < 0 for _, size in split_sizes
    ):
        raise ValueError("split sizes must be non-negative integers")
    if sum(size for _, size in split_sizes) != len(tasks):
        raise ValueError("split sizes must exhaust the tasks")

    remaining: dict[Stratum, list[SplitTask]] = defaultdict(list)
    for task in tasks:
        remaining[task.stratum].append(task)

    assignments: dict[str, str] = {}
    for split_name, split_size in split_sizes:
        capacities = {stratum: len(items) for stratum, items in remaining.items()}
        quotas = largest_remainder_allocation(capacities, split_size, split_name=split_name)
        for stratum in sorted(remaining):
            quota = quotas[stratum]
            ranked = sorted(
                remaining[stratum],
                key=lambda task: (
                    _hash_rank(
                        "selection",
                        split_name,
                        _stratum_token(stratum),
                        task.id,
                    ),
                    task.manifest_index,
                ),
            )
            selected = ranked[:quota]
            selected_ids = {task.id for task in selected}
            for task in selected:
                assignments[task.id] = split_name
            remaining[stratum] = [
                task for task in remaining[stratum] if task.id not in selected_ids
            ]

    if len(assignments) != len(tasks):
        raise AssertionError("split assignment did not exhaust the tasks")
    return assignments


def split_payload(
    tasks: Sequence[SplitTask],
    *,
    manifest_sha256: str,
    initial_workbook_hashes: Sequence[tuple[str, str]],
) -> dict[str, Any]:
    """Build the canonical, manifest-ordered public split payload."""

    task_ids = [task.id for task in tasks]
    binding_ids = [task_id for task_id, _ in initial_workbook_hashes]
    if binding_ids != task_ids:
        raise ValueError("initial-workbook hashes must match task IDs in manifest order")
    workbook_binding_sha256 = initial_workbooks_aggregate_sha256(initial_workbook_hashes)

    assignments = assign_splits(tasks)
    by_split = {
        split_name: [task.id for task in tasks if assignments[task.id] == split_name]
        for split_name, _ in SPLIT_SIZES
    }
    bucket_order = {bucket: index for index, bucket in enumerate(TARGET_SIZE_PROXY_BUCKETS)}
    strata = sorted(
        {task.stratum for task in tasks},
        key=lambda stratum: (stratum[0], bucket_order[stratum[1]]),
    )
    per_stratum_counts = []
    for instruction_type, bucket in strata:
        stratum_tasks = [
            task
            for task in tasks
            if task.instruction_type == instruction_type and task.target_size_proxy_bucket == bucket
        ]
        counts = Counter(assignments[task.id] for task in stratum_tasks)
        per_stratum_counts.append(
            {
                "counts": {
                    "development": counts["development"],
                    "final_only": counts["final_only"],
                    "held_out": counts["held_out"],
                    "total": len(stratum_tasks),
                },
                "instruction_type": instruction_type,
                "initial_workbook_target_size_proxy_bucket": bucket,
            }
        )

    return {
        "algorithm": {
            "allocation": (
                "Sequential Hamilton largest-remainder allocation over remaining stratum "
                "capacities in development, held_out, final_only order."
            ),
            "hash": "SHA-256",
            "hash_input": (
                "UTF-8 NUL-separated seed and operation fields. Allocation uses purpose, "
                "split, and stratum; selection also appends task ID."
            ),
            "name": "stratified-largest-remainder-sha256",
            "seed": SPLIT_SEED,
            "initial_workbook_target_size_proxy_buckets": list(TARGET_SIZE_PROXY_BUCKETS),
            "version": 1,
        },
        "counts": {
            "development": len(by_split["development"]),
            "final_only": len(by_split["final_only"]),
            "held_out": len(by_split["held_out"]),
            "total": len(tasks),
        },
        "dataset_manifest_sha256": manifest_sha256,
        "initial_workbook_binding": {
            "aggregate_sha256": workbook_binding_sha256,
            "construction": (
                "For each task in dataset manifest order: UTF-8 task ID, one NUL byte "
                "(0x00), the lowercase 64-character ASCII SHA-256 of the exact initial "
                "workbook file bytes, then one LF byte (0x0a). The aggregate is SHA-256 "
                "over the concatenation of all records."
            ),
            "hash": "SHA-256",
            "record_count": len(initial_workbook_hashes),
        },
        "development_ids": by_split["development"],
        "final_only_ids": by_split["final_only"],
        "held_out_ids": by_split["held_out"],
        "method": {
            "golden_files_opened": False,
            "inputs": ["dataset.json", "initial workbooks"],
            "output_order": "dataset manifest order",
            "strata": [
                "instruction_type",
                "initial-workbook-resolved target-size proxy bucket",
            ],
        },
        "per_stratum_counts": per_stratum_counts,
        "schema_version": SCHEMA_VERSION,
    }


def build_public_split(dataset_dir: Path) -> dict[str, Any]:
    """Load the official public release and create its reproducible split."""

    tasks = load_dataset_manifest(dataset_dir)
    if len(tasks) != PUBLIC_TASK_COUNT:
        raise ValueError(f"public split requires exactly {PUBLIC_TASK_COUNT} tasks")

    split_tasks: list[SplitTask] = []
    initial_workbook_hashes: list[tuple[str, str]] = []
    for index, task in enumerate(tasks):
        metadata = task.as_dict()
        instruction_type = metadata.get("instruction_type")
        if not isinstance(instruction_type, str) or not instruction_type.strip():
            raise ValueError(f"task {task.id} has no non-empty instruction_type")
        workbook = openpyxl.load_workbook(
            task.init_xlsx,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        try:
            target_size_proxy = initial_workbook_target_size_proxy(metadata, workbook)
        finally:
            workbook.close()
        initial_workbook_hashes.append((task.id, sha256_file(task.init_xlsx)))
        split_tasks.append(
            SplitTask(
                id=task.id,
                instruction_type=instruction_type.strip(),
                target_size_proxy=target_size_proxy,
                manifest_index=index,
            )
        )

    manifest_path = confined_path(dataset_dir, "dataset.json", reject_symlinks=True)
    return split_payload(
        split_tasks,
        manifest_sha256=sha256_file(manifest_path),
        initial_workbook_hashes=initial_workbook_hashes,
    )


def canonical_json(payload: Mapping[str, Any]) -> str:
    """Return the exact checked-in JSON representation."""

    return json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n"


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dataset-dir", required=True, type=Path)
    args = parser.parse_args()
    print(canonical_json(build_public_split(args.dataset_dir)), end="")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
