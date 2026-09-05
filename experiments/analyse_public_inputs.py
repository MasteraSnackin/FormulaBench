"""Measure public initial-workbook constraints without reading any golden workbook."""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from formulabench.artifacts import load_dataset_manifest
from formulabench.contract import build_target_contract


def _is_formula(value: object, data_type: str | None) -> bool:
    return (
        isinstance(value, (ArrayFormula, DataTableFormula))
        or data_type == "f"
        or (isinstance(value, str) and value.startswith("="))
    )


def _target_area(cell_range: str) -> int | None:
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    if None in (min_col, min_row, max_col, max_row):
        return None
    return (max_col - min_col + 1) * (max_row - min_row + 1)


def analyse(dataset_dir: Path) -> dict[str, Any]:
    tasks = load_dataset_manifest(dataset_dir)
    type_counts: Counter[str] = Counter()
    finite_target_counts: list[tuple[int, str]] = []
    dynamic_tasks: list[str] = []
    formula_cells = 0
    missing_formula_caches = 0
    formula_tasks: set[str] = set()
    oversized_tasks: set[str] = set()
    nonempty_outside_crop_tasks: set[str] = set()
    formula_outside_crop_tasks: set[str] = set()
    merged_non_anchor_targets: Counter[str] = Counter()

    for task in tasks:
        metadata = task.as_dict()
        type_counts[str(metadata.get("instruction_type"))] += 1
        formula_wb = openpyxl.load_workbook(
            task.init_xlsx,
            data_only=False,
            read_only=False,
            keep_links=False,
        )
        cached_wb = openpyxl.load_workbook(
            task.init_xlsx,
            data_only=True,
            read_only=False,
            keep_links=False,
        )
        try:
            contract = build_target_contract(metadata, formula_wb)
            task_area = 0
            task_dynamic = False
            for target in contract.ranges:
                area = _target_area(target.cell_range)
                if area is None:
                    task_dynamic = True
                    continue
                task_area += area

                if target.sheet not in formula_wb.sheetnames:
                    continue
                min_col, min_row, max_col, max_row = range_boundaries(target.cell_range)
                worksheet = formula_wb[target.sheet]
                for merged in worksheet.merged_cells.ranges:
                    overlap_min_col = max(min_col, merged.min_col)
                    overlap_min_row = max(min_row, merged.min_row)
                    overlap_max_col = min(max_col, merged.max_col)
                    overlap_max_row = min(max_row, merged.max_row)
                    if overlap_min_col > overlap_max_col or overlap_min_row > overlap_max_row:
                        continue
                    overlap = (overlap_max_col - overlap_min_col + 1) * (
                        overlap_max_row - overlap_min_row + 1
                    )
                    anchor_inside = (
                        overlap_min_col <= merged.min_col <= overlap_max_col
                        and overlap_min_row <= merged.min_row <= overlap_max_row
                    )
                    non_anchor_cells = overlap - int(anchor_inside)
                    if non_anchor_cells:
                        merged_non_anchor_targets[task.id] += non_anchor_cells

            finite_target_counts.append((task_area, task.id))
            if task_dynamic:
                dynamic_tasks.append(task.id)

            for worksheet in formula_wb.worksheets:
                if worksheet.max_row > 120 or worksheet.max_column > 30:
                    oversized_tasks.add(task.id)
                cached_ws = cached_wb[worksheet.title]
                coordinates = set(worksheet._cells).union(cached_ws._cells)
                for row, column in coordinates:
                    formula_cell = worksheet._cells.get((row, column))
                    cached_cell = cached_ws._cells.get((row, column))
                    if isinstance(formula_cell, MergedCell):
                        continue
                    value = None if formula_cell is None else formula_cell.value
                    data_type = None if formula_cell is None else formula_cell.data_type
                    cached_value = None if cached_cell is None else cached_cell.value
                    formula = _is_formula(value, data_type)
                    if formula:
                        formula_cells += 1
                        formula_tasks.add(task.id)
                        if cached_value is None:
                            missing_formula_caches += 1
                        if row > 120 or column > 30:
                            formula_outside_crop_tasks.add(task.id)
                    if (value is not None or cached_value is not None) and (
                        row > 120 or column > 30
                    ):
                        nonempty_outside_crop_tasks.add(task.id)
        finally:
            formula_wb.close()
            cached_wb.close()

    maximum, maximum_id = max(finite_target_counts)
    return {
        "dataset_tasks": len(tasks),
        "instruction_types": dict(sorted(type_counts.items())),
        "targets": {
            "dynamic_task_ids": sorted(dynamic_tasks),
            "finite_cells_total": sum(count for count, _ in finite_target_counts),
            "largest_finite_task": {"cells": maximum, "id": maximum_id},
            "tasks_over_100_cells": sum(count > 100 for count, _ in finite_target_counts),
            "tasks_over_500_cells": sum(count > 500 for count, _ in finite_target_counts),
            "tasks_over_1000_cells": sum(count > 1000 for count, _ in finite_target_counts),
        },
        "workbooks": {
            "formula_cells": formula_cells,
            "formula_cells_without_readable_cache": missing_formula_caches,
            "formula_tasks": len(formula_tasks),
            "oversized_for_baseline_crop_tasks": len(oversized_tasks),
            "tasks_with_formula_outside_baseline_crop": len(formula_outside_crop_tasks),
            "tasks_with_nonempty_cell_outside_baseline_crop": len(nonempty_outside_crop_tasks),
        },
        "merged_non_anchor_answer_cells": {
            "cells": sum(merged_non_anchor_targets.values()),
            "tasks": dict(sorted(merged_non_anchor_targets.items())),
        },
        "method": {
            "baseline_crop": "rows 1:120, columns A:AD",
            "golden_files_opened": False,
            "inputs": ["dataset.json", "initial workbooks"],
        },
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dataset-dir", required=True, type=Path)
    args = parser.parse_args()
    print(json.dumps(analyse(args.dataset_dir), indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
