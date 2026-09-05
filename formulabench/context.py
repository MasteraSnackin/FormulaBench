"""Deterministic, formula-aware workbook evidence for model prompts.

The serializer deliberately opens only ``task["init_xlsx"]``.  Dataset golden
files are neither discovered nor read.  Both workbook views are loaded from the
same input: one retains formulas and the other exposes any cached values stored
in the file.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import json
import math
import re
from collections import deque
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

try:  # openpyxl 3.1+
    from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
except ImportError:  # pragma: no cover - project dependency is openpyxl >=3.1
    ArrayFormula = DataTableFormula = ()  # type: ignore[assignment,misc]


DEFAULT_CONTEXT_CHARS = 16_000
SMALL_RANGE_CELL_LIMIT = 1_024
RANGE_SAMPLE_AXIS = 7
NEAR_TARGET_RADIUS = 5
MAX_EXCEL_ROW = 1_048_576
MAX_EXCEL_COLUMN = 16_384

_CELL = r"\$?[A-Za-z]{1,3}\$?\d+"
_COLUMN = r"\$?[A-Za-z]{1,3}"
_ROW = r"\$?\d+"
_RANGE_RE = re.compile(
    rf"(?<![A-Za-z0-9_])"
    rf"({_CELL}\s*:\s*(?:{_CELL}|{_COLUMN}|{_ROW})|{_CELL}|"
    rf"{_COLUMN}\s*:\s*{_COLUMN}|{_ROW}\s*:\s*{_ROW})"
    rf"(?![A-Za-z0-9_])(?!\s*!)",
    flags=re.IGNORECASE,
)

_ROLE_ORDER = {
    "answer": 0,
    "answer_sample": 1,
    "data": 2,
    "data_sample": 3,
    "near_answer_formula": 4,
    "data_boundary": 5,
    "instruction_ref": 6,
    "formula": 7,
    "value": 8,
}

_PRIORITY_WEIGHTS = (8, 4, 2, 3, 4, 2)


class ContextBudgetError(ValueError):
    """Raised when mandatory workbook records cannot fit the given budget."""


@dataclass(frozen=True, slots=True)
class PositionRef:
    """A parsed A1 range and how its worksheet name was resolved."""

    sheet: str
    a1_range: str
    requested_sheet: str | None
    resolution: str
    dynamic: bool


@dataclass(frozen=True, slots=True)
class ContextDocument:
    """Serialized JSONL plus deterministic accounting metadata."""

    jsonl: str
    char_budget: int
    used_chars: int
    record_count: int
    included_cells: int
    omitted_cells: int
    truncated: bool
    input_sha256: str

    def __str__(self) -> str:
        return self.jsonl


@dataclass(frozen=True, slots=True)
class _Bounds:
    min_col: int
    min_row: int
    max_col: int
    max_row: int

    @property
    def cell_count(self) -> int:
        return (self.max_col - self.min_col + 1) * (self.max_row - self.min_row + 1)

    def contains(self, row: int, col: int) -> bool:
        return self.min_row <= row <= self.max_row and self.min_col <= col <= self.max_col

    def distance(self, row: int, col: int) -> int:
        row_delta = max(self.min_row - row, 0, row - self.max_row)
        col_delta = max(self.min_col - col, 0, col - self.max_col)
        return row_delta + col_delta


@dataclass(frozen=True, slots=True)
class _CompactRange:
    """One all-or-nothing compact rendering of a finite declared range."""

    records: tuple[dict[str, Any], ...]
    coordinates: frozenset[tuple[str, int, int]]


def _normalise_a1_range(token: str) -> str:
    cleaned = re.sub(r"\s+", "", token).replace("$", "").upper()
    if ":" in cleaned:
        start, end = cleaned.split(":", 1)
        if end.isdigit() and re.fullmatch(r"[A-Z]{1,3}\d+", start):
            start_col = re.match(r"[A-Z]+", start)
            if start_col is not None:
                cleaned = f"{start}:{start_col.group(0)}{end}"
    range_boundaries(cleaned)
    return cleaned


def _normalise_sheet_token(token: str, sheet_names: Sequence[str]) -> str:
    for name in sorted(sheet_names, key=len, reverse=True):
        escaped = name.replace("'", "''")
        variants = (name, escaped, f"'{name}'", f"'{escaped}'", f"'{name}", f"'{escaped}")
        if any(
            token.endswith(variant)
            and (not token[: -len(variant)] or not token[: -len(variant)][-1].isalnum())
            for variant in variants
        ):
            return name
    cleaned = token.strip(" \t\r\n,;'\"")
    if cleaned.startswith("'") and cleaned.endswith("'") and len(cleaned) >= 2:
        cleaned = cleaned[1:-1]
    else:
        cleaned = cleaned.lstrip("'").rstrip("'")
    cleaned = cleaned.replace("''", "'").strip()
    if cleaned in sheet_names:
        return cleaned
    casefold_matches = [name for name in sheet_names if name.casefold() == cleaned.casefold()]
    return casefold_matches[0] if len(casefold_matches) == 1 else cleaned


def _default_sheet_name(default_sheet: str | None, sheet_names: Sequence[str]) -> str:
    if default_sheet:
        normalised = _normalise_sheet_token(str(default_sheet), sheet_names)
        if normalised in sheet_names:
            return normalised
    if not sheet_names:
        raise ValueError("workbook has no worksheets")
    return sheet_names[0]


def parse_position_references(
    value: str | Sequence[str] | None,
    *,
    sheet_names: Sequence[str],
    default_sheet: str | None = None,
    default_resolution: str = "default",
) -> tuple[PositionRef, ...]:
    """Parse forgiving SpreadsheetBench position text without losing sheet punctuation.

    Range matches are identified first.  The worksheet token immediately before
    each ``!`` is then reconciled against the workbook's actual titles.  This
    accepts the dataset's stray quotes and missing separators while preserving
    legitimate commas and apostrophes in worksheet names.
    """

    if value is None:
        return ()
    texts = [value] if isinstance(value, str) else [str(item) for item in value]
    names = tuple(str(name) for name in sheet_names)
    fallback = _default_sheet_name(default_sheet, names)
    parsed: list[PositionRef] = []

    for text in texts:
        previous_end = 0
        for match in _RANGE_RE.finditer(text):
            between = text[previous_end : match.start()]
            requested: str | None = None
            if "!" in between:
                requested = _normalise_sheet_token(between.rsplit("!", 1)[0], names)
            sheet = requested if requested in names else fallback
            resolution = "explicit" if requested in names else default_resolution
            if requested is not None and requested not in names:
                resolution = "missing_sheet_fallback"
            try:
                a1_range = _normalise_a1_range(match.group(1))
            except (TypeError, ValueError):
                previous_end = match.end()
                continue
            min_col, min_row, max_col, max_row = range_boundaries(a1_range)
            parsed.append(
                PositionRef(
                    sheet=sheet,
                    a1_range=a1_range,
                    requested_sheet=requested,
                    resolution=resolution,
                    dynamic=any(part is None for part in (min_col, min_row, max_col, max_row)),
                )
            )
            previous_end = match.end()
    return tuple(parsed)


def parse_data_positions(
    data_position: str | Sequence[str] | None,
    *,
    sheet_names: Sequence[str],
    default_sheet: str | None = None,
    default_resolution: str = "default",
) -> tuple[PositionRef, ...]:
    """Public, workbook-aware parser for ``task["data_position"]``."""

    return parse_position_references(
        data_position,
        sheet_names=sheet_names,
        default_sheet=default_sheet,
        default_resolution=default_resolution,
    )


def _instruction_mentions_sheet(instruction: str, sheet: str) -> bool:
    """Match one exact existing title in prose without substring collisions."""

    return bool(
        re.search(
            rf"(?<![A-Za-z0-9_]){re.escape(sheet)}(?![A-Za-z0-9_])",
            instruction,
            flags=re.IGNORECASE,
        )
    )


def _data_default_sheet(
    task: Mapping[str, Any],
    workbook: Any,
    answer_refs: Sequence[PositionRef],
) -> tuple[str, str]:
    """Resolve an unqualified data range conservatively from instruction evidence.

    A uniquely instruction-named existing sheet outside the answer targets is
    strong source-sheet evidence. Ambiguous or absent evidence deliberately
    keeps the established answer/active-sheet fallback.
    """

    fallback = answer_refs[0].sheet if answer_refs else workbook.active.title
    answer_sheets = {ref.sheet for ref in answer_refs}
    instruction = str(task.get("instruction", ""))
    candidates = [
        sheet
        for sheet in workbook.sheetnames
        if sheet not in answer_sheets and _instruction_mentions_sheet(instruction, sheet)
    ]
    if len(candidates) == 1:
        return candidates[0], "instruction_named_default"
    return fallback, "default"


def _fallback_answer_positions(
    task: Mapping[str, Any], sheet_names: Sequence[str]
) -> tuple[PositionRef, ...]:
    default = task.get("answer_sheet")
    return parse_position_references(
        task.get("answer_position"),
        sheet_names=sheet_names,
        default_sheet=str(default) if default is not None else None,
    )


def _contract_answer_positions(task: Mapping[str, Any], workbook: Any) -> tuple[PositionRef, ...]:
    """Use the output contract when available so context and evaluator targets agree."""

    try:
        from .contract import build_target_contract
    except ImportError:
        return _fallback_answer_positions(task, tuple(workbook.sheetnames))

    contract = build_target_contract(task, workbook)
    refs: list[PositionRef] = []
    for target in contract.ranges:
        cell_range = getattr(target, "cell_range", getattr(target, "a1_range", None))
        if cell_range is None:
            raise TypeError("target contract range has no cell_range")
        resolution = getattr(target, "resolution", "existing")
        if hasattr(resolution, "value"):
            resolution = resolution.value
        provenance = getattr(target, "provenance", None)
        if hasattr(provenance, "value"):
            provenance = provenance.value
        requested = getattr(target, "requested_sheet", None)
        refs.append(
            PositionRef(
                sheet=str(target.sheet),
                a1_range=_normalise_a1_range(str(cell_range)),
                requested_sheet=str(requested) if requested is not None else None,
                resolution=str(resolution if provenance is None else f"{provenance}:{resolution}"),
                dynamic=bool(getattr(target, "dynamic", False)),
            )
        )
    return tuple(refs)


def _bounds_for(ref: PositionRef, worksheet: Any) -> _Bounds:
    min_col, min_row, max_col, max_row = range_boundaries(ref.a1_range)
    min_col = min_col or 1
    min_row = min_row or 1
    max_col = max_col or max(worksheet.max_column, min_col)
    max_row = max_row or max(worksheet.max_row, min_row)
    return _Bounds(
        min_col=max(1, min(int(min_col), MAX_EXCEL_COLUMN)),
        min_row=max(1, min(int(min_row), MAX_EXCEL_ROW)),
        max_col=max(1, min(int(max_col), MAX_EXCEL_COLUMN)),
        max_row=max(1, min(int(max_row), MAX_EXCEL_ROW)),
    )


def _typed_value(value: Any) -> dict[str, Any]:
    if value is None:
        return {"type": "blank", "value": None}
    if isinstance(value, bool):
        return {"type": "boolean", "value": value}
    if isinstance(value, int):
        return {"type": "integer", "value": value}
    if isinstance(value, float):
        if math.isfinite(value):
            return {"type": "number", "value": value}
        return {"type": "number_special", "value": str(value)}
    if isinstance(value, Decimal):
        return {"type": "decimal", "value": str(value)}
    if isinstance(value, dt.datetime):
        return {"type": "datetime", "value": value.isoformat()}
    if isinstance(value, dt.date):
        return {"type": "date", "value": value.isoformat()}
    if isinstance(value, dt.time):
        return {"type": "time", "value": value.isoformat()}
    if isinstance(value, dt.timedelta):
        return {"type": "duration_seconds", "value": value.total_seconds()}
    if isinstance(value, bytes):
        return {"type": "bytes_hex", "value": value.hex()}
    if isinstance(value, str):
        return {"type": "string", "value": value}
    return {"type": type(value).__name__, "value": str(value)}


def _formula_view(value: Any, cell_data_type: str | None) -> dict[str, Any]:
    if ArrayFormula and isinstance(value, ArrayFormula):
        return {
            "kind": "formula",
            "formula_kind": "array",
            "ref": str(value.ref),
            "text": value.text,
        }
    if DataTableFormula and isinstance(value, DataTableFormula):
        metadata = {str(key): val for key, val in value}
        return {
            "kind": "formula",
            "formula_kind": "data_table",
            "metadata": metadata,
        }
    if cell_data_type == "f" or (isinstance(value, str) and value.startswith("=")):
        return {"kind": "formula", "formula_kind": "standard", "text": str(value)}
    return {"kind": "value", "value": _typed_value(value)}


def _is_formula(value: Any, cell_data_type: str | None) -> bool:
    return (
        bool(ArrayFormula and isinstance(value, ArrayFormula))
        or bool(DataTableFormula and isinstance(value, DataTableFormula))
        or cell_data_type == "f"
        or (isinstance(value, str) and value.startswith("="))
    )


def _axis_samples(start: int, end: int, count: int = RANGE_SAMPLE_AXIS) -> tuple[int, ...]:
    size = end - start + 1
    if size <= count:
        return tuple(range(start, end + 1))
    if count <= 1:
        return (start,)
    samples = {start + round(index * (size - 1) / (count - 1)) for index in range(count)}
    return tuple(sorted(samples))


def _sample_coordinates(bounds: _Bounds) -> Iterable[tuple[int, int]]:
    if bounds.cell_count <= SMALL_RANGE_CELL_LIMIT:
        for row in range(bounds.min_row, bounds.max_row + 1):
            for col in range(bounds.min_col, bounds.max_col + 1):
                yield row, col
        return
    for row in _axis_samples(bounds.min_row, bounds.max_row):
        for col in _axis_samples(bounds.min_col, bounds.max_col):
            yield row, col


def _boundary_coordinates(bounds: _Bounds) -> tuple[tuple[int, int], ...]:
    middle_row = (bounds.min_row + bounds.max_row) // 2
    middle_col = (bounds.min_col + bounds.max_col) // 2
    return tuple(
        sorted(
            {
                (bounds.min_row, bounds.min_col),
                (bounds.min_row, bounds.max_col),
                (bounds.max_row, bounds.min_col),
                (bounds.max_row, bounds.max_col),
                (middle_row, bounds.min_col),
                (middle_row, bounds.max_col),
                (bounds.min_row, middle_col),
                (bounds.max_row, middle_col),
            }
        )
    )


def _spread(items: Sequence[tuple[Any, ...]]) -> Iterable[tuple[Any, ...]]:
    """Yield edges then recursively spaced midpoints, never just a top crop."""

    if not items:
        return
    yielded: set[int] = set()
    for index in (0, len(items) - 1):
        if index not in yielded:
            yielded.add(index)
            yield items[index]
    queue: deque[tuple[int, int]] = deque()
    if len(items) > 2:
        queue.append((1, len(items) - 2))
    while queue:
        low, high = queue.popleft()
        if low > high:
            continue
        middle = (low + high) // 2
        if middle not in yielded:
            yielded.add(middle)
            yield items[middle]
        queue.append((low, middle - 1))
        queue.append((middle + 1, high))


def _priority(roles: set[str]) -> int:
    if roles.intersection({"answer", "answer_sample", "data", "data_sample"}):
        return 0
    if "near_answer_formula" in roles:
        return 1
    if "data_boundary" in roles:
        return 2
    if "instruction_ref" in roles:
        return 3
    if "formula" in roles:
        return 4
    return 5


def _candidate_order(
    candidates: Mapping[tuple[str, int, int], set[str]],
    sheet_order: Mapping[str, int],
) -> Iterable[tuple[str, int, int]]:
    groups: list[deque[tuple[str, int, int]]] = []
    for priority in range(len(_PRIORITY_WEIGHTS)):
        ordered = sorted(
            (key for key, roles in candidates.items() if _priority(roles) == priority),
            key=lambda item: (sheet_order[item[0]], item[1], item[2]),
        )
        groups.append(deque(_spread(ordered)))
    while any(groups):
        for priority, weight in enumerate(_PRIORITY_WEIGHTS):
            for _ in range(weight):
                if groups[priority]:
                    yield groups[priority].popleft()


def _json_line(record: Mapping[str, Any]) -> str:
    return json.dumps(
        record,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        allow_nan=False,
    )


def _document_text(records: Sequence[Mapping[str, Any]], cut: Mapping[str, Any]) -> str:
    return "".join(f"{_json_line(record)}\n" for record in (*records, cut))


def _cut_from_record_chars(
    record_chars: int,
    *,
    char_budget: int,
    included_cells: int,
    omitted_cells: int,
    instruction_truncated: bool,
) -> tuple[dict[str, Any], int]:
    cut: dict[str, Any] = {
        "budget_chars": char_budget,
        "included_cells": included_cells,
        "instruction_truncated": instruction_truncated,
        "omitted_cells": omitted_cells,
        "truncated": bool(omitted_cells or instruction_truncated),
        "type": "cut",
        "used_chars": 0,
    }
    for _ in range(12):
        length = record_chars + len(_json_line(cut)) + 1
        if cut["used_chars"] == length:
            return cut, length
        cut["used_chars"] = length
    length = record_chars + len(_json_line(cut)) + 1
    if cut["used_chars"] != length:  # pragma: no cover - integer width converges immediately
        raise RuntimeError("CUT character accounting did not converge")
    return cut, length


def _final_cut(
    records: Sequence[Mapping[str, Any]],
    *,
    char_budget: int,
    included_cells: int,
    omitted_cells: int,
    instruction_truncated: bool,
) -> tuple[dict[str, Any], str]:
    record_chars = sum(len(_json_line(record)) + 1 for record in records)
    cut, _ = _cut_from_record_chars(
        record_chars,
        char_budget=char_budget,
        included_cells=included_cells,
        omitted_cells=omitted_cells,
        instruction_truncated=instruction_truncated,
    )
    text = _document_text(records, cut)
    return cut, text


def _mandatory_records(
    task: Mapping[str, Any],
    formula_wb: Any,
    answer_refs: Sequence[PositionRef],
    data_refs: Sequence[PositionRef],
    input_sha256: str,
) -> list[dict[str, Any]]:
    task_record: dict[str, Any] = {
        "answer_position": task.get("answer_position"),
        "answer_sheet": task.get("answer_sheet"),
        "data_position": task.get("data_position"),
        "id": str(task.get("id", "")),
        "instruction_type": task.get("instruction_type"),
        "type": "task",
    }
    records: list[dict[str, Any]] = [
        task_record,
        {
            "date_system": "1904" if str(formula_wb.epoch).startswith("1904") else "1900",
            "input_sha256": input_sha256,
            "sheet_count": len(formula_wb.sheetnames),
            "sheets": list(formula_wb.sheetnames),
            "type": "workbook",
        },
    ]
    relevant_bounds: dict[str, list[_Bounds]] = {name: [] for name in formula_wb.sheetnames}
    for ref in (*answer_refs, *data_refs):
        if ref.sheet in formula_wb.sheetnames:
            relevant_bounds[ref.sheet].append(_bounds_for(ref, formula_wb[ref.sheet]))
    for worksheet in formula_wb.worksheets:
        materialised = [
            cell
            for cell in worksheet._cells.values()
            if not isinstance(cell, MergedCell) and cell.value is not None
        ]
        sheet_record: dict[str, Any] = {
            "formula_cells": sum(_is_formula(cell.value, cell.data_type) for cell in materialised),
            "max_column": worksheet.max_column,
            "max_row": worksheet.max_row,
            "nonempty_cells": len(materialised),
            "sheet": worksheet.title,
            "state": worksheet.sheet_state,
            "type": "sheet",
        }
        merged_ranges = []
        for merged in sorted(worksheet.merged_cells.ranges, key=str):
            merged_bounds = _Bounds(
                min_col=merged.min_col,
                min_row=merged.min_row,
                max_col=merged.max_col,
                max_row=merged.max_row,
            )
            if any(
                not (
                    merged_bounds.max_row < bounds.min_row
                    or merged_bounds.min_row > bounds.max_row
                    or merged_bounds.max_col < bounds.min_col
                    or merged_bounds.min_col > bounds.max_col
                )
                for bounds in relevant_bounds[worksheet.title]
            ):
                merged_ranges.append(
                    {
                        "anchor": f"{get_column_letter(merged.min_col)}{merged.min_row}",
                        "range": str(merged),
                    }
                )
        if merged_ranges:
            sheet_record["relevant_merged_ranges"] = merged_ranges
        records.append(sheet_record)
    if answer_refs:
        for ref in answer_refs:
            worksheet = formula_wb[ref.sheet] if ref.sheet in formula_wb.sheetnames else None
            bounds = _bounds_for(ref, worksheet) if worksheet is not None else None
            records.append(
                {
                    "cell_count": bounds.cell_count if bounds is not None else None,
                    "dynamic": ref.dynamic,
                    "range": ref.a1_range,
                    "requested_sheet": ref.requested_sheet,
                    "resolution": ref.resolution,
                    "sheet": ref.sheet,
                    "type": "answer",
                }
            )
    else:
        records.append(
            {
                "declared": bool(task.get("answer_position")),
                "status": "unparsed",
                "type": "answer",
            }
        )
    if data_refs:
        for ref in data_refs:
            worksheet = formula_wb[ref.sheet] if ref.sheet in formula_wb.sheetnames else None
            bounds = _bounds_for(ref, worksheet) if worksheet is not None else None
            records.append(
                {
                    "cell_count": bounds.cell_count if bounds is not None else None,
                    "declared": True,
                    "dynamic": ref.dynamic,
                    "range": ref.a1_range,
                    "requested_sheet": ref.requested_sheet,
                    "resolution": ref.resolution,
                    "sheet": ref.sheet,
                    "type": "data",
                }
            )
    else:
        records.append({"declared": False, "type": "data"})
    return records


def _fit_mandatory_records(
    records: list[dict[str, Any]],
    *,
    char_budget: int,
    candidate_count: int,
) -> tuple[list[dict[str, Any]], bool]:
    _, text = _final_cut(
        records,
        char_budget=char_budget,
        included_cells=0,
        omitted_cells=candidate_count,
        instruction_truncated=False,
    )
    if len(text) <= char_budget:
        return records, False
    raise ContextBudgetError(
        f"mandatory workbook records require at least {len(text)} characters; "
        f"budget is {char_budget}"
    )


def _cell_record(
    *,
    sheet: str,
    row: int,
    col: int,
    roles: set[str],
    formula_ws: Any,
    cached_ws: Any,
) -> dict[str, Any]:
    formula_cell = formula_ws._cells.get((row, col))
    cached_cell = cached_ws._cells.get((row, col))
    formula_value = None if formula_cell is None else formula_cell.value
    formula_type = None if formula_cell is None else formula_cell.data_type
    cached_value = None if cached_cell is None else cached_cell.value
    cached_record = _typed_value(cached_value)
    if cached_value is None and _is_formula(formula_value, formula_type):
        cached_record = {
            "reason": "no_cached_result",
            "type": "unknown",
            "value": None,
        }
    coordinate = f"{get_column_letter(col)}{row}"
    return {
        "cached_value": cached_record,
        "cell": coordinate,
        "cell_type": formula_type or "n",
        "formula_view": _formula_view(formula_value, formula_type),
        "priority": _priority(roles),
        "roles": sorted(roles, key=lambda role: _ROLE_ORDER[role]),
        "sheet": sheet,
        "type": "cell",
    }


def _compact_typed_value(value: Any) -> Any:
    """Use native JSON scalars where their type is unambiguous."""

    typed = _typed_value(value)
    if typed["type"] in {"blank", "boolean", "integer", "number", "string"}:
        return typed["value"]
    return typed


def _compact_cell_view(formula_cell: Any, cached_cell: Any) -> Any:
    """Preserve a cell's value or formula semantics without per-cell field repetition."""

    formula_value = None if formula_cell is None else formula_cell.value
    formula_type = None if formula_cell is None else formula_cell.data_type
    cached_value = None if cached_cell is None else cached_cell.value
    if _is_formula(formula_value, formula_type):
        cached_record: Any = _compact_typed_value(cached_value)
        if cached_value is None:
            cached_record = {
                "reason": "no_cached_result",
                "type": "unknown",
                "value": None,
            }
        return {
            "cached_value": cached_record,
            "cell_type": formula_type or "f",
            "formula_view": _formula_view(formula_value, formula_type),
        }
    if formula_type == "e":
        return {"type": "error", "value": str(formula_value)}
    value = formula_value if formula_cell is not None else cached_value
    return _compact_typed_value(value)


def _compact_ranges(
    formula_wb: Any,
    cached_wb: Any,
    refs: Sequence[PositionRef],
    *,
    record_type: str,
) -> tuple[_CompactRange, ...]:
    """Build complete row-major records for small, finite declared ranges."""

    compact_ranges: list[_CompactRange] = []
    claimed: set[tuple[str, int, int]] = set()
    for ref in refs:
        if ref.dynamic or ref.sheet not in formula_wb.sheetnames:
            continue
        bounds = _bounds_for(ref, formula_wb[ref.sheet])
        if bounds.cell_count > SMALL_RANGE_CELL_LIMIT:
            continue
        coordinates = frozenset(
            (ref.sheet, row, col)
            for row in range(bounds.min_row, bounds.max_row + 1)
            for col in range(bounds.min_col, bounds.max_col + 1)
        )
        # Exact or fully contained repeats add no evidence. A partially
        # overlapping range is still rendered in full so its unclaimed cells
        # are never lost.
        if coordinates.issubset(claimed):
            continue
        formula_ws = formula_wb[ref.sheet]
        cached_ws = cached_wb[ref.sheet]
        records: list[dict[str, Any]] = []
        first_column = get_column_letter(bounds.min_col)
        last_column = get_column_letter(bounds.max_col)
        for row in range(bounds.min_row, bounds.max_row + 1):
            values = [
                _compact_cell_view(
                    formula_ws._cells.get((row, col)),
                    cached_ws._cells.get((row, col)),
                )
                for col in range(bounds.min_col, bounds.max_col + 1)
            ]
            records.append(
                {
                    "range": f"{first_column}{row}:{last_column}{row}",
                    "sheet": ref.sheet,
                    "type": record_type,
                    "values": values,
                }
            )
        compact_ranges.append(
            _CompactRange(
                records=tuple(records),
                coordinates=coordinates,
            )
        )
        claimed.update(coordinates)
    return tuple(compact_ranges)


def _compact_answer_ranges(
    formula_wb: Any,
    cached_wb: Any,
    answer_refs: Sequence[PositionRef],
) -> tuple[_CompactRange, ...]:
    """Build compact answer-template evidence before source data is fitted."""

    return _compact_ranges(
        formula_wb,
        cached_wb,
        answer_refs,
        record_type="answer_row",
    )


def _compact_data_ranges(
    formula_wb: Any,
    cached_wb: Any,
    data_refs: Sequence[PositionRef],
) -> tuple[_CompactRange, ...]:
    """Build compact source evidence for small, finite declared data ranges."""

    return _compact_ranges(
        formula_wb,
        cached_wb,
        data_refs,
        record_type="data_row",
    )


def _collect_candidates(
    task: Mapping[str, Any],
    formula_wb: Any,
    cached_wb: Any,
    answer_refs: Sequence[PositionRef],
    data_refs: Sequence[PositionRef],
) -> dict[tuple[str, int, int], set[str]]:
    candidates: dict[tuple[str, int, int], set[str]] = {}
    answer_bounds: dict[str, list[_Bounds]] = {name: [] for name in formula_wb.sheetnames}
    data_bounds: dict[str, list[_Bounds]] = {name: [] for name in formula_wb.sheetnames}

    def add(sheet: str, row: int, col: int, role: str) -> None:
        if 1 <= row <= MAX_EXCEL_ROW and 1 <= col <= MAX_EXCEL_COLUMN:
            candidates.setdefault((sheet, row, col), set()).add(role)

    for ref in answer_refs:
        if ref.sheet not in formula_wb.sheetnames:
            continue
        bounds = _bounds_for(ref, formula_wb[ref.sheet])
        answer_bounds[ref.sheet].append(bounds)
        for row, col in _sample_coordinates(bounds):
            add(ref.sheet, row, col, "answer_sample")

    for ref in data_refs:
        if ref.sheet not in formula_wb.sheetnames:
            continue
        bounds = _bounds_for(ref, formula_wb[ref.sheet])
        data_bounds[ref.sheet].append(bounds)
        for row, col in _sample_coordinates(bounds):
            add(ref.sheet, row, col, "data_sample")
        for row, col in _boundary_coordinates(bounds):
            add(ref.sheet, row, col, "data_boundary")

    default_instruction_sheet = answer_refs[0].sheet if answer_refs else formula_wb.active.title
    instruction_refs = parse_position_references(
        str(task.get("instruction", "")),
        sheet_names=tuple(formula_wb.sheetnames),
        default_sheet=default_instruction_sheet,
    )
    for ref in instruction_refs:
        if ref.sheet not in formula_wb.sheetnames:
            continue
        bounds = _bounds_for(ref, formula_wb[ref.sheet])
        for row, col in _sample_coordinates(bounds):
            add(ref.sheet, row, col, "instruction_ref")

    for worksheet in formula_wb.worksheets:
        cached_ws = cached_wb[worksheet.title]
        materialised = set(worksheet._cells).union(cached_ws._cells)
        for row, col in materialised:
            formula_cell = worksheet._cells.get((row, col))
            cached_cell = cached_ws._cells.get((row, col))
            formula_value = None if formula_cell is None else formula_cell.value
            cached_value = None if cached_cell is None else cached_cell.value
            formula_type = None if formula_cell is None else formula_cell.data_type
            has_formula = _is_formula(formula_value, formula_type)
            in_answer = any(bounds.contains(row, col) for bounds in answer_bounds[worksheet.title])
            in_data = any(bounds.contains(row, col) for bounds in data_bounds[worksheet.title])
            if in_answer:
                add(worksheet.title, row, col, "answer")
            if in_data:
                add(worksheet.title, row, col, "data")
            if has_formula:
                add(worksheet.title, row, col, "formula")
                if any(
                    bounds.distance(row, col) <= NEAR_TARGET_RADIUS
                    for bounds in answer_bounds[worksheet.title]
                ):
                    add(worksheet.title, row, col, "near_answer_formula")
            elif formula_value is not None or cached_value is not None:
                add(worksheet.title, row, col, "value")
    return candidates


def build_context(
    task: Mapping[str, Any],
    *,
    char_budget: int = DEFAULT_CONTEXT_CHARS,
) -> ContextDocument:
    """Build deterministic JSONL evidence without writing to the source workbook."""

    if not isinstance(char_budget, int) or isinstance(char_budget, bool) or char_budget <= 0:
        raise ValueError("char_budget must be a positive integer")
    if "init_xlsx" not in task:
        raise KeyError("task is missing init_xlsx")

    input_path = Path(str(task["init_xlsx"]))
    input_sha256 = hashlib.sha256(input_path.read_bytes()).hexdigest()
    formula_wb = openpyxl.load_workbook(
        input_path,
        data_only=False,
        read_only=False,
        keep_links=False,
    )
    cached_wb = openpyxl.load_workbook(
        input_path,
        data_only=True,
        read_only=False,
        keep_links=False,
    )
    try:
        answer_refs = _contract_answer_positions(task, formula_wb)
        default_data_sheet, default_data_resolution = _data_default_sheet(
            task,
            formula_wb,
            answer_refs,
        )
        data_refs = parse_data_positions(
            task.get("data_position"),
            sheet_names=tuple(formula_wb.sheetnames),
            default_sheet=default_data_sheet,
            default_resolution=default_data_resolution,
        )
        candidates = _collect_candidates(
            task,
            formula_wb,
            cached_wb,
            answer_refs,
            data_refs,
        )
        records = _mandatory_records(
            task,
            formula_wb,
            answer_refs,
            data_refs,
            input_sha256,
        )
        records, instruction_truncated = _fit_mandatory_records(
            records,
            char_budget=char_budget,
            candidate_count=len(candidates),
        )
        sheet_order = {name: index for index, name in enumerate(formula_wb.sheetnames)}
        candidate_coordinates = set(candidates)
        included_coordinates: set[tuple[str, int, int]] = set()
        record_chars = sum(len(_json_line(record)) + 1 for record in records)

        def fit_compact_ranges(compact_ranges: Sequence[_CompactRange]) -> None:
            nonlocal included_coordinates, record_chars
            for compact_range in compact_ranges:
                compact_chars = sum(len(_json_line(record)) + 1 for record in compact_range.records)
                prospective = included_coordinates.union(compact_range.coordinates)
                _, tentative_chars = _cut_from_record_chars(
                    record_chars + compact_chars,
                    char_budget=char_budget,
                    included_cells=len(prospective),
                    omitted_cells=len(candidate_coordinates.difference(prospective)),
                    instruction_truncated=instruction_truncated,
                )
                if tentative_chars <= char_budget:
                    records.extend(compact_range.records)
                    included_coordinates = prospective
                    record_chars += compact_chars

        # The existing output is often the only worked template for a task.
        # Fit it completely before spending budget on compact source ranges.
        fit_compact_ranges(_compact_answer_ranges(formula_wb, cached_wb, answer_refs))
        fit_compact_ranges(_compact_data_ranges(formula_wb, cached_wb, data_refs))

        detail_candidates = {
            coordinate: roles
            for coordinate, roles in candidates.items()
            if coordinate not in included_coordinates
        }
        for sheet, row, col in _candidate_order(detail_candidates, sheet_order):
            candidate = _cell_record(
                sheet=sheet,
                row=row,
                col=col,
                roles=detail_candidates[(sheet, row, col)],
                formula_ws=formula_wb[sheet],
                cached_ws=cached_wb[sheet],
            )
            candidate_chars = len(_json_line(candidate)) + 1
            prospective = included_coordinates.union({(sheet, row, col)})
            _, tentative_chars = _cut_from_record_chars(
                record_chars + candidate_chars,
                char_budget=char_budget,
                included_cells=len(prospective),
                omitted_cells=len(candidate_coordinates.difference(prospective)),
                instruction_truncated=instruction_truncated,
            )
            if tentative_chars <= char_budget:
                records.append(candidate)
                included_coordinates = prospective
                record_chars += candidate_chars
        included_cells = len(included_coordinates)
        omitted_cells = len(candidate_coordinates.difference(included_coordinates))
        _, jsonl = _final_cut(
            records,
            char_budget=char_budget,
            included_cells=included_cells,
            omitted_cells=omitted_cells,
            instruction_truncated=instruction_truncated,
        )
        if len(jsonl) > char_budget:  # defensive invariant
            raise RuntimeError("context exceeded its exact character budget")
        return ContextDocument(
            jsonl=jsonl,
            char_budget=char_budget,
            used_chars=len(jsonl),
            record_count=len(records) + 1,
            included_cells=included_cells,
            omitted_cells=omitted_cells,
            truncated=bool(omitted_cells or instruction_truncated),
            input_sha256=input_sha256,
        )
    finally:
        formula_wb.close()
        cached_wb.close()


def serialize_context(
    task: Mapping[str, Any],
    *,
    char_budget: int = DEFAULT_CONTEXT_CHARS,
) -> str:
    """Return only the JSONL text for prompt assembly."""

    return build_context(task, char_budget=char_budget).jsonl


__all__ = [
    "ContextBudgetError",
    "ContextDocument",
    "DEFAULT_CONTEXT_CHARS",
    "PositionRef",
    "build_context",
    "parse_data_positions",
    "parse_position_references",
    "serialize_context",
]
