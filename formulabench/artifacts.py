"""Safe dataset discovery and durable submission artefact helpers.

This module deliberately knows nothing about golden workbooks.  A private
evaluation dataset may not contain them at all, and inference code must never
look for them.
"""

from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import stat
import tempfile
import zipfile
from collections.abc import Iterable, Iterator, Mapping
from contextlib import contextmanager, suppress
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from types import MappingProxyType
from typing import Any

import openpyxl

from .constants import MODEL_PROVENANCE

try:  # POSIX in Docker; kept optional so the helpers remain importable elsewhere.
    import fcntl
except ImportError:  # pragma: no cover - exercised only on non-POSIX platforms.
    fcntl = None


_SAFE_TASK_ID = re.compile(r"[A-Za-z0-9](?:[A-Za-z0-9_-]{0,126}[A-Za-z0-9])?\Z")
_MAX_MANIFEST_BYTES = 32 * 1024 * 1024
_MAX_JSONL_LINE_BYTES = 16 * 1024 * 1024
_MAX_WORKBOOK_MEMBERS = 20_000
_MAX_WORKBOOK_UNCOMPRESSED_BYTES = 1024 * 1024 * 1024
_RETRY_ROOT_NAME = ".retry"
SUPERVISOR_LOCK_FD_ENV = "FORMULABENCH_SUPERVISOR_LOCK_FD"
_RESUME_ROOT_ENTRIES = frozenset(
    {"predictions.jsonl", "outputs", "traces", "run.log", _RETRY_ROOT_NAME}
)
_RETRY_FORMAT_VERSION = 1
_RETRY_REQUIRED_FILES = frozenset(
    {
        "old-output.xlsx",
        "old-trace.jsonl",
        "new-output.xlsx",
        "new-trace.jsonl",
        "meta.json",
    }
)
_RETRY_PUBLISH_FILES = frozenset(
    {"publish-output.xlsx", "publish-trace.jsonl", "publish-predictions.jsonl"}
)
_RETRY_ALL_FILES = _RETRY_REQUIRED_FILES | _RETRY_PUBLISH_FILES
_RETRY_HASH_NAMES = frozenset({"old_output", "old_trace", "new_output", "new_trace"})
_TRACE_REQUIRED_FIELDS = frozenset(
    {
        "step",
        "model",
        "prompt",
        "response",
        "input_tokens",
        "output_tokens",
        "latency_ms",
        "error",
    }
)


class ArtifactError(ValueError):
    """Raised when an artefact or path violates the submission contract."""


class DatasetManifestError(ArtifactError):
    """Raised when the evaluation dataset manifest is unsafe or malformed."""


class ResumeStateError(ArtifactError):
    """Raised when prior output cannot be trusted as a resume checkpoint."""

    def __init__(self, code: str, *, task_id: str | None = None) -> None:
        self.code = code
        self.task_id = task_id
        suffix = f" for task {task_id}" if task_id is not None else ""
        super().__init__(f"resume state invalid: {code}{suffix}")


@dataclass(frozen=True, slots=True)
class DatasetTask:
    """A validated task with one confined initial workbook."""

    id: str
    instruction: str
    spreadsheet_path: str
    init_xlsx: Path
    metadata: Mapping[str, Any]

    def __getitem__(self, key: str) -> Any:
        """Offer the small mapping surface expected by the upstream scaffold."""
        if key == "id":
            return self.id
        if key == "instruction":
            return self.instruction
        if key == "spreadsheet_path":
            return self.spreadsheet_path
        if key == "init_xlsx":
            return str(self.init_xlsx)
        return self.metadata[key]

    def get(self, key: str, default: Any = None) -> Any:
        try:
            return self[key]
        except KeyError:
            return default

    def as_dict(self) -> dict[str, Any]:
        task = dict(self.metadata)
        task.update(
            {
                "id": self.id,
                "instruction": self.instruction,
                "spreadsheet_path": self.spreadsheet_path,
                "init_xlsx": str(self.init_xlsx),
            }
        )
        return task


@dataclass(frozen=True, slots=True)
class ArtifactLayout:
    """Canonical paths under one output root."""

    root: Path
    predictions: Path
    outputs: Path
    traces: Path
    run_log: Path

    @classmethod
    def at(cls, root: str | os.PathLike[str]) -> ArtifactLayout:
        resolved = Path(root).resolve(strict=False)
        return cls(
            root=resolved,
            predictions=resolved / "predictions.jsonl",
            outputs=resolved / "outputs",
            traces=resolved / "traces",
            run_log=resolved / "run.log",
        )

    @classmethod
    def initialise(
        cls,
        root: str | os.PathLike[str],
        *,
        allow_supervisor_log: bool = False,
    ) -> ArtifactLayout:
        """Create the output structure without deleting or overwriting anything.

        ``capture`` creates ``run.log`` before the child CLI starts.  The child
        may therefore opt into exactly that one pre-existing, regular file.
        """
        layout = cls.at(root)
        allowed = {"run.log"} if allow_supervisor_log else set()
        require_writable_directory(layout.root, allowed_entries=allowed)
        layout.outputs.mkdir(mode=0o700)
        layout.traces.mkdir(mode=0o700)
        atomic_write_text(layout.predictions, "")
        if not layout.run_log.exists():
            atomic_write_text(layout.run_log, "")
        return layout

    @classmethod
    def resume(cls, root: str | os.PathLike[str]) -> ArtifactLayout:
        """Prepare canonical paths without replacing any prior artefact.

        A process can be interrupted after the supervisor creates ``run.log``
        but before the child creates the other paths.  Missing empty structure
        is therefore completed here; existing files are never overwritten.
        Detailed task-level validation is performed by :func:`load_resume_state`.
        """

        resolved = require_resume_directory(root)
        layout = cls.at(resolved)
        predictions_exist = layout.predictions.exists()
        predictions_have_bytes = predictions_exist and layout.predictions.stat().st_size > 0
        try:
            outputs_have_entries = layout.outputs.exists() and next(layout.outputs.iterdir(), None)
            traces_have_entries = layout.traces.exists() and next(layout.traces.iterdir(), None)
        except OSError as exc:
            raise ResumeStateError("canonical_directory_unreadable") from exc

        # Only the run.log-only state is recoverable automatically.  Once any
        # task artefact or prediction exists, missing canonical counterparts
        # are evidence of interruption or tampering and must remain untouched.
        if not predictions_exist and (outputs_have_entries or traces_have_entries):
            raise ResumeStateError("predictions_missing_with_task_artifacts")
        if not layout.outputs.exists() and (predictions_have_bytes or traces_have_entries):
            raise ResumeStateError("outputs_missing_with_task_artifacts")
        if not layout.traces.exists() and (predictions_have_bytes or outputs_have_entries):
            raise ResumeStateError("traces_missing_with_task_artifacts")
        try:
            layout.outputs.mkdir(mode=0o700, exist_ok=True)
            layout.traces.mkdir(mode=0o700, exist_ok=True)
        except OSError as exc:
            raise ResumeStateError("canonical_directory_unavailable") from exc
        if not predictions_exist:
            atomic_write_text(layout.predictions, "")
        # Re-run the shallow check after creating missing structure so a race
        # cannot replace a canonical path with a link or special file.
        require_resume_directory(resolved)
        return layout


@dataclass(frozen=True, slots=True)
class ResumeState:
    """Validated completed records from an interrupted or partial run."""

    predictions: Mapping[str, Mapping[str, Any]]
    completed_ids: frozenset[str]
    successful_ids: frozenset[str]


@dataclass(frozen=True, slots=True)
class RetryWorkspace:
    """Private paths used to build one replacement without touching canonical files."""

    task_id: str
    directory: Path
    new_output: Path
    new_trace: Path


def normalise_task_id(value: Any) -> str:
    """Return a filename-safe task id, rejecting ambiguous values."""
    if isinstance(value, bool) or not isinstance(value, (str, int)):
        raise DatasetManifestError("task id must be a string or integer")
    task_id = str(value)
    if not _SAFE_TASK_ID.fullmatch(task_id):
        raise DatasetManifestError("task id contains unsafe characters")
    return task_id


def _normalise_relative_path(value: Any, *, label: str) -> Path:
    if not isinstance(value, str) or not value or "\x00" in value:
        raise ArtifactError(f"{label} must be a non-empty relative path")
    # Backslashes are separators on Windows but ordinary characters on POSIX.
    # Rejecting them makes the same manifest safe on both.
    if "\\" in value:
        raise ArtifactError(f"{label} must use forward-slash separators")
    path = Path(value)
    if path.is_absolute() or value.startswith("/"):
        raise ArtifactError(f"{label} must be relative")
    if not path.parts or any(part in ("", ".", "..") for part in path.parts):
        raise ArtifactError(f"{label} contains an unsafe path component")
    return path


def _is_relative_to(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
    except ValueError:
        return False
    return True


def _reject_symlink_components(root: Path, relative: Path, *, include_leaf: bool) -> None:
    current = root
    parts = relative.parts if include_leaf else relative.parts[:-1]
    for part in parts:
        current = current / part
        if current.is_symlink():
            raise ArtifactError("path contains a symbolic link")


def confined_path(
    root: str | os.PathLike[str],
    relative: str | os.PathLike[str],
    *,
    must_exist: bool = True,
    reject_symlinks: bool = True,
) -> Path:
    """Resolve a relative path beneath ``root`` without traversal or symlinks."""
    root_path = Path(root).resolve(strict=True)
    if not root_path.is_dir():
        raise ArtifactError("path root is not a directory")
    rel = _normalise_relative_path(os.fspath(relative), label="path")
    if reject_symlinks:
        _reject_symlink_components(root_path, rel, include_leaf=must_exist)
    try:
        resolved = (root_path / rel).resolve(strict=must_exist)
    except (FileNotFoundError, RuntimeError) as exc:
        raise ArtifactError("path does not exist or cannot be resolved") from exc
    if not _is_relative_to(resolved, root_path):
        raise ArtifactError("path escapes its root")
    return resolved


def load_dataset_manifest(dataset_dir: str | os.PathLike[str]) -> list[DatasetTask]:
    """Load tasks and discover one official-convention initial workbook each.

    Only the manifest and filenames matching the initial-workbook suffix are
    inspected.  Golden files are neither opened nor globbed.
    """
    try:
        root = Path(dataset_dir).resolve(strict=True)
    except (FileNotFoundError, RuntimeError) as exc:
        raise DatasetManifestError("dataset directory does not exist") from exc
    if not root.is_dir():
        raise DatasetManifestError("dataset path is not a directory")

    try:
        manifest_path = confined_path(root, "dataset.json", reject_symlinks=True)
    except ArtifactError as exc:
        raise DatasetManifestError("dataset.json is missing or unsafe") from exc
    if not manifest_path.is_file():
        raise DatasetManifestError("dataset.json is not a regular file")
    if manifest_path.stat().st_size > _MAX_MANIFEST_BYTES:
        raise DatasetManifestError("dataset.json exceeds the size limit")
    try:
        payload = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise DatasetManifestError("dataset.json is not valid UTF-8 JSON") from exc
    if not isinstance(payload, list) or not payload:
        raise DatasetManifestError("dataset.json must contain a non-empty task list")

    tasks: list[DatasetTask] = []
    seen: set[str] = set()
    seen_casefolded: set[str] = set()
    for index, raw in enumerate(payload):
        if not isinstance(raw, dict):
            raise DatasetManifestError(f"task {index} is not an object")
        try:
            task_id = normalise_task_id(raw.get("id"))
        except DatasetManifestError as exc:
            raise DatasetManifestError(f"task {index} has an invalid id") from exc
        if task_id in seen or task_id.casefold() in seen_casefolded:
            raise DatasetManifestError(f"duplicate task id at task {index}")
        seen.add(task_id)
        seen_casefolded.add(task_id.casefold())

        instruction = raw.get("instruction")
        if not isinstance(instruction, str):
            raise DatasetManifestError(f"task {index} has no string instruction")
        try:
            relative_folder = _normalise_relative_path(
                raw.get("spreadsheet_path"), label="spreadsheet_path"
            )
            folder = confined_path(root, relative_folder, reject_symlinks=True)
        except ArtifactError as exc:
            raise DatasetManifestError(f"task {index} has an unsafe spreadsheet path") from exc
        if not folder.is_dir():
            raise DatasetManifestError(f"task {index} spreadsheet path is not a directory")

        # Both names occur in the official 400-task release.  The patterns are
        # intentionally narrow: do not enumerate arbitrary workbooks, and
        # especially do not inspect golden workbook content.
        init_candidates = sorted(folder.glob("*_init.xlsx"), key=lambda path: path.name)
        exact_initial = folder / "initial.xlsx"
        if exact_initial.exists() or exact_initial.is_symlink():
            init_candidates.append(exact_initial)
        if len(init_candidates) != 1:
            raise DatasetManifestError(f"task {index} must have exactly one initial workbook")
        candidate = init_candidates[0]
        try:
            relative_candidate = candidate.relative_to(root)
            init_xlsx = confined_path(root, relative_candidate, reject_symlinks=True)
        except (ValueError, ArtifactError) as exc:
            raise DatasetManifestError(f"task {index} initial workbook is unsafe") from exc
        if not init_xlsx.is_file():
            raise DatasetManifestError(f"task {index} initial workbook is not a regular file")

        metadata = dict(raw)
        metadata["id"] = task_id
        tasks.append(
            DatasetTask(
                id=task_id,
                instruction=instruction,
                spreadsheet_path=relative_folder.as_posix(),
                init_xlsx=init_xlsx,
                metadata=MappingProxyType(metadata),
            )
        )
    return tasks


def require_writable_directory(
    path: str | os.PathLike[str],
    *,
    allowed_entries: set[str] | frozenset[str] = frozenset(),
) -> Path:
    """Require a writable directory containing no entries except an allow-list.

    Missing directories are created.  Existing content is never deleted.
    """
    candidate = Path(path)
    if candidate.is_symlink():
        raise ArtifactError("output directory must not be a symbolic link")
    try:
        candidate.mkdir(parents=True, exist_ok=True)
        root = candidate.resolve(strict=True)
    except OSError as exc:
        raise ArtifactError("output directory cannot be created") from exc
    if not root.is_dir():
        raise ArtifactError("output path is not a directory")

    entries = list(root.iterdir())
    unexpected = [entry for entry in entries if entry.name not in allowed_entries]
    if unexpected:
        raise ArtifactError("output directory is not empty")
    for entry in entries:
        if entry.is_symlink() or not entry.is_file():
            raise ArtifactError("allowed output entry is not a regular file")

    # os.access is unreliable for elevated users.  A create-and-remove probe
    # tests the actual operation while leaving the directory as it was.
    try:
        fd, probe_name = tempfile.mkstemp(prefix=".write-probe-", dir=root)
        os.close(fd)
        Path(probe_name).unlink()
    except OSError as exc:
        raise ArtifactError("output directory is not writable") from exc
    return root


def require_resume_directory(path: str | os.PathLike[str]) -> Path:
    """Shallow-check an existing output root before appending or resuming.

    Only canonical FormulaBench entries are accepted.  This function does not
    delete, truncate or rewrite anything; task-level consistency is checked by
    :func:`load_resume_state` after the canonical layout is available.
    """

    candidate = Path(path)
    if candidate.is_symlink():
        raise ResumeStateError("output_root_symlink")
    try:
        root = candidate.resolve(strict=True)
    except (FileNotFoundError, OSError, RuntimeError) as exc:
        raise ResumeStateError("output_root_missing") from exc
    if not root.is_dir():
        raise ResumeStateError("output_root_not_directory")

    try:
        entries = {entry.name: entry for entry in root.iterdir()}
    except OSError as exc:
        raise ResumeStateError("output_root_unreadable") from exc
    if unexpected := set(entries).difference(_RESUME_ROOT_ENTRIES):
        del unexpected  # Names may contain private data; never reflect them.
        raise ResumeStateError("unexpected_root_entry")

    run_log = entries.get("run.log")
    if run_log is None:
        raise ResumeStateError("run_log_missing")
    if run_log.is_symlink() or not run_log.is_file():
        raise ResumeStateError("run_log_unsafe")

    predictions = entries.get("predictions.jsonl")
    if predictions is not None and (predictions.is_symlink() or not predictions.is_file()):
        raise ResumeStateError("predictions_unsafe")
    for name in ("outputs", "traces"):
        directory = entries.get(name)
        if directory is not None and (directory.is_symlink() or not directory.is_dir()):
            raise ResumeStateError(f"{name}_unsafe")
    retry_root = entries.get(_RETRY_ROOT_NAME)
    if retry_root is not None and (retry_root.is_symlink() or not retry_root.is_dir()):
        raise ResumeStateError("retry_root_unsafe")

    try:
        fd, probe_name = tempfile.mkstemp(prefix=".resume-write-probe-", dir=root)
        os.close(fd)
        Path(probe_name).unlink()
    except OSError as exc:
        raise ResumeStateError("output_root_not_writable") from exc
    return root


@contextmanager
def exclusive_output_lock(run_log: Path) -> Iterator[None]:
    """Hold the stable run.log lock for one CLI owner of an output root.

    The capture supervisor passes its already-locked descriptor to the child.
    A direct CLI opens and locks the same file itself.  This prevents two CLI
    coordinators from concurrently rewriting the shared predictions checkpoint.
    """

    if fcntl is None:  # pragma: no cover - supported production hosts are POSIX.
        yield
        return

    inherited = os.environ.get(SUPERVISOR_LOCK_FD_ENV)
    owns_lock = inherited is None
    if inherited is None:
        flags = os.O_RDWR | getattr(os, "O_CLOEXEC", 0) | getattr(os, "O_NOFOLLOW", 0)
        try:
            lock_fd = os.open(run_log, flags)
        except OSError as exc:
            raise ResumeStateError("run_log_unavailable") from exc
    else:
        try:
            lock_fd = int(inherited)
        except ValueError as exc:
            raise ResumeStateError("supervisor_lock_invalid") from exc
        if lock_fd < 0:
            raise ResumeStateError("supervisor_lock_invalid")

    try:
        try:
            if run_log.is_symlink():
                raise ResumeStateError("run_log_unsafe")
            descriptor_stat = os.fstat(lock_fd)
            path_stat = os.stat(run_log, follow_symlinks=False)
            if (
                not stat.S_ISREG(descriptor_stat.st_mode)
                or not stat.S_ISREG(path_stat.st_mode)
                or (descriptor_stat.st_dev, descriptor_stat.st_ino)
                != (path_stat.st_dev, path_stat.st_ino)
            ):
                raise ResumeStateError("supervisor_lock_invalid")
            fcntl.flock(lock_fd, fcntl.LOCK_EX | fcntl.LOCK_NB)
        except BlockingIOError as exc:
            raise ResumeStateError("run_already_active") from exc
        except OSError as exc:
            raise ResumeStateError("supervisor_lock_invalid") from exc
        yield
    finally:
        if owns_lock:
            with suppress(OSError):
                fcntl.flock(lock_fd, fcntl.LOCK_UN)
            with suppress(OSError):
                os.close(lock_fd)
        else:
            # Closing only the child's duplicate preserves the supervisor's
            # lock on the shared open-file description until capture finishes.
            with suppress(OSError):
                os.close(lock_fd)


def _atomic_temp_path(destination: Path) -> tuple[int, Path]:
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.is_symlink() or (destination.exists() and not destination.is_file()):
        raise ArtifactError("destination is not a regular file")
    fd, name = tempfile.mkstemp(prefix=f".{destination.name}.", dir=destination.parent)
    return fd, Path(name)


def _fsync_directory(directory: Path) -> None:
    """Persist a same-directory rename on the supported POSIX runtimes."""

    if os.name != "posix":  # pragma: no cover - Docker and supported host are POSIX.
        return
    flags = os.O_RDONLY | getattr(os, "O_DIRECTORY", 0) | getattr(os, "O_CLOEXEC", 0)
    directory_fd = os.open(directory, flags)
    try:
        os.fsync(directory_fd)
    finally:
        os.close(directory_fd)


def atomic_write_bytes(path: str | os.PathLike[str], data: bytes, *, mode: int = 0o600) -> Path:
    """Durably replace a file from a same-directory temporary file."""
    if not isinstance(data, bytes):
        raise TypeError("data must be bytes")
    destination = Path(path)
    fd, temporary = _atomic_temp_path(destination)
    try:
        os.fchmod(fd, mode)
        with os.fdopen(fd, "wb", closefd=True) as handle:
            handle.write(data)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(temporary, destination)
        _fsync_directory(destination.parent)
    except BaseException:
        with suppress(OSError):
            os.close(fd)
        temporary.unlink(missing_ok=True)
        raise
    return destination


def atomic_write_text(path: str | os.PathLike[str], text: str, *, mode: int = 0o600) -> Path:
    if not isinstance(text, str):
        raise TypeError("text must be a string")
    return atomic_write_bytes(path, text.encode("utf-8"), mode=mode)


def atomic_write_json(path: str | os.PathLike[str], value: Any, *, mode: int = 0o600) -> Path:
    encoded = json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    return atomic_write_text(path, encoded, mode=mode)


def atomic_write_jsonl(
    path: str | os.PathLike[str], records: Iterable[Mapping[str, Any]], *, mode: int = 0o600
) -> Path:
    lines = [json.dumps(record, ensure_ascii=False, separators=(",", ":")) for record in records]
    return atomic_write_text(path, "".join(f"{line}\n" for line in lines), mode=mode)


def atomic_copy(
    source: str | os.PathLike[str],
    destination: str | os.PathLike[str],
    *,
    mode: int = 0o600,
) -> Path:
    """Copy exact bytes into place without exposing a partial workbook."""
    source_path = Path(source)
    if not source_path.is_file():
        raise ArtifactError("copy source is not a regular file")
    destination_path = Path(destination)
    fd, temporary = _atomic_temp_path(destination_path)
    try:
        os.fchmod(fd, mode)
        with (
            source_path.open("rb") as source_handle,
            os.fdopen(fd, "wb", closefd=True) as destination_handle,
        ):
            shutil.copyfileobj(source_handle, destination_handle, length=1024 * 1024)
            destination_handle.flush()
            os.fsync(destination_handle.fileno())
        os.replace(temporary, destination_path)
        _fsync_directory(destination_path.parent)
    except BaseException:
        with suppress(OSError):
            os.close(fd)
        temporary.unlink(missing_ok=True)
        raise
    return destination_path


def append_jsonl(path: str | os.PathLike[str], record: Mapping[str, Any]) -> None:
    """Append exactly one compact JSON record under an advisory file lock."""
    if not isinstance(record, Mapping):
        raise TypeError("record must be a mapping")
    payload = (json.dumps(record, ensure_ascii=False, separators=(",", ":")) + "\n").encode("utf-8")
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.is_symlink():
        raise ArtifactError("JSONL destination must not be a symbolic link")
    flags = os.O_APPEND | os.O_CREAT | os.O_WRONLY
    flags |= getattr(os, "O_CLOEXEC", 0)
    flags |= getattr(os, "O_NOFOLLOW", 0)
    fd = os.open(destination, flags, 0o600)
    try:
        if fcntl is not None:
            fcntl.flock(fd, fcntl.LOCK_EX)
        view = memoryview(payload)
        while view:
            written = os.write(fd, view)
            if written <= 0:  # Defensive: os.write should raise instead.
                raise OSError("short JSONL append")
            view = view[written:]
        os.fsync(fd)
    finally:
        if fcntl is not None:
            fcntl.flock(fd, fcntl.LOCK_UN)
        os.close(fd)


def iter_jsonl(path: str | os.PathLike[str]) -> Iterator[tuple[int, Any]]:
    """Yield one decoded JSON value per non-blank line, with line numbers."""
    with Path(path).open("r", encoding="utf-8") as handle:
        for line_number, line in enumerate(handle, start=1):
            if not line.strip():
                continue
            try:
                yield line_number, json.loads(line)
            except json.JSONDecodeError as exc:
                raise ArtifactError(f"invalid JSON on line {line_number}") from exc


def sha256_file(path: str | os.PathLike[str], *, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        while chunk := handle.read(chunk_size):
            digest.update(chunk)
    return digest.hexdigest()


def _read_resume_jsonl(path: Path, *, code: str, task_id: str | None = None) -> list[Any]:
    records: list[Any] = []
    try:
        with path.open("r", encoding="utf-8") as handle:
            for line in handle:
                if not line.strip():
                    continue
                if len(line.encode("utf-8")) > _MAX_JSONL_LINE_BYTES:
                    raise ResumeStateError(f"{code}_line_too_large", task_id=task_id)
                try:
                    records.append(json.loads(line))
                except json.JSONDecodeError as exc:
                    raise ResumeStateError(f"{code}_invalid_json", task_id=task_id) from exc
    except ResumeStateError:
        raise
    except (OSError, UnicodeError) as exc:
        raise ResumeStateError(f"{code}_unreadable", task_id=task_id) from exc
    return records


def _resume_workbook_is_readable(path: Path) -> bool:
    """Read every worksheet while bounding hostile OOXML expansion."""

    if path.is_symlink() or not path.is_file() or not zipfile.is_zipfile(path):
        return False
    workbook = None
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > _MAX_WORKBOOK_MEMBERS:
                return False
            if sum(member.file_size for member in members) > _MAX_WORKBOOK_UNCOMPRESSED_BYTES:
                return False
            for member in members:
                member_path = PurePosixPath(member.filename)
                if member_path.is_absolute() or ".." in member_path.parts:
                    return False
            if archive.testzip() is not None:
                return False
        workbook = openpyxl.load_workbook(
            path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        if not workbook.sheetnames:
            return False
        for worksheet in workbook.worksheets:
            for _ in worksheet.iter_rows(values_only=True):
                pass
        return True
    except Exception:
        return False
    finally:
        if workbook is not None:
            workbook.close()


def _valid_optional_non_negative_int(value: Any) -> bool:
    return value is None or (isinstance(value, int) and not isinstance(value, bool) and value >= 0)


def _validate_resume_trace(
    path: Path,
    *,
    task_id: str,
    expected_model: str,
    expected_model_provenance: str | None = None,
    require_successful_call: bool,
) -> None:
    records = _read_resume_jsonl(path, code="trace", task_id=task_id)
    successful_call = False
    for expected_step, record in enumerate(records, start=1):
        if not isinstance(record, dict):
            raise ResumeStateError("trace_record_type", task_id=task_id)
        if _TRACE_REQUIRED_FIELDS.difference(record):
            raise ResumeStateError("trace_missing_fields", task_id=task_id)
        step = record.get("step")
        if isinstance(step, bool) or not isinstance(step, int) or step != expected_step:
            raise ResumeStateError("trace_step_order", task_id=task_id)
        if record.get("model") != expected_model:
            raise ResumeStateError("trace_model", task_id=task_id)
        if expected_model_provenance is not None:
            recorded_provenance = record.get("model_provenance")
            if recorded_provenance is None and expected_model_provenance == MODEL_PROVENANCE:
                # Base-model traces created before provenance binding remain
                # resumable. Checkpoint traces never use this legacy default.
                recorded_provenance = MODEL_PROVENANCE
            request = record.get("request")
            request_provenance = (
                request.get("model_provenance") if isinstance(request, dict) else None
            )
            if request_provenance is None and expected_model_provenance == MODEL_PROVENANCE:
                request_provenance = MODEL_PROVENANCE
            if (
                recorded_provenance != expected_model_provenance
                or request_provenance != expected_model_provenance
            ):
                raise ResumeStateError("trace_model_provenance", task_id=task_id)
        for field in ("prompt", "response", "error"):
            if record.get(field) is not None and not isinstance(record.get(field), str):
                raise ResumeStateError("trace_field_type", task_id=task_id)
        for field in ("input_tokens", "output_tokens", "latency_ms"):
            if not _valid_optional_non_negative_int(record.get(field)):
                raise ResumeStateError("trace_metric_type", task_id=task_id)
        failure_codes = record.get("failure_codes", [])
        if not isinstance(failure_codes, list) or any(
            not isinstance(value, str) or not value for value in failure_codes
        ):
            raise ResumeStateError("trace_failure_codes", task_id=task_id)
        if record.get("error") is None and isinstance(record.get("response"), str):
            if failure_codes:
                raise ResumeStateError("trace_success_mismatch", task_id=task_id)
            successful_call = True
    if require_successful_call and not successful_call:
        raise ResumeStateError("trace_no_successful_call", task_id=task_id)


def _validate_resume_prediction(
    record: Any,
    *,
    task_by_id: Mapping[str, DatasetTask],
) -> tuple[str, dict[str, Any]]:
    if not isinstance(record, dict):
        raise ResumeStateError("prediction_record_type")
    if set(record) != {"id", "output", "status", "failure_codes"}:
        raise ResumeStateError("prediction_schema")
    raw_id = record.get("id")
    if not isinstance(raw_id, str):
        raise ResumeStateError("prediction_id_type")
    try:
        task_id = normalise_task_id(raw_id)
    except ArtifactError as exc:
        raise ResumeStateError("prediction_id_unsafe") from exc
    if task_id not in task_by_id:
        raise ResumeStateError("prediction_task_not_selected", task_id=task_id)

    expected_output = f"outputs/{task_id}.xlsx"
    if record.get("output") != expected_output:
        raise ResumeStateError("prediction_output_layout", task_id=task_id)
    status = record.get("status")
    failure_codes = record.get("failure_codes")
    if not isinstance(status, str) or not isinstance(failure_codes, list):
        raise ResumeStateError("prediction_status_schema", task_id=task_id)
    if any(not isinstance(code, str) or not code for code in failure_codes):
        raise ResumeStateError("prediction_failure_codes", task_id=task_id)
    if len(failure_codes) != len(set(failure_codes)):
        raise ResumeStateError("prediction_failure_codes", task_id=task_id)
    if status == "ok":
        if failure_codes:
            raise ResumeStateError("prediction_status_mismatch", task_id=task_id)
    elif not failure_codes or status != f"error:{','.join(failure_codes)}":
        raise ResumeStateError("prediction_status_mismatch", task_id=task_id)
    return task_id, dict(record)


@dataclass(frozen=True, slots=True)
class _RetryTransaction:
    workspace: RetryWorkspace
    old_prediction: dict[str, Any]
    new_prediction: dict[str, Any]
    hashes: Mapping[str, str]


def _retry_root(layout: ArtifactLayout) -> Path:
    return layout.root / _RETRY_ROOT_NAME


def _retry_filename_is_temporary(name: str) -> bool:
    return any(name.startswith(f".{base}.") for base in _RETRY_ALL_FILES)


def _retry_directory_entries(directory: Path) -> dict[str, Path]:
    try:
        entries = list(directory.iterdir())
    except OSError as exc:
        raise ResumeStateError("retry_workspace_unreadable", task_id=directory.name) from exc
    result: dict[str, Path] = {}
    for entry in entries:
        if entry.is_symlink() or not entry.is_file():
            raise ResumeStateError("retry_workspace_entry_unsafe", task_id=directory.name)
        if entry.name not in _RETRY_ALL_FILES and not _retry_filename_is_temporary(entry.name):
            raise ResumeStateError("retry_workspace_entry_unexpected", task_id=directory.name)
        result[entry.name] = entry
    return result


def _remove_retry_root_if_empty(layout: ArtifactLayout) -> None:
    root = _retry_root(layout)
    if not root.exists():
        return
    if root.is_symlink() or not root.is_dir():
        raise ResumeStateError("retry_root_unsafe")
    try:
        if next(root.iterdir(), None) is not None:
            return
        root.rmdir()
        _fsync_directory(layout.root)
    except OSError as exc:
        raise ResumeStateError("retry_root_cleanup_failed") from exc


def _remove_retry_workspace(layout: ArtifactLayout, workspace: RetryWorkspace) -> None:
    expected = _retry_root(layout) / workspace.task_id
    if workspace.directory != expected:
        raise ResumeStateError("retry_workspace_path", task_id=workspace.task_id)
    directory = workspace.directory
    if not directory.exists():
        _remove_retry_root_if_empty(layout)
        return
    if directory.is_symlink() or not directory.is_dir():
        raise ResumeStateError("retry_workspace_unsafe", task_id=workspace.task_id)
    entries = _retry_directory_entries(directory)

    # Removing the durable marker first makes cleanup itself recoverable.  If
    # the process stops afterwards, the remaining files are known scratch and
    # an ordinary resume may remove them without interpreting partial state.
    marker = entries.pop("meta.json", None)
    try:
        if marker is not None:
            marker.unlink()
            _fsync_directory(directory)
        for entry in entries.values():
            entry.unlink()
        _fsync_directory(directory)
        directory.rmdir()
        _fsync_directory(directory.parent)
    except OSError as exc:
        raise ResumeStateError("retry_workspace_cleanup_failed", task_id=workspace.task_id) from exc
    _remove_retry_root_if_empty(layout)


def discard_retry_workspace(layout: ArtifactLayout, workspace: RetryWorkspace) -> None:
    """Remove an unprepared retry workspace without touching canonical files."""

    if (workspace.directory / "meta.json").exists():
        raise ResumeStateError("retry_workspace_prepared", task_id=workspace.task_id)
    _remove_retry_workspace(layout, workspace)


def create_retry_workspace(layout: ArtifactLayout, task_id: str) -> RetryWorkspace:
    """Copy the old canonical artefacts aside before a paid retry starts."""

    try:
        safe_task_id = normalise_task_id(task_id)
    except ArtifactError as exc:
        raise ResumeStateError("retry_task_id_unsafe") from exc
    root = _retry_root(layout)
    if root.exists() and (root.is_symlink() or not root.is_dir()):
        raise ResumeStateError("retry_root_unsafe")
    if not root.exists():
        try:
            root.mkdir(mode=0o700)
            _fsync_directory(layout.root)
        except OSError as exc:
            raise ResumeStateError("retry_root_unavailable") from exc
    directory = root / safe_task_id
    try:
        directory.mkdir(mode=0o700)
        _fsync_directory(root)
    except FileExistsError as exc:
        raise ResumeStateError("retry_workspace_exists", task_id=safe_task_id) from exc
    except OSError as exc:
        raise ResumeStateError("retry_workspace_unavailable", task_id=safe_task_id) from exc

    workspace = RetryWorkspace(
        task_id=safe_task_id,
        directory=directory,
        new_output=directory / "new-output.xlsx",
        new_trace=directory / "new-trace.jsonl",
    )
    old_output = layout.outputs / f"{safe_task_id}.xlsx"
    old_trace = layout.traces / f"{safe_task_id}.jsonl"
    try:
        if old_output.is_symlink() or not old_output.is_file():
            raise ResumeStateError("retry_old_output_unsafe", task_id=safe_task_id)
        if old_trace.is_symlink() or not old_trace.is_file():
            raise ResumeStateError("retry_old_trace_unsafe", task_id=safe_task_id)
        atomic_copy(old_output, directory / "old-output.xlsx")
        atomic_copy(old_trace, directory / "old-trace.jsonl")
    except BaseException:
        with suppress(ArtifactError, ResumeStateError, OSError):
            discard_retry_workspace(layout, workspace)
        raise
    return workspace


def _validate_retry_payloads(
    workspace: RetryWorkspace,
    *,
    task: DatasetTask,
    old_prediction: Mapping[str, Any],
    new_prediction: Mapping[str, Any],
    expected_model: str,
    expected_model_provenance: str | None = None,
) -> dict[str, str]:
    task_by_id = {task.id: task}
    old_id, validated_old = _validate_resume_prediction(old_prediction, task_by_id=task_by_id)
    new_id, validated_new = _validate_resume_prediction(new_prediction, task_by_id=task_by_id)
    if old_id != workspace.task_id or new_id != workspace.task_id:
        raise ResumeStateError("retry_prediction_id", task_id=workspace.task_id)
    if validated_old["status"] == "ok":
        raise ResumeStateError("retry_old_prediction_successful", task_id=workspace.task_id)

    paths = {
        "old_output": workspace.directory / "old-output.xlsx",
        "old_trace": workspace.directory / "old-trace.jsonl",
        "new_output": workspace.new_output,
        "new_trace": workspace.new_trace,
    }
    for path in paths.values():
        if path.is_symlink() or not path.is_file():
            raise ResumeStateError("retry_payload_missing", task_id=workspace.task_id)
    try:
        for path in paths.values():
            with path.open("rb") as handle:
                os.fsync(handle.fileno())
        _fsync_directory(workspace.directory)
    except OSError as exc:
        raise ResumeStateError("retry_payload_sync_failed", task_id=workspace.task_id) from exc
    if not _resume_workbook_is_readable(paths["old_output"]):
        raise ResumeStateError("retry_old_workbook_unreadable", task_id=workspace.task_id)
    if not _resume_workbook_is_readable(paths["new_output"]):
        raise ResumeStateError("retry_new_workbook_unreadable", task_id=workspace.task_id)
    try:
        input_hash = sha256_file(task.init_xlsx)
        hashes = {name: sha256_file(path) for name, path in paths.items()}
    except OSError as exc:
        raise ResumeStateError("retry_payload_unreadable", task_id=workspace.task_id) from exc
    if hashes["old_output"] != input_hash:
        raise ResumeStateError("retry_old_fallback_hash", task_id=workspace.task_id)
    if validated_new["status"] != "ok" and hashes["new_output"] != input_hash:
        raise ResumeStateError("retry_new_fallback_hash", task_id=workspace.task_id)
    _validate_resume_trace(
        paths["old_trace"],
        task_id=workspace.task_id,
        expected_model=expected_model,
        expected_model_provenance=expected_model_provenance,
        require_successful_call=False,
    )
    _validate_resume_trace(
        paths["new_trace"],
        task_id=workspace.task_id,
        expected_model=expected_model,
        expected_model_provenance=expected_model_provenance,
        require_successful_call=validated_new["status"] == "ok",
    )
    return hashes


def prepare_retry_transaction(
    workspace: RetryWorkspace,
    *,
    task: DatasetTask,
    old_prediction: Mapping[str, Any],
    new_prediction: Mapping[str, Any],
    expected_model: str,
    expected_model_provenance: str | None = None,
) -> None:
    """Durably mark a complete replacement as recoverable before publication."""

    if workspace.task_id != task.id:
        raise ResumeStateError("retry_task_mismatch", task_id=workspace.task_id)
    if (workspace.directory / "meta.json").exists():
        raise ResumeStateError("retry_already_prepared", task_id=workspace.task_id)
    hashes = _validate_retry_payloads(
        workspace,
        task=task,
        old_prediction=old_prediction,
        new_prediction=new_prediction,
        expected_model=expected_model,
        expected_model_provenance=expected_model_provenance,
    )
    meta = {
        "version": _RETRY_FORMAT_VERSION,
        "task_id": workspace.task_id,
        "old_prediction": dict(old_prediction),
        "new_prediction": dict(new_prediction),
        "hashes": hashes,
    }
    atomic_write_json(workspace.directory / "meta.json", meta)


def _load_retry_transaction(
    workspace: RetryWorkspace,
    *,
    task: DatasetTask,
    expected_model: str,
    expected_model_provenance: str | None = None,
) -> _RetryTransaction:
    entries = _retry_directory_entries(workspace.directory)
    if "meta.json" not in entries:
        raise ResumeStateError("retry_not_prepared", task_id=workspace.task_id)

    # Publish copies and same-directory temporary files are disposable.  The
    # four immutable payloads and marker remain the sole recovery authority.
    try:
        for name, entry in tuple(entries.items()):
            if name in _RETRY_PUBLISH_FILES or _retry_filename_is_temporary(name):
                entry.unlink()
                entries.pop(name)
        _fsync_directory(workspace.directory)
    except OSError as exc:
        raise ResumeStateError("retry_scratch_cleanup_failed", task_id=workspace.task_id) from exc
    if set(entries) != _RETRY_REQUIRED_FILES:
        raise ResumeStateError("retry_payload_set", task_id=workspace.task_id)

    meta_path = entries["meta.json"]
    try:
        if meta_path.stat().st_size > _MAX_JSONL_LINE_BYTES:
            raise ResumeStateError("retry_meta_too_large", task_id=workspace.task_id)
        with meta_path.open("r", encoding="utf-8") as handle:
            meta = json.load(handle)
    except ResumeStateError:
        raise
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise ResumeStateError("retry_meta_unreadable", task_id=workspace.task_id) from exc
    if not isinstance(meta, dict) or set(meta) != {
        "version",
        "task_id",
        "old_prediction",
        "new_prediction",
        "hashes",
    }:
        raise ResumeStateError("retry_meta_schema", task_id=workspace.task_id)
    if meta.get("version") != _RETRY_FORMAT_VERSION or meta.get("task_id") != workspace.task_id:
        raise ResumeStateError("retry_meta_identity", task_id=workspace.task_id)
    raw_hashes = meta.get("hashes")
    if not isinstance(raw_hashes, dict) or set(raw_hashes) != _RETRY_HASH_NAMES:
        raise ResumeStateError("retry_hash_schema", task_id=workspace.task_id)
    if any(
        not isinstance(value, str) or re.fullmatch(r"[0-9a-f]{64}", value) is None
        for value in raw_hashes.values()
    ):
        raise ResumeStateError("retry_hash_schema", task_id=workspace.task_id)

    computed_hashes = _validate_retry_payloads(
        workspace,
        task=task,
        old_prediction=meta["old_prediction"],
        new_prediction=meta["new_prediction"],
        expected_model=expected_model,
        expected_model_provenance=expected_model_provenance,
    )
    if computed_hashes != raw_hashes:
        raise ResumeStateError("retry_payload_hash", task_id=workspace.task_id)
    return _RetryTransaction(
        workspace=workspace,
        old_prediction=dict(meta["old_prediction"]),
        new_prediction=dict(meta["new_prediction"]),
        hashes=MappingProxyType(dict(raw_hashes)),
    )


def _validated_prediction_records(
    predictions_path: Path,
    tasks: list[DatasetTask],
) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    task_by_id = {task.id: task for task in tasks}
    raw_records = _read_resume_jsonl(predictions_path, code="predictions")
    records: list[dict[str, Any]] = []
    by_id: dict[str, dict[str, Any]] = {}
    for raw_record in raw_records:
        task_id, record = _validate_resume_prediction(raw_record, task_by_id=task_by_id)
        if task_id in by_id:
            raise ResumeStateError("prediction_duplicate", task_id=task_id)
        by_id[task_id] = record
        records.append(record)
    expected_order = [task.id for task in tasks if task.id in by_id]
    if [record["id"] for record in records] != expected_order:
        raise ResumeStateError("prediction_order")
    return records, by_id


def _publish_retry_file(source: Path, destination: Path, publish_path: Path) -> None:
    """Publish via a journal-local copy so SIGKILL leaves no canonical temp."""

    atomic_copy(source, publish_path)
    os.replace(publish_path, destination)
    _fsync_directory(destination.parent)


def commit_retry_transaction(
    layout: ArtifactLayout,
    tasks: list[DatasetTask],
    *,
    task_id: str,
    expected_model: str,
    expected_model_provenance: str | None = None,
) -> dict[str, Any]:
    """Idempotently publish one prepared retry and remove its journal."""

    task_by_id = {task.id: task for task in tasks}
    task = task_by_id.get(task_id)
    if task is None:
        raise ResumeStateError("retry_task_not_selected", task_id=task_id)
    workspace = RetryWorkspace(
        task_id=task_id,
        directory=_retry_root(layout) / task_id,
        new_output=_retry_root(layout) / task_id / "new-output.xlsx",
        new_trace=_retry_root(layout) / task_id / "new-trace.jsonl",
    )
    transaction = _load_retry_transaction(
        workspace,
        task=task,
        expected_model=expected_model,
        expected_model_provenance=expected_model_provenance,
    )
    records, predictions = _validated_prediction_records(layout.predictions, tasks)
    current = predictions.get(task_id)
    if current not in (transaction.old_prediction, transaction.new_prediction):
        raise ResumeStateError("retry_prediction_conflict", task_id=task_id)

    canonical_output = layout.outputs / f"{task_id}.xlsx"
    canonical_trace = layout.traces / f"{task_id}.jsonl"
    for path, hash_names, error_code in (
        (
            canonical_output,
            ("old_output", "new_output"),
            "retry_canonical_output_mismatch",
        ),
        (
            canonical_trace,
            ("old_trace", "new_trace"),
            "retry_canonical_trace_mismatch",
        ),
    ):
        if path.is_symlink() or not path.is_file():
            raise ResumeStateError(error_code, task_id=task_id)
        try:
            current_hash = sha256_file(path)
        except OSError as exc:
            raise ResumeStateError(error_code, task_id=task_id) from exc
        if current_hash not in {transaction.hashes[name] for name in hash_names}:
            raise ResumeStateError(error_code, task_id=task_id)

    _publish_retry_file(
        workspace.new_output,
        canonical_output,
        workspace.directory / "publish-output.xlsx",
    )
    _publish_retry_file(
        workspace.new_trace,
        canonical_trace,
        workspace.directory / "publish-trace.jsonl",
    )
    updated_records = [
        transaction.new_prediction if record["id"] == task_id else record for record in records
    ]
    publish_predictions = workspace.directory / "publish-predictions.jsonl"
    atomic_write_jsonl(publish_predictions, updated_records)
    os.replace(publish_predictions, layout.predictions)
    _fsync_directory(layout.root)
    try:
        _remove_retry_workspace(layout, workspace)
    except ResumeStateError:
        # Publication is already durable.  Once cleanup has removed meta.json,
        # propagating its error could let a later checkpoint reinstall the old
        # in-memory failure record with no journal left to repair it.  Residual
        # markerless scratch is safe for the next resume (or runner finally) to
        # remove idempotently.
        if (workspace.directory / "meta.json").exists():
            raise
    return dict(transaction.new_prediction)


def recover_retry_transactions(
    layout: ArtifactLayout,
    tasks: list[DatasetTask],
    *,
    expected_model: str,
    expected_model_provenance: str | None = None,
) -> None:
    """Finish durable local retry commits before validating an ordinary resume."""

    root = _retry_root(layout)
    if not root.exists():
        return
    if root.is_symlink() or not root.is_dir():
        raise ResumeStateError("retry_root_unsafe")
    task_by_id = {task.id: task for task in tasks}
    if len(task_by_id) != len(tasks):
        raise ResumeStateError("retry_task_duplicate")
    try:
        entries = list(root.iterdir())
    except OSError as exc:
        raise ResumeStateError("retry_root_unreadable") from exc
    directories: dict[str, Path] = {}
    for entry in entries:
        if entry.is_symlink() or not entry.is_dir():
            raise ResumeStateError("retry_root_entry_unsafe")
        try:
            task_id = normalise_task_id(entry.name)
        except ArtifactError as exc:
            raise ResumeStateError("retry_task_id_unsafe") from exc
        if task_id not in task_by_id:
            raise ResumeStateError("retry_task_not_selected", task_id=task_id)
        directories[task_id] = entry

    for task in tasks:
        directory = directories.get(task.id)
        if directory is None:
            continue
        workspace = RetryWorkspace(
            task_id=task.id,
            directory=directory,
            new_output=directory / "new-output.xlsx",
            new_trace=directory / "new-trace.jsonl",
        )
        entries_by_name = _retry_directory_entries(directory)
        if "meta.json" not in entries_by_name:
            _remove_retry_workspace(layout, workspace)
            continue
        commit_retry_transaction(
            layout,
            tasks,
            task_id=task.id,
            expected_model=expected_model,
            expected_model_provenance=expected_model_provenance,
        )
    _remove_retry_root_if_empty(layout)


def _canonical_regular_entries(directory: Path, *, kind: str) -> dict[str, Path]:
    try:
        entries = list(directory.iterdir())
    except OSError as exc:
        raise ResumeStateError(f"{kind}_unreadable") from exc
    result: dict[str, Path] = {}
    for entry in entries:
        if entry.is_symlink() or not entry.is_file():
            raise ResumeStateError(f"{kind}_entry_unsafe")
        result[entry.name] = entry
    return result


def load_resume_state(
    layout: ArtifactLayout,
    tasks: list[DatasetTask],
    *,
    expected_model: str,
    expected_model_provenance: str | None = None,
) -> ResumeState:
    """Validate and return the task records that may safely be skipped.

    The selected ``tasks`` must be in manifest order.  Any uncheckpointed
    output or trace is treated as an orphan and left untouched so an operator
    can inspect it.  This fail-closed choice avoids guessing a prediction
    status after a process was terminated between atomic publications.
    """

    if not tasks:
        raise ResumeStateError("no_selected_tasks")
    if not isinstance(expected_model, str) or not expected_model:
        raise ValueError("expected_model must be a non-empty string")
    if expected_model_provenance is not None and (
        not isinstance(expected_model_provenance, str) or not expected_model_provenance
    ):
        raise ValueError("expected_model_provenance must be a non-empty string or None")
    require_resume_directory(layout.root)
    for path, code in (
        (layout.predictions, "predictions_missing"),
        (layout.outputs, "outputs_missing"),
        (layout.traces, "traces_missing"),
    ):
        if path.is_symlink() or not path.exists():
            raise ResumeStateError(code)
    if not layout.predictions.is_file():
        raise ResumeStateError("predictions_unsafe")
    if not layout.outputs.is_dir():
        raise ResumeStateError("outputs_unsafe")
    if not layout.traces.is_dir():
        raise ResumeStateError("traces_unsafe")

    task_by_id = {task.id: task for task in tasks}
    raw_records = _read_resume_jsonl(layout.predictions, code="predictions")
    predictions: dict[str, dict[str, Any]] = {}
    record_order: list[str] = []
    for raw_record in raw_records:
        task_id, record = _validate_resume_prediction(raw_record, task_by_id=task_by_id)
        if task_id in predictions:
            raise ResumeStateError("prediction_duplicate", task_id=task_id)
        predictions[task_id] = record
        record_order.append(task_id)

    expected_record_order = [task.id for task in tasks if task.id in predictions]
    if record_order != expected_record_order:
        raise ResumeStateError("prediction_order")

    output_entries = _canonical_regular_entries(layout.outputs, kind="outputs")
    trace_entries = _canonical_regular_entries(layout.traces, kind="traces")
    expected_output_names = {f"{task_id}.xlsx" for task_id in predictions}
    expected_trace_names = {f"{task_id}.jsonl" for task_id in predictions}
    if set(output_entries).difference(expected_output_names):
        raise ResumeStateError("orphan_output")
    if set(trace_entries).difference(expected_trace_names):
        raise ResumeStateError("orphan_trace")
    if expected_output_names.difference(output_entries):
        missing_name = next(iter(expected_output_names.difference(output_entries)))
        raise ResumeStateError("workbook_missing", task_id=missing_name[: -len(".xlsx")])
    if expected_trace_names.difference(trace_entries):
        missing_name = next(iter(expected_trace_names.difference(trace_entries)))
        raise ResumeStateError("trace_missing", task_id=missing_name[: -len(".jsonl")])

    successful_ids: set[str] = set()
    for task in tasks:
        prediction = predictions.get(task.id)
        if prediction is None:
            continue
        output_path = output_entries[f"{task.id}.xlsx"]
        if not _resume_workbook_is_readable(output_path):
            raise ResumeStateError("workbook_unreadable", task_id=task.id)
        status = prediction["status"]
        if status != "ok":
            try:
                fallback_matches = sha256_file(output_path) == sha256_file(task.init_xlsx)
            except OSError as exc:
                raise ResumeStateError("failure_fallback_unreadable", task_id=task.id) from exc
            if not fallback_matches:
                raise ResumeStateError("failure_fallback_hash", task_id=task.id)
        else:
            successful_ids.add(task.id)
        _validate_resume_trace(
            trace_entries[f"{task.id}.jsonl"],
            task_id=task.id,
            expected_model=expected_model,
            expected_model_provenance=expected_model_provenance,
            require_successful_call=status == "ok",
        )

    frozen_predictions = MappingProxyType(
        {task_id: MappingProxyType(dict(record)) for task_id, record in predictions.items()}
    )
    return ResumeState(
        predictions=frozen_predictions,
        completed_ids=frozenset(predictions),
        successful_ids=frozenset(successful_ids),
    )
