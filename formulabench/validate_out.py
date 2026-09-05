"""Read-only validation for a FormulaBench ``/out`` directory."""

from __future__ import annotations

import argparse
import json
import os
import stat
import zipfile
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Any

import openpyxl

from .artifacts import (
    ArtifactError,
    DatasetTask,
    confined_path,
    load_dataset_manifest,
    normalise_task_id,
    sha256_file,
)

_TRACE_REQUIRED_FIELDS = frozenset(
    {
        "step",
        "model",
        "prompt",
        "response",
        "input_tokens",
        "output_tokens",
        "latency_ms",
        "error",
    }
)
_MAX_JSONL_LINE_BYTES = 16 * 1024 * 1024
_MAX_WORKBOOK_MEMBERS = 20_000
_MAX_WORKBOOK_UNCOMPRESSED_BYTES = 1024 * 1024 * 1024


@dataclass(frozen=True, slots=True)
class ValidationIssue:
    code: str
    message: str
    task_id: str | None = None
    path: str | None = None

    def as_dict(self) -> dict[str, str]:
        result = {"code": self.code, "message": self.message}
        if self.task_id is not None:
            result["task_id"] = self.task_id
        if self.path is not None:
            result["path"] = self.path
        return result


@dataclass(frozen=True, slots=True)
class ValidationReport:
    expected_tasks: int
    prediction_records: int
    readable_workbooks: int
    valid_traces: int
    issues: tuple[ValidationIssue, ...]

    @property
    def ok(self) -> bool:
        return not self.issues

    def as_dict(self) -> dict[str, Any]:
        return {
            "ok": self.ok,
            "expected_tasks": self.expected_tasks,
            "prediction_records": self.prediction_records,
            "readable_workbooks": self.readable_workbooks,
            "valid_traces": self.valid_traces,
            "issues": [issue.as_dict() for issue in self.issues],
        }


def _issue(
    issues: list[ValidationIssue],
    code: str,
    message: str,
    *,
    task_id: str | None = None,
    path: str | None = None,
) -> None:
    issues.append(ValidationIssue(code, message, task_id=task_id, path=path))


def _jsonl_records(
    path: Path,
    issues: list[ValidationIssue],
    *,
    code_prefix: str,
    task_id: str | None = None,
) -> list[Any]:
    records: list[Any] = []
    try:
        with path.open("r", encoding="utf-8") as handle:
            for line_number, line in enumerate(handle, start=1):
                if not line.strip():
                    continue
                if len(line.encode("utf-8")) > _MAX_JSONL_LINE_BYTES:
                    _issue(
                        issues,
                        f"{code_prefix}_line_too_large",
                        f"JSONL line {line_number} exceeds the validation limit",
                        task_id=task_id,
                        path=path.name,
                    )
                    continue
                try:
                    records.append(json.loads(line))
                except json.JSONDecodeError:
                    _issue(
                        issues,
                        f"{code_prefix}_invalid_json",
                        f"JSONL line {line_number} is invalid",
                        task_id=task_id,
                        path=path.name,
                    )
    except (OSError, UnicodeError):
        _issue(
            issues,
            f"{code_prefix}_unreadable",
            "JSONL file cannot be read as UTF-8",
            task_id=task_id,
            path=path.name,
        )
    return records


def _resolve_required_file(
    root: Path,
    relative: str,
    issues: list[ValidationIssue],
    *,
    code: str,
    task_id: str | None = None,
) -> Path | None:
    try:
        path = confined_path(root, relative, must_exist=True, reject_symlinks=True)
    except ArtifactError:
        _issue(
            issues,
            code,
            "Required path is missing, unsafe or outside the output directory",
            task_id=task_id,
            path=relative,
        )
        return None
    if not path.is_file():
        _issue(
            issues,
            code,
            "Required path is not a regular file",
            task_id=task_id,
            path=relative,
        )
        return None
    return path


def _workbook_is_readable(path: Path) -> bool:
    if not zipfile.is_zipfile(path):
        return False
    workbook = None
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > _MAX_WORKBOOK_MEMBERS:
                return False
            if sum(member.file_size for member in members) > _MAX_WORKBOOK_UNCOMPRESSED_BYTES:
                return False
            for member in members:
                member_path = PurePosixPath(member.filename)
                if member_path.is_absolute() or ".." in member_path.parts:
                    return False
            if archive.testzip() is not None:
                return False
        workbook = openpyxl.load_workbook(
            path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
        # Force every worksheet XML stream to be parsed without ever saving or
        # recalculating the file.
        if not workbook.sheetnames:
            return False
        for worksheet in workbook.worksheets:
            for _ in worksheet.iter_rows(values_only=True):
                pass
        return True
    except Exception:
        return False
    finally:
        if workbook is not None:
            workbook.close()


def _valid_optional_non_negative_int(value: Any) -> bool:
    return value is None or (isinstance(value, int) and not isinstance(value, bool) and value >= 0)


def _validate_trace_records(
    records: list[Any],
    *,
    task_id: str,
    expected_model: str,
    issues: list[ValidationIssue],
    require_successful_call: bool,
) -> bool:
    start_issue_count = len(issues)
    for expected_step, record in enumerate(records, start=1):
        if not isinstance(record, dict):
            _issue(
                issues,
                "trace_record_type",
                "Every trace line must be a JSON object",
                task_id=task_id,
                path=f"traces/{task_id}.jsonl",
            )
            continue
        if missing := _TRACE_REQUIRED_FIELDS.difference(record):
            # Report the field names: they are schema constants, never model or
            # secret content.
            _issue(
                issues,
                "trace_missing_fields",
                f"Trace record is missing fields: {', '.join(sorted(missing))}",
                task_id=task_id,
                path=f"traces/{task_id}.jsonl",
            )
        step = record.get("step")
        if isinstance(step, bool) or not isinstance(step, int) or step != expected_step:
            _issue(
                issues,
                "trace_step_order",
                "Trace steps must be contiguous integers starting at 1",
                task_id=task_id,
                path=f"traces/{task_id}.jsonl",
            )
        if record.get("model") != expected_model:
            _issue(
                issues,
                "trace_model",
                "Trace model does not match the fixed submission model",
                task_id=task_id,
                path=f"traces/{task_id}.jsonl",
            )
        for name in ("prompt", "response", "error"):
            if record.get(name) is not None and not isinstance(record.get(name), str):
                _issue(
                    issues,
                    "trace_field_type",
                    f"Trace field {name} must be a string or null",
                    task_id=task_id,
                    path=f"traces/{task_id}.jsonl",
                )
        for name in ("input_tokens", "output_tokens", "latency_ms"):
            if not _valid_optional_non_negative_int(record.get(name)):
                _issue(
                    issues,
                    "trace_metric_type",
                    f"Trace field {name} must be a non-negative integer or null",
                    task_id=task_id,
                    path=f"traces/{task_id}.jsonl",
                )
    if require_successful_call and not any(
        isinstance(record, dict)
        and record.get("error") is None
        and isinstance(record.get("response"), str)
        for record in records
    ):
        _issue(
            issues,
            "trace_no_successful_call",
            "An ok prediction must include at least one successful model-call record",
            task_id=task_id,
            path=f"traces/{task_id}.jsonl",
        )
    return len(issues) == start_issue_count


def _normalise_secrets(secret_values: Iterable[str | bytes]) -> tuple[bytes, ...]:
    needles: list[bytes] = []
    seen: set[bytes] = set()
    for value in secret_values:
        if isinstance(value, str):
            encoded = value.encode("utf-8")
        elif isinstance(value, bytes):
            encoded = value
        else:
            raise TypeError("secret values must be strings or bytes")
        # Very short environment values create unusable false positives.  API
        # credentials are substantially longer; eight bytes remains conservative.
        if len(encoded) < 8 or encoded in seen:
            continue
        seen.add(encoded)
        needles.append(encoded)
    return tuple(needles)


def _file_contains_secret(path: Path, needles: tuple[bytes, ...]) -> bool:
    if not needles:
        return False
    overlap_size = max(len(needle) for needle in needles) - 1
    overlap = b""
    with path.open("rb") as handle:
        while chunk := handle.read(1024 * 1024):
            window = overlap + chunk
            if any(needle in window for needle in needles):
                return True
            overlap = window[-overlap_size:] if overlap_size else b""
    return False


def _scan_for_secrets(
    root: Path,
    secret_values: Iterable[str | bytes],
    issues: list[ValidationIssue],
) -> None:
    needles = _normalise_secrets(secret_values)
    if not needles:
        return
    for current_root, directory_names, file_names in os.walk(root, followlinks=False):
        current = Path(current_root)
        for directory_name in list(directory_names):
            candidate = current / directory_name
            if candidate.is_symlink():
                directory_names.remove(directory_name)
                _issue(
                    issues,
                    "symlink_artifact",
                    "Output artefacts must not contain symbolic links",
                    path=candidate.relative_to(root).as_posix(),
                )
        for file_name in file_names:
            candidate = current / file_name
            relative = candidate.relative_to(root).as_posix()
            try:
                mode = candidate.lstat().st_mode
            except OSError:
                _issue(
                    issues,
                    "secret_scan_error",
                    "An output artefact could not be inspected",
                    path=relative,
                )
                continue
            if stat.S_ISLNK(mode):
                _issue(
                    issues,
                    "symlink_artifact",
                    "Output artefacts must not contain symbolic links",
                    path=relative,
                )
                continue
            if not stat.S_ISREG(mode):
                continue
            try:
                leaked = _file_contains_secret(candidate, needles)
            except OSError:
                _issue(
                    issues,
                    "secret_scan_error",
                    "An output artefact could not be scanned",
                    path=relative,
                )
                continue
            if leaked:
                _issue(
                    issues,
                    "secret_leak",
                    "An output artefact contains a configured secret value",
                    path=relative,
                )


def _sanitise_issues(
    issues: Iterable[ValidationIssue], secret_values: Iterable[str | bytes]
) -> tuple[ValidationIssue, ...]:
    secret_texts: list[str] = []
    for value in secret_values:
        if isinstance(value, str) and len(value.encode("utf-8")) >= 8:
            secret_texts.append(value)
        elif isinstance(value, bytes) and len(value) >= 8:
            try:
                secret_texts.append(value.decode("utf-8"))
            except UnicodeDecodeError:
                continue

    def sanitise(value: str | None) -> str | None:
        if value is None:
            return None
        for secret in secret_texts:
            value = value.replace(secret, "[configured-secret]")
        return value

    return tuple(
        ValidationIssue(
            code=issue.code,
            message=sanitise(issue.message) or "",
            task_id=sanitise(issue.task_id),
            path=sanitise(issue.path),
        )
        for issue in issues
    )


def validate_output(
    dataset_dir: str | os.PathLike[str],
    out_dir: str | os.PathLike[str],
    *,
    expected_model: str,
    secret_values: Iterable[str | bytes] = (),
    require_traces: bool = True,
) -> ValidationReport:
    """Validate an output tree without modifying or recalculating workbooks."""
    if not isinstance(expected_model, str) or not expected_model:
        raise ValueError("expected_model must be a non-empty string")
    secret_values = tuple(secret_values)
    tasks = load_dataset_manifest(dataset_dir)
    task_by_id: dict[str, DatasetTask] = {task.id: task for task in tasks}
    expected_ids = set(task_by_id)
    issues: list[ValidationIssue] = []
    prediction_records = 0
    readable_workbooks = 0
    valid_traces = 0

    output_candidate = Path(out_dir)
    if output_candidate.is_symlink():
        _issue(
            issues,
            "output_root_symlink",
            "Output directory must not be a symbolic link",
        )
    try:
        output_root = output_candidate.resolve(strict=True)
    except (FileNotFoundError, RuntimeError):
        _issue(issues, "output_missing", "Output directory does not exist")
        return ValidationReport(len(tasks), 0, 0, 0, tuple(issues))
    if not output_root.is_dir():
        _issue(issues, "output_not_directory", "Output path is not a directory")
        return ValidationReport(len(tasks), 0, 0, 0, tuple(issues))

    predictions_path = _resolve_required_file(
        output_root,
        "predictions.jsonl",
        issues,
        code="predictions_missing",
    )
    records: list[Any] = []
    if predictions_path is not None:
        records = _jsonl_records(predictions_path, issues, code_prefix="predictions")
    prediction_records = len(records)

    predictions: dict[str, dict[str, Any]] = {}
    for record in records:
        if not isinstance(record, dict):
            _issue(
                issues,
                "prediction_record_type",
                "Every prediction line must be a JSON object",
                path="predictions.jsonl",
            )
            continue
        raw_id = record.get("id")
        if not isinstance(raw_id, str):
            _issue(
                issues,
                "prediction_id_type",
                "Prediction id must be a string",
                path="predictions.jsonl",
            )
            continue
        try:
            task_id = normalise_task_id(raw_id)
        except ArtifactError:
            _issue(
                issues,
                "prediction_id_unsafe",
                "Prediction id is not filename-safe",
                path="predictions.jsonl",
            )
            continue
        if task_id in predictions:
            _issue(
                issues,
                "prediction_duplicate",
                "Task has more than one prediction record",
                task_id=task_id,
                path="predictions.jsonl",
            )
            continue
        predictions[task_id] = record

    for missing_id in sorted(expected_ids.difference(predictions)):
        _issue(
            issues,
            "prediction_missing",
            "Dataset task has no prediction record",
            task_id=missing_id,
            path="predictions.jsonl",
        )
    for extra_id in sorted(set(predictions).difference(expected_ids)):
        _issue(
            issues,
            "prediction_extra",
            "Prediction id is not present in the dataset",
            task_id=extra_id,
            path="predictions.jsonl",
        )

    seen_output_paths: dict[Path, str] = {}
    for task_id in sorted(expected_ids.intersection(predictions)):
        task = task_by_id[task_id]
        prediction = predictions[task_id]
        status_value = prediction.get("status")
        if not isinstance(status_value, str) or not status_value:
            _issue(
                issues,
                "prediction_status",
                "Prediction status must be a non-empty string",
                task_id=task_id,
                path="predictions.jsonl",
            )
        output_value = prediction.get("output")
        if not isinstance(output_value, str) or not output_value:
            _issue(
                issues,
                "prediction_output",
                "Prediction output must be a non-empty relative path",
                task_id=task_id,
                path="predictions.jsonl",
            )
            continue
        expected_relative = f"outputs/{task_id}.xlsx"
        if output_value != expected_relative:
            _issue(
                issues,
                "prediction_output_layout",
                "Prediction output must use the canonical task-derived path",
                task_id=task_id,
                path="predictions.jsonl",
            )
        output_path = _resolve_required_file(
            output_root,
            output_value,
            issues,
            code="workbook_missing_or_unsafe",
            task_id=task_id,
        )
        if output_path is None:
            continue
        if output_path.suffix.lower() != ".xlsx":
            _issue(
                issues,
                "workbook_extension",
                "Prediction output must be an .xlsx workbook",
                task_id=task_id,
                path=output_value,
            )
        prior_task = seen_output_paths.get(output_path)
        if prior_task is not None and prior_task != task_id:
            _issue(
                issues,
                "workbook_reused",
                "Different tasks must not reference the same output workbook",
                task_id=task_id,
                path=output_value,
            )
        seen_output_paths[output_path] = task_id

        if _workbook_is_readable(output_path):
            readable_workbooks += 1
        else:
            _issue(
                issues,
                "workbook_unreadable",
                "Output is not a readable OOXML workbook",
                task_id=task_id,
                path=output_value,
            )

        # Every non-success record must be a byte-exact, deterministic fallback.
        # This makes failures scorable without partially edited workbooks.
        if status_value != "ok":
            try:
                fallback_matches = sha256_file(output_path) == sha256_file(task.init_xlsx)
            except OSError:
                fallback_matches = False
            if not fallback_matches:
                _issue(
                    issues,
                    "failure_fallback_hash",
                    "Failed task output is not an exact copy of its initial workbook",
                    task_id=task_id,
                    path=output_value,
                )

    run_log = _resolve_required_file(output_root, "run.log", issues, code="run_log_missing")
    del run_log  # Presence and confinement are the complete structural check.

    try:
        traces_dir = confined_path(output_root, "traces", must_exist=True, reject_symlinks=True)
    except ArtifactError:
        traces_dir = None
        if require_traces:
            _issue(
                issues,
                "traces_missing",
                "Trace directory is missing or unsafe",
                path="traces",
            )
    if traces_dir is not None and not traces_dir.is_dir():
        _issue(
            issues,
            "traces_not_directory",
            "Trace path is not a directory",
            path="traces",
        )
        traces_dir = None

    if traces_dir is not None:
        expected_trace_names = {f"{task_id}.jsonl" for task_id in expected_ids}
        try:
            actual_trace_names = {
                entry.name for entry in traces_dir.iterdir() if entry.name.endswith(".jsonl")
            }
        except OSError:
            actual_trace_names = set()
            _issue(
                issues,
                "traces_unreadable",
                "Trace directory cannot be enumerated",
                path="traces",
            )
        for extra_name in sorted(actual_trace_names.difference(expected_trace_names)):
            _issue(
                issues,
                "trace_extra",
                "Trace filename is not derived from a dataset task id",
                path=f"traces/{extra_name}",
            )

        for task_id in sorted(expected_ids):
            relative_trace = f"traces/{task_id}.jsonl"
            try:
                trace_path = confined_path(
                    output_root,
                    relative_trace,
                    must_exist=True,
                    reject_symlinks=True,
                )
                if not trace_path.is_file():
                    raise ArtifactError("trace is not a regular file")
            except ArtifactError:
                trace_path = None
                if require_traces:
                    _issue(
                        issues,
                        "trace_missing",
                        "Required trace is missing, unsafe or not a regular file",
                        task_id=task_id,
                        path=relative_trace,
                    )
            if trace_path is None:
                continue
            before_read = len(issues)
            trace_records = _jsonl_records(
                trace_path,
                issues,
                code_prefix="trace",
                task_id=task_id,
            )
            parse_was_valid = len(issues) == before_read
            structure_valid = _validate_trace_records(
                trace_records,
                task_id=task_id,
                expected_model=expected_model,
                issues=issues,
                require_successful_call=(
                    require_traces and predictions.get(task_id, {}).get("status") == "ok"
                ),
            )
            if parse_was_valid and structure_valid:
                valid_traces += 1

    _scan_for_secrets(output_root, secret_values, issues)
    return ValidationReport(
        expected_tasks=len(tasks),
        prediction_records=prediction_records,
        readable_workbooks=readable_workbooks,
        valid_traces=valid_traces,
        issues=_sanitise_issues(issues, secret_values),
    )


def parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Validate FormulaBench output artefacts.")
    parser.add_argument("--dataset-dir", required=True)
    parser.add_argument("--out-dir", required=True)
    parser.add_argument("--expected-model", required=True)
    parser.add_argument(
        "--secret-env",
        action="append",
        default=[],
        metavar="NAME",
        help="scan for the value of an environment variable without printing it",
    )
    parser.add_argument(
        "--optional-traces",
        action="store_true",
        help="make trace files optional (existing trace records are still validated)",
    )
    return parser.parse_args(argv)


def main(argv: Sequence[str] | None = None) -> int:
    args = parse_args(argv)
    secrets = [os.environ[name] for name in args.secret_env if os.environ.get(name)]
    try:
        report = validate_output(
            args.dataset_dir,
            args.out_dir,
            expected_model=args.expected_model,
            secret_values=secrets,
            require_traces=not args.optional_traces,
        )
    except (ArtifactError, OSError, ValueError) as exc:
        print(
            json.dumps(
                {
                    "ok": False,
                    "error": "validation could not start",
                    "type": type(exc).__name__,
                },
                sort_keys=True,
            )
        )
        return 2
    print(json.dumps(report.as_dict(), indent=2, sort_keys=True))
    return 0 if report.ok else 1


if __name__ == "__main__":
    raise SystemExit(main())
