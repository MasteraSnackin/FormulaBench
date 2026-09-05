"""Strict target and output contracts for SpreadsheetBench workbooks.

The official evaluator identifies targets with ``answer_position`` and
``answer_sheet``.  This module deliberately delegates the answer-position
grammar to :func:`sb.parse_answer_position`, then makes the implicit sheet
selection and response-validation rules explicit.

No function in this module reads a golden workbook.  ``write_response_atomic``
always starts from the task's pristine input workbook and installs either a
fully validated result or a byte-identical copy of that input.
"""

from __future__ import annotations

import datetime as dt
import json
import math
import os
import re
import shutil
import tempfile
from collections.abc import Iterable, Mapping
from dataclasses import dataclass
from enum import StrEnum
from pathlib import Path
from typing import Literal, TypeAlias

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.compat.strings import safe_string
from openpyxl.formula.tokenizer import Tokenizer, TokenizerError
from openpyxl.formula.translate import Translator, TranslatorError
from openpyxl.utils.cell import column_index_from_string, get_column_letter
from openpyxl.utils.datetime import from_excel
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula
from pydantic import (
    BaseModel,
    ConfigDict,
    Field,
    StrictBool,
    StrictFloat,
    StrictInt,
    StrictStr,
    field_validator,
)
from pydantic import ValidationError as PydanticValidationError

from sb import parse_answer_position

EXCEL_MAX_ROW = 1_048_576
EXCEL_MAX_COLUMN = 16_384
EXCEL_MAX_CELL_CHARACTERS = 32_767
EXCEL_MAX_FORMULA_CHARACTERS = 8_192
# The largest public answer has 104,110 cells.  This permits legitimate results
# of more than twice that size while bounding model-controlled expansion.
MAX_EXPANDED_RESPONSE_CELLS = 250_000
# Expanded strings are also bounded independently: a legal 32,767-character
# scalar repeated 250,000 times would otherwise require multiple gigabytes.
MAX_EXPANDED_RESPONSE_CHARACTERS = 128 * 1024 * 1024

_CELL_RE = re.compile(r"^\$?([A-Za-z]{1,3})\$?([1-9][0-9]*)$")
_COLUMN_RANGE_RE = re.compile(r"^\$?([A-Za-z]{1,3}):\$?([A-Za-z]{1,3})$")
_FINITE_RANGE_RE = re.compile(
    r"^\$?([A-Za-z]{1,3})\$?([1-9][0-9]*):"
    r"\$?([A-Za-z]{1,3})\$?([1-9][0-9]*)$"
)
_INVALID_SHEET_CHARS_RE = re.compile(r"[\\*?:/\[\]]")
_DATE_VALUE_RE = re.compile(r"^[0-9]{4}-[0-9]{2}-[0-9]{2}$")
_DATETIME_VALUE_RE = re.compile(
    r"^[0-9]{4}-[0-9]{2}-[0-9]{2}T"
    r"[0-9]{2}:[0-9]{2}:[0-9]{2}(?:\.[0-9]{1,3}0{0,3})?$"
)
_UNSAFE_FORMULA_FUNCTIONS = frozenset(
    {
        "CALL",
        "CUBEKPIMEMBER",
        "CUBEMEMBER",
        "CUBEMEMBERPROPERTY",
        "CUBERANKEDMEMBER",
        "CUBESET",
        "CUBESETCOUNT",
        "CUBEVALUE",
        "DDE",
        "EVALUATE",
        "EXEC",
        "FOPEN",
        "FREAD",
        "FREADLN",
        "FWRITE",
        "FWRITELN",
        "HYPERLINK",
        "IMAGE",
        "IMPORTDATA",
        "IMPORTFEED",
        "IMPORTHTML",
        "IMPORTRANGE",
        "IMPORTXML",
        "PY",
        "REGISTER",
        "REGISTER.ID",
        "RTD",
        "RUN",
        "SQL.REQUEST",
        "STOCKHISTORY",
        "WEBSERVICE",
    }
)
_EXTERNAL_URI_RE = re.compile(r"(?i)(?:https?|ftp|file|smb)://")
_EXTERNAL_DRIVE_PATH_RE = re.compile(r"(?i)(?:^|[' ])(?:[A-Z]:\\)")
_EXTERNAL_WORKBOOK_RE = re.compile(r"\[[^\]\r\n]+\][^!\r\n]*!")
_EXTERNAL_WORKBOOK_RAW_RE = re.compile(
    r"(?i)\[[^\]\r\n]+\.(?:xlsx?|xlsm|xlsb|xlam|xltx|xltm|ods|csv)\]"
    r"[^!+\-*/^&(),<>=;\r\n]*!"
)
_EXTERNAL_WORKBOOK_QUOTED_RE = re.compile(r"'(?:[^'\r\n]|'')*\[[^\]\r\n]+\](?:[^'\r\n]|'')*'!")
_EXTERNAL_WORKBOOK_BARE_RE = re.compile(r"(?i)(?:^|[=+\-*/^&,(<>:; ])\[[^\]\r\n]+\][A-Z0-9_.]+!")
_DDE_REFERENCE_RE = re.compile(r"\|(?:'[^'\r\n]{0,255}'|[^!\r\n]{0,255})!")
_FORMULA_FUNCTION_CALL_RE = re.compile(
    r"(?i)(?<![A-Z0-9_.])"
    r"(?P<name>@?(?:(?:_XLFN|_XLWS)\.)*[A-Z][A-Z0-9._]*)\s*\("
)


class ExcelDateValue(BaseModel):
    """An exact, timezone-free ISO calendar date from model JSON."""

    model_config = ConfigDict(extra="forbid", strict=True, frozen=True)

    type: Literal["date"]
    value: StrictStr

    @field_validator("value")
    @classmethod
    def validate_date(cls, value: str) -> str:
        if _DATE_VALUE_RE.fullmatch(value) is None:
            raise ValueError("date value must use exact YYYY-MM-DD syntax")
        try:
            dt.date.fromisoformat(value)
        except ValueError as exc:
            raise ValueError("date value must be a valid Gregorian date") from exc
        return value

    def materialise(self) -> dt.datetime:
        # The supplied evaluator recognises native Excel dates as datetimes.
        # Materialise a date at midnight rather than serialising an ISO date
        # string, which would reload as datetime.date and fail that contract.
        return dt.datetime.combine(dt.date.fromisoformat(self.value), dt.time())


class ExcelDateTimeValue(BaseModel):
    """An exact, timezone-free ISO local date and time from model JSON."""

    model_config = ConfigDict(extra="forbid", strict=True, frozen=True)

    type: Literal["datetime"]
    value: StrictStr

    @field_validator("value")
    @classmethod
    def validate_datetime(cls, value: str) -> str:
        if _DATETIME_VALUE_RE.fullmatch(value) is None:
            raise ValueError(
                "datetime value must use exact YYYY-MM-DDTHH:MM:SS[.fraction] syntax "
                "without a timezone"
            )
        try:
            parsed = dt.datetime.fromisoformat(value)
        except ValueError as exc:
            raise ValueError("datetime value must be a valid local date and time") from exc
        if parsed.tzinfo is not None:  # Defensive; the exact pattern excludes offsets.
            raise ValueError("datetime value must not include a timezone")
        if parsed.microsecond % 1_000:
            raise ValueError("datetime value must use exact millisecond precision")
        return value

    def materialise(self) -> dt.datetime:
        return dt.datetime.fromisoformat(self.value)


ScalarValue: TypeAlias = (
    StrictStr | StrictInt | StrictFloat | StrictBool | None | ExcelDateValue | ExcelDateTimeValue
)
MaterialisedValue: TypeAlias = str | int | float | bool | None | dt.datetime
PreservedCellSignature: TypeAlias = tuple[object, ...]


class FailureCode(StrEnum):
    """Stable, machine-readable reasons a response was not materialised."""

    INVALID_RESPONSE_JSON = "invalid_response_json"
    INVALID_RESPONSE_SCHEMA = "invalid_response_schema"
    INVALID_TARGET_METADATA = "invalid_target_metadata"
    INVALID_TARGET_COORDINATE = "invalid_target_coordinate"
    INVALID_TARGET_SHEET = "invalid_target_sheet"
    DUPLICATE_CELLS = "duplicate_cells"
    MISSING_CELLS = "missing_cells"
    EXTRA_CELLS = "extra_cells"
    DYNAMIC_RANGE_EMPTY = "dynamic_range_empty"
    DUPLICATE_UNMERGE_RANGES = "duplicate_unmerge_ranges"
    EXTRA_UNMERGE_RANGES = "extra_unmerge_ranges"
    EXPANSION_LIMIT_EXCEEDED = "expansion_limit_exceeded"
    MODEL_CALL_FAILED = "model_call_failed"
    CONTEXT_BUILD_FAILED = "context_build_failed"
    INPUT_LOAD_FAILED = "input_load_failed"
    INVALID_INPUT_WORKBOOK = "invalid_input_workbook"
    INVALID_OUTPUT_PATH = "invalid_output_path"
    UNSAFE_FORMULA = "unsafe_formula"
    WORKBOOK_WRITE_FAILED = "workbook_write_failed"


class TargetProvenance(StrEnum):
    """Where a target sheet name came from in the task metadata."""

    EXPLICIT = "explicit"
    ANSWER_SHEET = "answer_sheet"


class SheetResolution(StrEnum):
    """How a requested target sheet maps to the input workbook."""

    EXISTING = "existing"
    CREATED = "created"
    ACTIVE_FALLBACK = "active_fallback"


class ContractError(ValueError):
    """A contract error with a stable code suitable for traces."""

    def __init__(self, code: FailureCode, message: str) -> None:
        super().__init__(message)
        self.code = code


class ResponseCell(BaseModel):
    """One sheet-qualified value returned by the model."""

    model_config = ConfigDict(extra="forbid", strict=True)

    sheet: StrictStr
    cell: StrictStr
    value: ScalarValue

    @field_validator("sheet")
    @classmethod
    def validate_sheet(cls, value: str) -> str:
        if not value:
            raise ValueError("sheet must not be empty")
        return value

    @field_validator("cell")
    @classmethod
    def normalise_cell(cls, value: str) -> str:
        try:
            return normalise_coordinate(value)
        except ContractError as exc:
            raise ValueError(str(exc)) from exc

    @field_validator("value")
    @classmethod
    def require_finite_number(cls, value: ScalarValue) -> ScalarValue:
        if isinstance(value, float) and not math.isfinite(value):
            raise ValueError("numeric values must be finite")
        if (
            isinstance(value, str)
            and value.startswith("=")
            and len(value) > EXCEL_MAX_FORMULA_CHARACTERS
        ):
            raise ValueError(
                f"formula values must not exceed {EXCEL_MAX_FORMULA_CHARACTERS} characters"
            )
        if isinstance(value, str) and len(value) > EXCEL_MAX_CELL_CHARACTERS:
            raise ValueError(
                f"string values must not exceed {EXCEL_MAX_CELL_CHARACTERS} characters"
            )
        return value


class ResponseFill(BaseModel):
    """A compact rectangular fill, anchored at its top-left cell."""

    model_config = ConfigDict(extra="forbid", strict=True)

    sheet: StrictStr
    range: StrictStr
    value: ScalarValue

    @field_validator("sheet")
    @classmethod
    def validate_sheet(cls, value: str) -> str:
        if not value:
            raise ValueError("sheet must not be empty")
        return value

    @field_validator("range")
    @classmethod
    def normalise_range(cls, value: str) -> str:
        if not isinstance(value, str) or _FINITE_RANGE_RE.fullmatch(value) is None:
            raise ValueError("fill range must be a finite rectangular A1 range")
        try:
            normalised, dynamic = _normalise_range(value)
        except ContractError as exc:
            raise ValueError(str(exc)) from exc
        if dynamic:
            raise ValueError("fill range must not be a whole-column range")
        return normalised

    @field_validator("value")
    @classmethod
    def validate_value(cls, value: ScalarValue) -> ScalarValue:
        return ResponseCell.require_finite_number(value)


class ResponsePreserveRange(BaseModel):
    """A finite target range whose existing contents must remain unchanged."""

    model_config = ConfigDict(extra="forbid", strict=True)

    sheet: StrictStr
    range: StrictStr

    @field_validator("sheet")
    @classmethod
    def validate_sheet(cls, value: str) -> str:
        if not value:
            raise ValueError("sheet must not be empty")
        return value

    @field_validator("range")
    @classmethod
    def normalise_range(cls, value: str) -> str:
        if not isinstance(value, str) or _FINITE_RANGE_RE.fullmatch(value) is None:
            raise ValueError("preserve range must be a finite rectangular A1 range")
        try:
            normalised, dynamic = _normalise_range(value)
        except ContractError as exc:
            raise ValueError(str(exc)) from exc
        if dynamic:
            raise ValueError("preserve range must not be a whole-column range")
        return normalised


class ResponseUnmergeRange(BaseModel):
    """One exact, existing merged range the model asks to unmerge."""

    model_config = ConfigDict(extra="forbid", strict=True)

    sheet: StrictStr
    range: StrictStr

    @field_validator("sheet")
    @classmethod
    def validate_sheet(cls, value: str) -> str:
        if not value:
            raise ValueError("sheet must not be empty")
        return value

    @field_validator("range")
    @classmethod
    def normalise_range(cls, value: str) -> str:
        if not isinstance(value, str) or _FINITE_RANGE_RE.fullmatch(value) is None:
            raise ValueError("unmerge range must be a finite rectangular A1 range")
        try:
            normalised, dynamic = _normalise_range(value)
        except ContractError as exc:
            raise ValueError(str(exc)) from exc
        if dynamic:
            raise ValueError("unmerge range must not be a whole-column range")
        return normalised


class SpreadsheetResponse(BaseModel):
    """The complete structured response for a single task."""

    model_config = ConfigDict(extra="forbid", strict=True)

    cells: list[ResponseCell]
    fills: list[ResponseFill] = Field(default_factory=list)
    preserve_ranges: list[ResponsePreserveRange] = Field(default_factory=list)
    unmerge_ranges: list[ResponseUnmergeRange] = Field(default_factory=list)


@dataclass(frozen=True)
class CellKey:
    """A cell identity; the sheet is part of the key."""

    sheet: str
    cell: str

    def as_text(self) -> str:
        return f"{self.sheet}!{self.cell}"


@dataclass(frozen=True)
class RangeKey:
    """A sheet-qualified finite A1 range."""

    sheet: str
    cell_range: str

    def as_text(self) -> str:
        return f"{self.sheet}!{self.cell_range}"


@dataclass(frozen=True)
class MaterialisedCell:
    """A validated response cell after compact fills have been expanded."""

    sheet: str
    cell: str
    value: MaterialisedValue


@dataclass(frozen=True)
class TargetRange:
    """One normalised answer range and its sheet-resolution provenance."""

    sheet: str
    requested_sheet: str | None
    cell_range: str
    provenance: TargetProvenance
    resolution: SheetResolution
    dynamic: bool


@dataclass(frozen=True)
class TargetContract:
    """All target ranges for one task, in metadata order."""

    ranges: tuple[TargetRange, ...]
    active_sheet: str


@dataclass(frozen=True)
class DynamicExtent:
    """The candidate worksheet's last row used for a whole-column target."""

    sheet: str
    cell_range: str
    max_row: int


@dataclass(frozen=True)
class ValidationResult:
    """All set-level response defects, reported together."""

    expected: tuple[CellKey, ...]
    duplicates: tuple[CellKey, ...]
    missing: tuple[CellKey, ...]
    extra: tuple[CellKey, ...]
    dynamic_extents: tuple[DynamicExtent, ...]
    unmerge_duplicates: tuple[RangeKey, ...]
    unmerge_extras: tuple[RangeKey, ...]
    materialised_cells: tuple[MaterialisedCell, ...]
    preserved_cells: tuple[CellKey, ...]
    failure_codes: tuple[FailureCode, ...]

    @property
    def ok(self) -> bool:
        return not self.failure_codes


@dataclass(frozen=True)
class WriteResult:
    """Outcome of installing a response into an output workbook."""

    output_path: Path
    success: bool
    failure_codes: tuple[FailureCode, ...]
    validation: ValidationResult | None = None
    error: str | None = None

    @property
    def status(self) -> str:
        if self.success:
            return "ok"
        joined = ",".join(code.value for code in self.failure_codes)
        return f"error:{joined}"


def normalise_coordinate(coordinate: str) -> str:
    """Return an absolute-marker-free, uppercase Excel A1 coordinate."""

    if not isinstance(coordinate, str):
        raise ContractError(FailureCode.INVALID_TARGET_COORDINATE, "cell must be a string")
    match = _CELL_RE.fullmatch(coordinate)
    if not match:
        raise ContractError(
            FailureCode.INVALID_TARGET_COORDINATE,
            f"invalid cell coordinate: {coordinate!r}",
        )
    column, row_text = match.groups()
    column_index = column_index_from_string(column.upper())
    row = int(row_text)
    if column_index > EXCEL_MAX_COLUMN or row > EXCEL_MAX_ROW:
        raise ContractError(
            FailureCode.INVALID_TARGET_COORDINATE,
            f"cell coordinate outside Excel limits: {coordinate!r}",
        )
    return f"{get_column_letter(column_index)}{row}"


def _normalise_range(cell_range: str) -> tuple[str, bool]:
    if not isinstance(cell_range, str):
        raise ContractError(
            FailureCode.INVALID_TARGET_COORDINATE,
            "answer range must be a string",
        )

    column_match = _COLUMN_RANGE_RE.fullmatch(cell_range)
    if column_match:
        start, end = (column_index_from_string(value.upper()) for value in column_match.groups())
        if start > end or end > EXCEL_MAX_COLUMN:
            raise ContractError(
                FailureCode.INVALID_TARGET_COORDINATE,
                f"invalid whole-column range: {cell_range!r}",
            )
        return f"{get_column_letter(start)}:{get_column_letter(end)}", True

    finite_match = _FINITE_RANGE_RE.fullmatch(cell_range)
    if finite_match:
        start_column, start_row_text, end_column, end_row_text = finite_match.groups()
        start_column_index = column_index_from_string(start_column.upper())
        end_column_index = column_index_from_string(end_column.upper())
        start_row = int(start_row_text)
        end_row = int(end_row_text)
        if (
            start_column_index > end_column_index
            or start_row > end_row
            or end_column_index > EXCEL_MAX_COLUMN
            or end_row > EXCEL_MAX_ROW
        ):
            raise ContractError(
                FailureCode.INVALID_TARGET_COORDINATE,
                f"invalid finite range: {cell_range!r}",
            )
        return (
            f"{get_column_letter(start_column_index)}{start_row}:"
            f"{get_column_letter(end_column_index)}{end_row}",
            False,
        )

    return normalise_coordinate(cell_range), False


def _validate_creatable_sheet_name(name: str, existing_sheets: Iterable[str]) -> None:
    if not name or len(name) > 31 or _INVALID_SHEET_CHARS_RE.search(name):
        raise ContractError(
            FailureCode.INVALID_TARGET_SHEET,
            f"invalid explicit target sheet: {name!r}",
        )
    casefolded = {sheet.casefold() for sheet in existing_sheets}
    if name.casefold() in casefolded:
        raise ContractError(
            FailureCode.INVALID_TARGET_SHEET,
            f"explicit target sheet differs only by case from an existing sheet: {name!r}",
        )


def build_target_contract(
    task: Mapping[str, object], workbook: openpyxl.Workbook
) -> TargetContract:
    """Resolve official task metadata without changing ``workbook``.

    Explicit sheet names use that exact sheet and are created later if absent.
    Unqualified ranges use ``answer_sheet`` only when it exists exactly;
    otherwise they target the workbook's active sheet, matching the shipped
    SpreadsheetBench helper.
    """

    answer_position = task.get("answer_position")
    if not isinstance(answer_position, str) or not answer_position:
        raise ContractError(
            FailureCode.INVALID_TARGET_METADATA,
            "task answer_position must be a non-empty string",
        )
    if not workbook.sheetnames:
        raise ContractError(
            FailureCode.INVALID_INPUT_WORKBOOK,
            "input workbook contains no worksheets",
        )

    try:
        parsed_ranges = parse_answer_position(answer_position)
    except Exception as exc:
        raise ContractError(
            FailureCode.INVALID_TARGET_METADATA,
            f"could not parse answer_position: {exc}",
        ) from exc
    if not parsed_ranges:
        raise ContractError(
            FailureCode.INVALID_TARGET_METADATA,
            "answer_position contains no ranges",
        )

    answer_sheet = task.get("answer_sheet")
    if answer_sheet is not None and not isinstance(answer_sheet, str):
        raise ContractError(
            FailureCode.INVALID_TARGET_METADATA,
            "task answer_sheet must be a string or null",
        )

    active_sheet = workbook.active.title
    existing_sheets = tuple(workbook.sheetnames)
    reserved_sheets = list(existing_sheets)
    ranges: list[TargetRange] = []
    for explicit_sheet, raw_range in parsed_ranges:
        cell_range, dynamic = _normalise_range(raw_range)
        if explicit_sheet is not None:
            if not explicit_sheet:
                raise ContractError(
                    FailureCode.INVALID_TARGET_SHEET,
                    "explicit target sheet must not be empty",
                )
            requested_sheet = explicit_sheet
            sheet = explicit_sheet
            provenance = TargetProvenance.EXPLICIT
            if sheet in existing_sheets:
                resolution = SheetResolution.EXISTING
            else:
                if sheet not in reserved_sheets:
                    _validate_creatable_sheet_name(sheet, reserved_sheets)
                    reserved_sheets.append(sheet)
                resolution = SheetResolution.CREATED
        else:
            requested_sheet = answer_sheet
            provenance = TargetProvenance.ANSWER_SHEET
            if answer_sheet and answer_sheet in existing_sheets:
                sheet = answer_sheet
                resolution = SheetResolution.EXISTING
            else:
                sheet = active_sheet
                resolution = SheetResolution.ACTIVE_FALLBACK

        ranges.append(
            TargetRange(
                sheet=sheet,
                requested_sheet=requested_sheet,
                cell_range=cell_range,
                provenance=provenance,
                resolution=resolution,
                dynamic=dynamic,
            )
        )
    return TargetContract(ranges=tuple(ranges), active_sheet=active_sheet)


def _reject_json_constant(value: str) -> None:
    raise ValueError(f"non-finite JSON number is not allowed: {value}")


def _strict_json_object(pairs: list[tuple[str, object]]) -> dict[str, object]:
    result: dict[str, object] = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"duplicate JSON key: {key!r}")
        result[key] = value
    return result


def parse_response(payload: str | bytes | bytearray | Mapping[str, object]) -> SpreadsheetResponse:
    """Parse an exact JSON response; prose, fences, coercion and extras fail."""

    if isinstance(payload, (str, bytes, bytearray)):
        try:
            data = json.loads(
                payload,
                object_pairs_hook=_strict_json_object,
                parse_constant=_reject_json_constant,
            )
        except (json.JSONDecodeError, UnicodeDecodeError, ValueError) as exc:
            raise ContractError(FailureCode.INVALID_RESPONSE_JSON, str(exc)) from exc
    elif isinstance(payload, Mapping):
        data = dict(payload)
    else:
        raise ContractError(
            FailureCode.INVALID_RESPONSE_SCHEMA,
            "response must be a JSON object or its encoded text",
        )

    try:
        return SpreadsheetResponse.model_validate(data, strict=True)
    except PydanticValidationError as exc:
        raise ContractError(FailureCode.INVALID_RESPONSE_SCHEMA, str(exc)) from exc


def _coordinate_parts(coordinate: str) -> tuple[int, int]:
    match = _CELL_RE.fullmatch(coordinate)
    if match is None:  # Coordinates in validated models should make this unreachable.
        raise ContractError(
            FailureCode.INVALID_TARGET_COORDINATE,
            f"invalid cell coordinate: {coordinate!r}",
        )
    column, row = match.groups()
    return column_index_from_string(column), int(row)


def _cell_sort_key(key: CellKey) -> tuple[str, int, int]:
    column, row = _coordinate_parts(key.cell)
    return key.sheet, row, column


def _range_sort_key(key: RangeKey) -> tuple[str, int, int, int, int]:
    min_column, min_row, max_column, max_row = _finite_range_bounds(key.cell_range)
    return key.sheet, min_row, min_column, max_row, max_column


def _finite_range_bounds(cell_range: str) -> tuple[int, int, int, int]:
    match = _FINITE_RANGE_RE.fullmatch(cell_range)
    if match is None:
        raise ContractError(
            FailureCode.INVALID_TARGET_COORDINATE,
            f"not a finite rectangular A1 range: {cell_range!r}",
        )
    start_column, start_row, end_column, end_row = match.groups()
    return (
        column_index_from_string(start_column),
        int(start_row),
        column_index_from_string(end_column),
        int(end_row),
    )


def _bounded_add(current: int, additional: int, *, label: str) -> int:
    total = current + additional
    if total > MAX_EXPANDED_RESPONSE_CELLS:
        raise ContractError(
            FailureCode.EXPANSION_LIMIT_EXCEEDED,
            f"{label} expands beyond {MAX_EXPANDED_RESPONSE_CELLS} cells",
        )
    return total


def _bounded_character_add(current: int, additional: int) -> int:
    total = current + additional
    if total > MAX_EXPANDED_RESPONSE_CHARACTERS:
        raise ContractError(
            FailureCode.EXPANSION_LIMIT_EXCEEDED,
            "expanded string and formula values exceed the character-work limit",
        )
    return total


def _materialise_value(value: ScalarValue) -> MaterialisedValue:
    if isinstance(value, (ExcelDateValue, ExcelDateTimeValue)):
        return value.materialise()
    return value


def _formula_parts(formula: str) -> tuple[str, tuple[str, ...]]:
    """Return formula syntax with text masked plus decoded text literals."""

    masked: list[str] = []
    literals: list[str] = []
    literal: list[str] = []
    index = 0
    in_text = False
    while index < len(formula):
        character = formula[index]
        if character != '"':
            if in_text:
                literal.append(character)
            masked.append(" " if in_text else character)
            index += 1
            continue
        if in_text and index + 1 < len(formula) and formula[index + 1] == '"':
            literal.append('"')
            masked.extend((" ", " "))
            index += 2
            continue
        if in_text:
            literals.append("".join(literal))
            literal = []
        in_text = not in_text
        masked.append(" ")
        index += 1
    if in_text:
        raise ContractError(
            FailureCode.UNSAFE_FORMULA,
            "formula contains an unterminated text literal",
        )
    return "".join(masked), tuple(literals)


def _normalise_formula_function(value: str) -> str:
    name = value.removesuffix("(").strip().upper().lstrip("@")
    while name.startswith(("_XLFN.", "_XLWS.")):
        name = name.partition(".")[2].lstrip("@")
    return name


def _formula_outside_single_quoted_identifiers(formula: str) -> str:
    """Mask quoted sheet/workbook identifiers for raw function scanning."""

    masked: list[str] = []
    index = 0
    in_identifier = False
    while index < len(formula):
        character = formula[index]
        if character != "'":
            masked.append(" " if in_identifier else character)
            index += 1
            continue
        if in_identifier and index + 1 < len(formula) and formula[index + 1] == "'":
            masked.extend((" ", " "))
            index += 2
            continue
        in_identifier = not in_identifier
        masked.append(" ")
        index += 1
    if in_identifier:
        raise ContractError(
            FailureCode.UNSAFE_FORMULA,
            "formula contains an unterminated quoted identifier",
        )
    return "".join(masked)


def _has_external_reference(value: str) -> bool:
    return bool(
        _EXTERNAL_URI_RE.search(value)
        or "\\\\" in value
        or _EXTERNAL_DRIVE_PATH_RE.search(value)
        or _EXTERNAL_WORKBOOK_RE.search(value)
    )


def _has_raw_external_reference(value: str) -> bool:
    """Detect high-confidence external syntax in an un-tokenised formula."""

    return bool(
        _EXTERNAL_URI_RE.search(value)
        or "\\\\" in value
        or _EXTERNAL_DRIVE_PATH_RE.search(value)
        or _EXTERNAL_WORKBOOK_RAW_RE.search(value)
        or _EXTERNAL_WORKBOOK_QUOTED_RE.search(value)
        or _EXTERNAL_WORKBOOK_BARE_RE.search(value)
    )


def _require_safe_formula(value: object) -> None:
    """Reject model formulas with explicit external or executable behaviour.

    This is a deliberately narrow static boundary, not an Excel sandbox. It
    preserves ordinary cross-sheet, structured-reference and dynamic-array
    formulas while blocking the external behaviours relevant to automatic
    recalculation. Unknown syntax is allowed after the raw DDE/external scan
    because openpyxl 3.1.5 cannot tokenise valid spill syntax such as ``A1#``.
    """

    if not isinstance(value, str) or not value.startswith("="):
        return
    code, text_literals = _formula_parts(value)
    if _DDE_REFERENCE_RE.search(code) or _has_raw_external_reference(code):
        raise ContractError(
            FailureCode.UNSAFE_FORMULA,
            "formula contains an external workbook, path, URI, or DDE reference",
        )
    function_code = _formula_outside_single_quoted_identifiers(code)
    raw_functions = {
        _normalise_formula_function(match.group("name"))
        for match in _FORMULA_FUNCTION_CALL_RE.finditer(function_code)
    }
    if raw_functions.intersection(_UNSAFE_FORMULA_FUNCTIONS):
        raise ContractError(
            FailureCode.UNSAFE_FORMULA,
            "formula calls an external-data, executable, or link function",
        )
    if "INDIRECT" in raw_functions and any(
        _has_external_reference(literal) for literal in text_literals
    ):
        raise ContractError(
            FailureCode.UNSAFE_FORMULA,
            "INDIRECT formula contains an external reference",
        )
    try:
        tokens = Tokenizer(value).items
    except (IndexError, TokenizerError, ValueError):
        return

    functions = {
        _normalise_formula_function(token.value)
        for token in tokens
        if token.type == "FUNC" and token.subtype == "OPEN"
    }
    if functions.intersection(_UNSAFE_FORMULA_FUNCTIONS):
        raise ContractError(
            FailureCode.UNSAFE_FORMULA,
            "formula calls an external-data, executable, or link function",
        )
    if "INDIRECT" in functions and any(
        token.type == "OPERAND" and token.subtype == "TEXT" and _has_external_reference(token.value)
        for token in tokens
    ):
        raise ContractError(
            FailureCode.UNSAFE_FORMULA,
            "INDIRECT formula contains an external reference",
        )
    if any(
        token.type == "OPERAND"
        and token.subtype == "RANGE"
        and _has_external_reference(token.value)
        for token in tokens
    ):
        raise ContractError(
            FailureCode.UNSAFE_FORMULA,
            "formula contains an external range reference",
        )


def _materialise_response_cells(
    response: SpreadsheetResponse,
) -> tuple[tuple[MaterialisedCell, ...], tuple[CellKey, ...]]:
    """Expand authored and preserved coverage under one bounded cell budget."""

    expanded_count = _bounded_add(0, len(response.cells), label="response")
    materialised: list[MaterialisedCell] = []
    for cell in response.cells:
        value = _materialise_value(cell.value)
        _require_safe_formula(value)
        materialised.append(MaterialisedCell(sheet=cell.sheet, cell=cell.cell, value=value))
    preserved: list[CellKey] = []
    expanded_characters = sum(
        len(cell.value) for cell in response.cells if isinstance(cell.value, str)
    )
    expanded_characters = _bounded_character_add(0, expanded_characters)
    for fill in response.fills:
        min_column, min_row, max_column, max_row = _finite_range_bounds(fill.range)
        area = (max_column - min_column + 1) * (max_row - min_row + 1)
        expanded_count = _bounded_add(expanded_count, area, label="response")
        origin = f"{get_column_letter(min_column)}{min_row}"
        translator: Translator | None = None
        if isinstance(fill.value, str):
            # Reserve the full repeated base size before entering either loop.
            # Formula translations adjust this reservation by their exact delta.
            expanded_characters = _bounded_character_add(
                expanded_characters,
                len(fill.value) * area,
            )
            if fill.value.startswith("="):
                try:
                    translator = Translator(fill.value, origin=origin)
                except (IndexError, TokenizerError, TranslatorError, ValueError) as exc:
                    raise ContractError(
                        FailureCode.INVALID_RESPONSE_SCHEMA,
                        f"could not parse fill formula at {fill.sheet}!{origin}",
                    ) from exc
        for row in range(min_row, max_row + 1):
            for column in range(min_column, max_column + 1):
                coordinate = f"{get_column_letter(column)}{row}"
                value = fill.value
                if translator is not None:
                    try:
                        value = translator.translate_formula(coordinate)
                    except (IndexError, TokenizerError, TranslatorError, ValueError) as exc:
                        raise ContractError(
                            FailureCode.INVALID_RESPONSE_SCHEMA,
                            f"could not translate fill formula at {fill.sheet}!{coordinate}",
                        ) from exc
                    if len(value) > EXCEL_MAX_FORMULA_CHARACTERS:
                        raise ContractError(
                            FailureCode.INVALID_RESPONSE_SCHEMA,
                            f"translated formula at {fill.sheet}!{coordinate} exceeds Excel limits",
                        )
                    expanded_characters = _bounded_character_add(
                        expanded_characters,
                        len(value) - len(fill.value),
                    )
                _require_safe_formula(value)
                materialised.append(
                    MaterialisedCell(
                        sheet=fill.sheet,
                        cell=coordinate,
                        value=_materialise_value(value),
                    )
                )
    for preserve in response.preserve_ranges:
        min_column, min_row, max_column, max_row = _finite_range_bounds(preserve.range)
        area = (max_column - min_column + 1) * (max_row - min_row + 1)
        expanded_count = _bounded_add(expanded_count, area, label="response")
        preserved.extend(
            CellKey(preserve.sheet, f"{get_column_letter(column)}{row}")
            for row in range(min_row, max_row + 1)
            for column in range(min_column, max_column + 1)
        )
    assert len(materialised) + len(preserved) == expanded_count
    return tuple(materialised), tuple(preserved)


def _range_bounds(target_range: TargetRange) -> tuple[int, int, int, int | None]:
    if target_range.dynamic:
        match = _COLUMN_RANGE_RE.fullmatch(target_range.cell_range)
        if match is None:
            raise AssertionError("dynamic target is not a whole-column range")
        start, end = (column_index_from_string(value) for value in match.groups())
        return start, 1, end, None

    if ":" not in target_range.cell_range:
        column, row = _coordinate_parts(target_range.cell_range)
        return column, row, column, row

    match = _FINITE_RANGE_RE.fullmatch(target_range.cell_range)
    if match is None:
        raise AssertionError("finite target is not a normalised A1 range")
    start_column, start_row, end_column, end_row = match.groups()
    return (
        column_index_from_string(start_column),
        int(start_row),
        column_index_from_string(end_column),
        int(end_row),
    )


def _expand_target(target_range: TargetRange, max_row: int | None = None) -> Iterable[CellKey]:
    min_column, min_row, max_column, finite_max_row = _range_bounds(target_range)
    last_row = finite_max_row if finite_max_row is not None else max_row
    if last_row is None:
        return ()
    return (
        CellKey(target_range.sheet, f"{get_column_letter(column)}{row}")
        for row in range(min_row, last_row + 1)
        for column in range(min_column, max_column + 1)
    )


def _target_area(target_range: TargetRange, *, max_row: int | None = None) -> int:
    min_column, min_row, max_column, finite_max_row = _range_bounds(target_range)
    last_row = finite_max_row if finite_max_row is not None else max_row
    if last_row is None:
        return 0
    return (max_column - min_column + 1) * (last_row - min_row + 1)


def _target_intersects_range(target_range: TargetRange, cell_range: str) -> bool:
    target_min_col, target_min_row, target_max_col, target_max_row = _range_bounds(target_range)
    range_min_col, range_min_row, range_max_col, range_max_row = _finite_range_bounds(cell_range)
    target_last_row = EXCEL_MAX_ROW if target_max_row is None else target_max_row
    return not (
        target_max_col < range_min_col
        or range_max_col < target_min_col
        or target_last_row < range_min_row
        or range_max_row < target_min_row
    )


def _validate_unmerge_ranges(
    contract: TargetContract,
    response: SpreadsheetResponse,
    workbook: openpyxl.Workbook | None,
) -> tuple[tuple[RangeKey, ...], tuple[RangeKey, ...]]:
    if not response.unmerge_ranges:
        return (), ()
    if workbook is None:
        raise ContractError(
            FailureCode.INVALID_INPUT_WORKBOOK,
            "input workbook is required to validate unmerge actions",
        )

    requested = [RangeKey(action.sheet, action.range) for action in response.unmerge_ranges]
    seen: set[RangeKey] = set()
    duplicate_set: set[RangeKey] = set()
    for key in requested:
        if key in seen:
            duplicate_set.add(key)
        seen.add(key)

    existing: set[RangeKey] = set()
    for worksheet in workbook.worksheets:
        for merged_range in worksheet.merged_cells.ranges:
            normalised, dynamic = _normalise_range(str(merged_range))
            if dynamic:  # pragma: no cover - openpyxl merged ranges are always finite.
                continue
            existing.add(RangeKey(worksheet.title, normalised))

    extra_set = {
        key
        for key in requested
        if key not in existing
        or not any(
            target.sheet == key.sheet and _target_intersects_range(target, key.cell_range)
            for target in contract.ranges
        )
    }
    return (
        tuple(sorted(duplicate_set, key=_range_sort_key)),
        tuple(sorted(extra_set, key=_range_sort_key)),
    )


def validate_response(
    contract: TargetContract,
    response: SpreadsheetResponse | Mapping[str, object] | str | bytes | bytearray,
    *,
    workbook: openpyxl.Workbook | None = None,
) -> ValidationResult:
    """Validate response coverage and return every set-level defect at once."""

    parsed = response if isinstance(response, SpreadsheetResponse) else parse_response(response)
    if parsed.preserve_ranges:
        if workbook is None:
            raise ContractError(
                FailureCode.INVALID_INPUT_WORKBOOK,
                "input workbook is required to validate preserve ranges",
            )
        created_target_sheets = {
            target.sheet
            for target in contract.ranges
            if target.resolution is SheetResolution.CREATED
        }
        invalid_created_sheets = sorted(
            {
                preserve.sheet
                for preserve in parsed.preserve_ranges
                if preserve.sheet in created_target_sheets
            }
        )
        if invalid_created_sheets:
            raise ContractError(
                FailureCode.INVALID_RESPONSE_SCHEMA,
                "preserve ranges cannot refer to newly created target sheets: "
                + ", ".join(repr(sheet) for sheet in invalid_created_sheets),
            )
        dynamic_target_sheets = {target.sheet for target in contract.ranges if target.dynamic}
        extending_dynamic_sheets = sorted(
            {
                preserve.sheet
                for preserve in parsed.preserve_ranges
                if preserve.sheet in dynamic_target_sheets
                and _finite_range_bounds(preserve.range)[3] > workbook[preserve.sheet].max_row
            }
        )
        if extending_dynamic_sheets:
            raise ContractError(
                FailureCode.INVALID_RESPONSE_SCHEMA,
                "preserve ranges cannot extend a dynamic target beyond its existing worksheet "
                "extent: " + ", ".join(repr(sheet) for sheet in extending_dynamic_sheets),
            )

    materialised, preserved = _materialise_response_cells(parsed)
    returned = [CellKey(cell.sheet, cell.cell) for cell in materialised]
    returned.extend(preserved)
    seen: set[CellKey] = set()
    duplicate_set: set[CellKey] = set()
    for key in returned:
        if key in seen:
            duplicate_set.add(key)
        seen.add(key)

    expected_set: set[CellKey] = set()
    dynamic_extents: list[DynamicExtent] = []
    expected_count = 0
    for target_range in contract.ranges:
        if not target_range.dynamic:
            expected_count = _bounded_add(
                expected_count,
                _target_area(target_range),
                label="answer target",
            )
            expected_set.update(_expand_target(target_range))
            continue

        if workbook is None:
            raise ContractError(
                FailureCode.INVALID_INPUT_WORKBOOK,
                "input workbook is required to validate dynamic answer ranges",
            )
        if target_range.sheet in workbook.sheetnames:
            initial_max_row = workbook[target_range.sheet].max_row
        elif target_range.resolution is SheetResolution.CREATED:
            # A newly created openpyxl worksheet reports a one-row extent, as
            # does the candidate worksheet inspected by the shipped evaluator.
            initial_max_row = 1
        else:  # A resolved existing/fallback target must still exist exactly.
            raise ContractError(
                FailureCode.INVALID_INPUT_WORKBOOK,
                f"resolved target sheet is absent: {target_range.sheet!r}",
            )
        returned_sheet_max_row = max(
            (_coordinate_parts(key.cell)[1] for key in returned if key.sheet == target_range.sheet),
            default=1,
        )
        max_row = max(initial_max_row, returned_sheet_max_row)
        dynamic_extents.append(
            DynamicExtent(
                sheet=target_range.sheet,
                cell_range=target_range.cell_range,
                max_row=max_row,
            )
        )
        expected_count = _bounded_add(
            expected_count,
            _target_area(target_range, max_row=max_row),
            label="answer target",
        )
        expected_set.update(_expand_target(target_range, max_row=max_row))

    returned_set = set(returned)
    duplicates = tuple(sorted(duplicate_set, key=_cell_sort_key))
    missing = tuple(sorted(expected_set - returned_set, key=_cell_sort_key))
    extra = tuple(sorted(returned_set - expected_set, key=_cell_sort_key))
    failure_codes: list[FailureCode] = []
    if duplicates:
        failure_codes.append(FailureCode.DUPLICATE_CELLS)
    if missing:
        failure_codes.append(FailureCode.MISSING_CELLS)
    if extra:
        failure_codes.append(FailureCode.EXTRA_CELLS)
    unmerge_duplicates, unmerge_extras = _validate_unmerge_ranges(
        contract,
        parsed,
        workbook,
    )
    if unmerge_duplicates:
        failure_codes.append(FailureCode.DUPLICATE_UNMERGE_RANGES)
    if unmerge_extras:
        failure_codes.append(FailureCode.EXTRA_UNMERGE_RANGES)

    return ValidationResult(
        expected=tuple(sorted(expected_set, key=_cell_sort_key)),
        duplicates=duplicates,
        missing=missing,
        extra=extra,
        dynamic_extents=tuple(dynamic_extents),
        unmerge_duplicates=unmerge_duplicates,
        unmerge_extras=unmerge_extras,
        materialised_cells=materialised,
        preserved_cells=preserved,
        failure_codes=tuple(failure_codes),
    )


def _atomic_copy(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        dir=destination.parent,
        prefix=f".{destination.name}.",
        suffix=".tmp",
    )
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        shutil.copyfile(source, temporary_path)
        os.replace(temporary_path, destination)
    finally:
        temporary_path.unlink(missing_ok=True)


def _failure_result(
    input_path: Path,
    output_path: Path,
    failure_codes: tuple[FailureCode, ...],
    *,
    validation: ValidationResult | None = None,
    error: str | None = None,
) -> WriteResult:
    if input_path.resolve() != output_path.resolve():
        _atomic_copy(input_path, output_path)
    return WriteResult(
        output_path=output_path,
        success=False,
        failure_codes=failure_codes,
        validation=validation,
        error=error,
    )


def _is_finite_excel_number(value: object) -> bool:
    """Return whether a value is a finite Excel number rather than a boolean."""

    return (isinstance(value, int) and not isinstance(value, bool)) or (
        isinstance(value, float) and math.isfinite(value)
    )


def _writer_round_trip_number(value: int | float) -> int | float:
    """Replay openpyxl's numeric writer and reader conversion exactly."""

    serialised = safe_string(value)
    if "." in serialised or "e" in serialised.lower():
        return float(serialised)
    return int(serialised)


def _assigned_value_round_trips(
    expected_value: MaterialisedValue,
    expected_data_type: str,
    actual_value: object,
    actual_data_type: str,
    epoch: dt.datetime,
) -> bool:
    """Compare an assigned value with the exact value openpyxl reloads."""

    if actual_value == expected_value and actual_data_type == expected_data_type:
        return True
    if expected_data_type != "n" or not _is_finite_excel_number(expected_value):
        return False

    canonical_number = _writer_round_trip_number(expected_value)
    if actual_data_type == "n" and _is_finite_excel_number(actual_value):
        return actual_value == canonical_number
    if actual_data_type == "d" and isinstance(
        actual_value,
        (dt.datetime, dt.date, dt.time, dt.timedelta),
    ):
        expected_temporal = from_excel(
            canonical_number,
            epoch,
            timedelta=isinstance(actual_value, dt.timedelta),
        )
        return actual_value == expected_temporal
    return False


def _verify_saved_workbook(
    workbook_path: Path,
    *,
    created_sheets: tuple[str, ...],
    assigned_cells: tuple[tuple[str, str, MaterialisedValue, str], ...],
    preserved_cells: tuple[tuple[str, str, PreservedCellSignature], ...],
    unmerged_ranges: tuple[RangeKey, ...],
) -> None:
    """Reopen a candidate output and verify all authored structural and cell changes."""

    try:
        verified = openpyxl.load_workbook(workbook_path, data_only=False)
    except Exception as exc:
        raise ContractError(
            FailureCode.WORKBOOK_WRITE_FAILED,
            "saved workbook could not be reopened",
        ) from exc
    try:
        missing_sheets = [sheet for sheet in created_sheets if sheet not in verified.sheetnames]
        if missing_sheets:
            raise ContractError(
                FailureCode.WORKBOOK_WRITE_FAILED,
                "saved workbook is missing an exact created sheet",
            )
        for sheet, coordinate, expected_value, expected_data_type in assigned_cells:
            actual = verified[sheet][coordinate]
            if not _assigned_value_round_trips(
                expected_value,
                expected_data_type,
                actual.value,
                actual.data_type,
                verified.epoch,
            ):
                raise ContractError(
                    FailureCode.WORKBOOK_WRITE_FAILED,
                    f"saved cell did not round-trip at {sheet}!{coordinate}",
                )
            if (
                isinstance(expected_value, str)
                and expected_value.startswith("=")
                and actual.data_type != "f"
            ):
                raise ContractError(
                    FailureCode.WORKBOOK_WRITE_FAILED,
                    f"saved formula type did not round-trip at {sheet}!{coordinate}",
                )
        for sheet, coordinate, expected_signature in preserved_cells:
            actual = verified[sheet][coordinate]
            if _preserved_cell_signature(actual) != expected_signature:
                raise ContractError(
                    FailureCode.WORKBOOK_WRITE_FAILED,
                    f"preserved cell changed during save at {sheet}!{coordinate}",
                )
        for key in unmerged_ranges:
            if key.cell_range in {str(item) for item in verified[key.sheet].merged_cells.ranges}:
                raise ContractError(
                    FailureCode.WORKBOOK_WRITE_FAILED,
                    f"merged range remained after save at {key.as_text()}",
                )
    finally:
        verified.close()


def _canonical_formula_value(value: object) -> object:
    """Return a stable value for formula objects that compare by identity."""

    if isinstance(value, ArrayFormula):
        return ("array", value.ref, value.text)
    if isinstance(value, DataTableFormula):
        # Iteration follows openpyxl's own serialised attributes.  It also
        # normalises constructor booleans and reloaded XML strings alike.
        return ("data_table", tuple(value))
    return value


def _preserved_cell_signature(cell: object) -> PreservedCellSignature:
    """Capture round-trip-equivalent content and presentation for a preserve.

    openpyxl's writer serialises finite floats through ``safe_string`` and can
    normalise three equivalent blank encodings and an all-zero style array.
    Canonicalising exactly those cases avoids rejecting its own saved output;
    every meaningful value, type, style, format, link and comment remains exact.
    """

    style = getattr(cell, "_style", None)
    style_signature = tuple(style) if style is not None else None
    if style_signature is not None and not any(style_signature):
        style_signature = None

    value = _canonical_formula_value(getattr(cell, "value", None))
    data_type = getattr(cell, "data_type", None)
    if value in (None, "") and data_type in {"n", "s", "inlineStr"}:
        value = None
        data_type = "n"
    elif (isinstance(value, int) and not isinstance(value, bool)) or (
        isinstance(value, float) and math.isfinite(value)
    ):
        value = safe_string(value)

    hyperlink = getattr(cell, "hyperlink", None)
    hyperlink_signature = None
    if hyperlink is not None:
        hyperlink_signature = (
            hyperlink.target,
            hyperlink.location,
            hyperlink.tooltip,
            hyperlink.display,
        )
    comment = getattr(cell, "comment", None)
    comment_signature = None
    if comment is not None:
        comment_signature = (
            comment.text,
            comment.author,
            comment.width,
            comment.height,
        )
    return (
        type(cell).__name__,
        value,
        data_type,
        style_signature,
        getattr(cell, "number_format", None),
        hyperlink_signature,
        comment_signature,
    )


def write_response_atomic(
    task: Mapping[str, object],
    response: SpreadsheetResponse | Mapping[str, object] | str | bytes | bytearray,
    input_path: str | Path,
    output_path: str | Path,
) -> WriteResult:
    """Install a complete response atomically, otherwise copy the input bytes.

    The source workbook is opened directly for every call.  Explicit missing
    target sheets are created only after validation succeeds.  Unqualified
    missing ``answer_sheet`` values never cause sheet creation.
    """

    source = Path(input_path)
    destination = Path(output_path)
    if not source.is_file():
        return WriteResult(
            output_path=destination,
            success=False,
            failure_codes=(FailureCode.INVALID_INPUT_WORKBOOK,),
            error="input workbook is not a file",
        )
    if source.resolve() == destination.resolve():
        return WriteResult(
            output_path=destination,
            success=False,
            failure_codes=(FailureCode.INVALID_OUTPUT_PATH,),
            error="input_path and output_path must differ",
        )

    try:
        parsed = response if isinstance(response, SpreadsheetResponse) else parse_response(response)
    except ContractError as exc:
        return _failure_result(
            source,
            destination,
            (exc.code,),
            error=str(exc),
        )

    try:
        workbook = openpyxl.load_workbook(source, data_only=False)
    except Exception as exc:
        # The file may still be a valid opaque submission artefact.  Preserve it
        # byte-for-byte rather than leaving a partial or stale destination.
        return _failure_result(
            source,
            destination,
            (FailureCode.INVALID_INPUT_WORKBOOK,),
            error=f"{type(exc).__name__}: {exc}",
        )

    try:
        contract = build_target_contract(task, workbook)
        validation = validate_response(contract, parsed, workbook=workbook)
        if not validation.ok:
            workbook.close()
            return _failure_result(
                source,
                destination,
                validation.failure_codes,
                validation=validation,
            )

        created_sheets: list[str] = []
        for target_range in contract.ranges:
            if (
                target_range.resolution is SheetResolution.CREATED
                and target_range.sheet not in workbook.sheetnames
            ):
                created = workbook.create_sheet(title=target_range.sheet)
                if created.title != target_range.sheet:
                    raise ContractError(
                        FailureCode.INVALID_TARGET_SHEET,
                        f"could not create exact target sheet {target_range.sheet!r}",
                    )
                created_sheets.append(target_range.sheet)

        preserved_cells = tuple(
            (
                key.sheet,
                key.cell,
                _preserved_cell_signature(workbook[key.sheet][key.cell]),
            )
            for key in validation.preserved_cells
        )

        unmerged_ranges = tuple(
            RangeKey(action.sheet, action.range) for action in parsed.unmerge_ranges
        )
        for action in parsed.unmerge_ranges:
            workbook[action.sheet].unmerge_cells(action.range)

        assigned_cells: list[tuple[str, str, MaterialisedValue, str]] = []
        for cell in validation.materialised_cells:
            target_cell = workbook[cell.sheet][cell.cell]
            if isinstance(target_cell, MergedCell):
                if cell.value is None:
                    continue
                raise ContractError(
                    FailureCode.WORKBOOK_WRITE_FAILED,
                    f"cannot write a non-anchor merged cell at {cell.sheet}!{cell.cell}",
                )
            target_cell.value = cell.value
            assigned_cells.append((cell.sheet, cell.cell, cell.value, target_cell.data_type))

        destination.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temporary_name = tempfile.mkstemp(
            dir=destination.parent,
            prefix=f".{destination.name}.",
            suffix=destination.suffix or ".xlsx",
        )
        os.close(descriptor)
        temporary_path = Path(temporary_name)
        try:
            workbook.save(temporary_path)
            workbook.close()
            _verify_saved_workbook(
                temporary_path,
                created_sheets=tuple(created_sheets),
                assigned_cells=tuple(assigned_cells),
                preserved_cells=preserved_cells,
                unmerged_ranges=unmerged_ranges,
            )
            os.replace(temporary_path, destination)
        finally:
            temporary_path.unlink(missing_ok=True)
        return WriteResult(
            output_path=destination,
            success=True,
            failure_codes=(),
            validation=validation,
        )
    except ContractError as exc:
        workbook.close()
        return _failure_result(
            source,
            destination,
            (exc.code,),
            error=str(exc),
        )
    except Exception as exc:
        workbook.close()
        return _failure_result(
            source,
            destination,
            (FailureCode.WORKBOOK_WRITE_FAILED,),
            error=f"{type(exc).__name__}: {exc}",
        )


def make_trace_record(
    *,
    step: int,
    model: str,
    prompt: str | None,
    response: str | None,
    input_tokens: int | None,
    output_tokens: int | None,
    latency_ms: int | None,
    write_result: WriteResult | None = None,
    error: str | None = None,
) -> dict[str, object]:
    """Return a stable JSON-serialisable model-call trace record."""

    failure_codes = []
    if write_result is not None:
        failure_codes = [code.value for code in write_result.failure_codes]
        error = error or write_result.error
    return {
        "step": step,
        "model": model,
        "prompt": prompt,
        "response": response,
        "input_tokens": input_tokens,
        "output_tokens": output_tokens,
        "latency_ms": latency_ms,
        "error": error,
        "failure_codes": failure_codes,
    }


def make_prediction_record(
    task_id: str | int,
    output: str | Path,
    write_result: WriteResult,
) -> dict[str, object]:
    """Return a stable ``predictions.jsonl`` record for one task."""

    return {
        "id": str(task_id),
        "output": Path(output).as_posix(),
        "status": write_result.status,
        "failure_codes": [code.value for code in write_result.failure_codes],
    }


__all__ = [
    "CellKey",
    "ContractError",
    "DynamicExtent",
    "EXCEL_MAX_CELL_CHARACTERS",
    "EXCEL_MAX_COLUMN",
    "EXCEL_MAX_FORMULA_CHARACTERS",
    "EXCEL_MAX_ROW",
    "ExcelDateTimeValue",
    "ExcelDateValue",
    "FailureCode",
    "MAX_EXPANDED_RESPONSE_CELLS",
    "MAX_EXPANDED_RESPONSE_CHARACTERS",
    "MaterialisedCell",
    "MaterialisedValue",
    "RangeKey",
    "ResponseCell",
    "ResponseFill",
    "ResponsePreserveRange",
    "ResponseUnmergeRange",
    "ScalarValue",
    "SheetResolution",
    "SpreadsheetResponse",
    "TargetContract",
    "TargetProvenance",
    "TargetRange",
    "ValidationResult",
    "WriteResult",
    "build_target_contract",
    "make_prediction_record",
    "make_trace_record",
    "normalise_coordinate",
    "parse_response",
    "validate_response",
    "write_response_atomic",
]
